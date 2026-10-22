#!/usr/bin/env python3
"""Standalone scientific figures and supervisor-facing PDF/DOCX from computed tables."""
from pathlib import Path
from xml.sax.saxutils import escape
import json,shutil,textwrap,re
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap,Normalize
from matplotlib.collections import LineCollection
from pyproj import Transformer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak,KeepTogether
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results';FIG=ROOT/'figures';DELIVERY=ROOT/'delivery'
TYPES=['decline:deconcentration','decline:stable_share','decline:concentration','stable:deconcentration','stable:stable_share','stable:concentration','growth:deconcentration','growth:stable_share','growth:concentration']
LABELS={'decline':'Сокращение','stable':'Стабильность','growth':'Рост','deconcentration':'Доля центра ↓','stable_share':'Доля центра ≈','concentration':'Доля центра ↑'}
PALETTE=['#a5627f','#b89eaa','#7c4662','#d1b55b','#d4cec2','#9c8c4e','#739ba7','#a0c4c9','#27667e']
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white','savefig.facecolor':'white','axes.labelcolor':'#23333d','text.color':'#23333d','axes.titlesize':14,'svg.fonttype':'none'})

def save(fig,name):
 FIG.mkdir(exist_ok=True)
 for ext in ['png','svg','pdf']:fig.savefig(FIG/f'{name}.{ext}',dpi=300,bbox_inches='tight')
 plt.close(fig)
 return name

def build_figures(mem,sys,sensitivity,reg,intervals):
 projection=Transformer.from_crs('EPSG:4326','+proj=aea +lat_1=50 +lat_2=70 +lat_0=56 +lon_0=100 +datum=WGS84',always_xy=True)
 x,y=projection.transform(mem.lon.to_numpy(),mem.lat.to_numpy());x,y=np.asarray(x)/1000,np.asarray(y)/1000
 world=ROOT/'coordination/world_outline.geojson'
 if not world.exists() and Path('/workspace/scratch/sgrt_map/world.geojson').exists():shutil.copy2('/workspace/scratch/sgrt_map/world.geojson',world)
 def base(ax):
  if world.exists():
   for f in json.loads(world.read_text())['features']:
    if f['properties'].get('ADMIN')!='Russia':continue
    geom=f['geometry'];polys=geom['coordinates'] if geom['type']=='MultiPolygon' else [geom['coordinates']]
    for poly in polys:
     coords=np.asarray(poly[0]);px,py=projection.transform(coords[:,0],coords[:,1]);ax.fill(np.asarray(px)/1000,np.asarray(py)/1000,fc='#f6f5f1',ec='#b8b6ae',lw=.45,zorder=0)
  ax.set_aspect('equal');ax.set_xlim(-5200,3800);ax.set_ylim(-1250,3650);ax.axis('off')
 for level in [1,2,3,4]:
  fig,ax=plt.subplots(figsize=(15,8));base(ax)
  valid=mem[f'center_id_l{level}'].notna().to_numpy();ax.scatter(x[~valid],y[~valid],s=3,c='#b8b8b8',rasterized=True,label='Нет назначения')
  if level==1:
   types=sys[sys.level==1].set_index('center_id').demographic_type.to_dict();tt=mem.center_id_l1.map(types)
   other=valid&~tt.isin(TYPES).to_numpy();ax.scatter(x[other],y[other],s=.7,c='#65757c',alpha=.6,rasterized=True)
   for t,col in zip(TYPES,PALETTE):
    a=tt.eq(t).to_numpy();ax.scatter(x[a],y[a],s=.5,c=col,alpha=.6,rasterized=True)
  else:
   ax.scatter(x[valid],y[valid],s=.4,c='#c3ccd1',rasterized=True)
  centers=sys[sys.level==level];cx,cy=projection.transform(centers.lon.to_numpy(),centers.lat.to_numpy())
  ax.scatter(np.asarray(cx)/1000,np.asarray(cy)/1000,s=4+centers.own_functional_mass.values*1.3,facecolors='none' if level==1 else '#2a6077',edgecolors='#20475e',linewidths=.35,alpha=.75)
  ax.set_title(f'Модельные СГРТ: уровень {level} · {len(centers):,} систем'.replace(',',' '),loc='left',pad=12)
  fig.text(.13,.10,'Фиксированная современная функциональная сетка; геодезические расстояния, не дорожное время.\nАльберс 50°/70°, центральный меридиан100°. Серый — вне модельного назначения. Ось142570НП.',fontsize=9,color='#59666e')
  from matplotlib.lines import Line2D
  if level==1:
   legend=[Line2D([0],[0],marker='o',lw=0,color=col,label=LABELS[t.split(':')[0]]+' / '+LABELS[t.split(':')[1]],markersize=5) for t,col in zip(TYPES,PALETTE)]
   legend += [Line2D([0],[0],marker='o',lw=0,color='#65757c',label='Федеральный контекст / несопоставимый тип',markersize=5),Line2D([0],[0],marker='o',lw=0,color='#b8b8b8',label='Нет назначения',markersize=5),Line2D([0],[0],marker='o',lw=0,markerfacecolor='none',color='#20475e',label='Кольцо: центр; размер растёт с массой',markersize=8)]
   fig.legend(handles=legend,loc='lower left',bbox_to_anchor=(.12,-.055),ncol=3,fontsize=8,frameon=False)
  else:
   fig.text(.13,.035,'Синие круги — центры; размер символа растёт с собственной функциональной массой. Светлые точки — НП с назначением.',fontsize=9)
  save(fig,f'01_national_level{level}')
 low=sys[(sys.level==1)&sys.has_rural_periphery&sys.center_axis_kind.eq('ordinary')].copy()
 grid=np.zeros((3,3),int)
 for i,tr in enumerate(['decline','stable','growth']):
  for j,sh in enumerate(['deconcentration','stable_share','concentration']):grid[i,j]=low.demographic_type.eq(tr+':'+sh).sum()
 fig,ax=plt.subplots(figsize=(10,6));ax.imshow(grid,cmap='Blues',vmin=0,vmax=grid.max());ax.set_xticks(range(3),['Доля центра ↓','Доля центра ≈','Доля центра ↑']);ax.set_yticks(range(3),['Система сокращается','Система стабильна','Система растёт'])
 for (i,j),v in np.ndenumerate(grid):ax.text(j,i,f'{v:,}\n{100*v/len(low):.1f}%'.replace(',',' '),ha='center',va='center',fontsize=16,color='white' if v>grid.max()*.5 else '#23333d')
 ax.set_title(f'Демографическая типология2002–2021 · {len(low)} низовых систем с сельским окружением',loc='left',pad=18);fig.text(.12,.01,'Стабильность:±0,25%/год; доля центра:±2п.п. Пороги сценарные. Современный фиксированный состав.',fontsize=9);fig.tight_layout(rect=[0,.06,1,1]);save(fig,'02_typology')
 fig,ax=plt.subplots(figsize=(10,6));ss=sensitivity.sort_values('changed');ax.barh(ss.scenario,100*ss.changed/142570,color='#2a6077');ax.set_xlabel('НП с другим центром относительно основной модели, % оси142570');ax.set_title('Насколько назначение зависит от спецификации',loc='left');ax.grid(axis='x',alpha=.15);fig.tight_layout();save(fig,'03_sensitivity')
 fig,ax=plt.subplots(figsize=(10,6));
 for period,marker,offset in [('2002-2021','o',-.12),('2010-2021','s',.12)]:
  z=intervals[intervals.period.eq(period)]
  for j,(_,r) in enumerate(z.iterrows()):
   p=r.relative_rmse_reduction*100;lo=r.region_bootstrap_p025*100;hi=r.region_bootstrap_p975*100
   ax.errorbar(p,j+offset,xerr=[[p-lo],[hi-p]],fmt=marker,c='#2a6077' if offset<0 else '#b68642',capsize=5,label=period if j==0 else None)
 ax.axvline(0,c='#777',lw=1);ax.set_yticks([0,1],['M2 кM1: аграрный контекст','M3 кM1: +организации2011']);ax.set_xlabel('Относительное снижение RMSE, %; 95% региональный bootstrap');ax.set_title('Прибавка признаков численно положительна, но интервалы включают ноль',loc='left');ax.legend();fig.tight_layout();save(fig,'04_model_comparison')
 good=low[(low.center_population_2002>0)&(low.center_population_2021>0)&(low.rural_external_population_2002>0)&(low.rural_external_population_2021>0)].copy()
 good['center_rate']=100*np.expm1(np.log(good.center_population_2021/good.center_population_2002)/19);good['rural_rate']=100*np.expm1(np.log(good.rural_external_population_2021/good.rural_external_population_2002)/19)
 fig,ax=plt.subplots(figsize=(9,7));ax.scatter(good.center_rate,good.rural_rate,s=10,alpha=.35,color='#2a6077');ax.axhline(0,c='#aaa',lw=.8);ax.axvline(0,c='#aaa',lw=.8);ax.set_xlabel('Среднегодовое изменение населения центра, %');ax.set_ylabel('Среднегодовое изменение сельского окружения, %');ax.set_title(f'Центр и сельское окружение меняются по-разному · n={len(good)}',loc='left');ax.set_xlim(-10,10);ax.set_ylim(-10,10);fig.text(.12,.01,'Показан диапазон±10%/год; все значения, включая крайние, сохранены в таблице систем.\nЭто демография фиксированного состава, не наблюдённое переселение между его частями.',fontsize=9);fig.tight_layout(rect=[0,.07,1,1]);save(fig,'05_center_rural_dynamics')
 # Representative-scale insets selected in advance by geographic setting, not agreement with documents.
 fig,axes=plt.subplots(1,3,figsize=(16,6))
 for ax,(title,box) in zip(axes,[('Нечерноземье: Тверь',(32,39,55,59)),('Урал',(57,63,53,60)),('Юг Западной Сибири',(76,88,50,57))]):
  mask=mem.lon.between(box[0],box[1])&mem.lat.between(box[2],box[3]);z=mem[mask]
  # A projected coordinate view, equal axes, no artificial polygon boundaries.
  ax.scatter(x[mask],y[mask],s=2,c='#9daab1',alpha=.6)
  for _,r in z.iterrows():
   j=int(r.center_index_l1)
   if j<0 or j>=len(mem):continue
   i=r.name
   if j==i:continue
   ax.plot([x[i],x[j]],[y[i],y[j]],lw=.15,c='#2a6077',alpha=.12)
  cc=z[z.candidate_center];ax.scatter(x[cc.index],y[cc.index],s=9,c='#254f63');ax.set_aspect('equal');ax.set_title(title);ax.set_xlabel('Альберс, км');ax.set_ylabel('Альберс, км')
 fig.suptitle('Региональные увеличения низовых модельных связей',fontsize=16);fig.text(.1,.01,'Линии — модельные отношения; не дороги, поездки или юридические границы.',fontsize=10);fig.tight_layout(rect=[0,.06,1,.94]);save(fig,'06_regional_insets')

def build_text(r,ext,scaffold,reg,intervals):
 n=r['systems_with_rural_periphery_l1'];shr=100*r['type_counts_l1_rural']['decline:concentration']/n
 rows=[]
 rows.append(('h','Динамическая типология сельско-городских районов потенциального тяготения России'))
 rows.append(('p','Предварительный национальный проход по исходной концепции Word ·21 октября2026. Материал для обсуждения с научным руководителем.'))
 rows.append(('h2','Что получено и что уже можно обсуждать'))
 rows.append(('p',f'Рассчитана многоуровневая система на фиксированной оси142570 населённых пунктов. В основной функциональной спецификации модельное назначение получили139287НП;3283 остались вне радиуса доступных кандидатов. На четырёх ступенях получены4891,2610,570 и161 система. Из низовых систем2006 имеют внешнее сельское окружение и обычный сопоставимый центр; именно они образуют основную демографическую типологию. Это национальный вычислительный результат, но не подтверждённая сеть реального обслуживания.'))
 rows.append(('p',f'В {r["type_counts_l1_rural"]["decline:concentration"]} из{n} сельских систем ({shr:.1f}%) сокращение общего населения2002–2021 сочетается с увеличением доли центра. Такое сочетание описывает концентрацию на фоне сжатия. Оно само по себе не устанавливает переезд жителей из окружения в центр.'))
 rows.append(('p','Ключевой отрицательный результат: основная модель плохо воспроизводит доступные нормативные школьные связи —2 совпадения из98 датированных утверждений для многопрофильного центра и0 для отдельного образовательного назначения. Среди16 известных поставщиков в модельных кандидатах присутствуют лишь3. Вариант пока нельзя использовать как готовые границы реального обслуживания или основу адресного распределения господдержки. Возможные причины: неполнота сельских функций, различия уровней услуг и дат, расстояние вместо дорожного времени, а также сама спецификация массы.'))
 rows.append(('p',f'Дополнительно построена близкая исходной постановке Word демографическая альтернатива: {scaffold["ordinary_candidate_centers"]} потенциальных обычных центров, {scaffold["rural_systems_l1"]} низовых систем с сельским окружением; уровни '+ ' / '.join(str(x) for x in scaffold['levels'].values())+'. Крупнейшие поселения муниципалитетов явно помечены provisional; они не объявляются доказанными райцентрами. Альтернатива показывает цену исходных предположений, не заменяет функциональную проверку.'))
 rows.append(('h2','1. Исследовательский вопрос и связь с исходным Word'))
 rows.append(('p','Исследование проверяет, как конфигурация системы «центр—окружение», её положение и доступные функциональные, экономические и аграрные признаки соотносятся с демографической устойчивостью. Сохраняются многоуровневое районирование снизу вверх, отдельная демография частей системы, функциональная неоднородность, сравнениеМ1/М2/М3 и картографическая типология. Пригородность/полицентричность и девять содержательных типов Word пока остаются проверяемыми гипотезами: близкие модельные веса не доказывают реальные потоки, а знак демографической динамики не доказывает этнический, дачный или вахтовый характер.'))
 rows.append(('p','Круговая калибровка устранена: собственная масса центра вычисляется до назначения окружения; полученный моделью размер района не используется ни для обучения массы, ни для её ранга. Основная масса состоит из11 групп функций, без населения и демографического исхода. Это явное изменение Word. В отдельных сценариях возвращены независимая людность2002 и смешанная масса, чтобы проверить влияние этого решения.'))
 rows.append(('h2','2. Источники, ось и наблюдаемость'))
 rows.append(('p','Побайтно восстановлен master142570×914, SHA2568c98d5a113eec4c3bf131232b1d7e5c6362dc3a50ee484e964420de58c377f85. Суммы обычной оси2002/2010/2021:127397236 /124684805 /125132003 человек. Это знаменатели выбранной сопоставляемой оси, не население всей РФ. Все21246 явных нулей2010 сохранены; качество промежуточных чисел и доля защищённых наблюдений доступны по каждой системе. Москва и Санкт-Петербург — отдельные контекстные территориальные узлы, не добавленные повторно к обычным НП.'))
 rows.append(('table',(['Слой','Фактически доступно','Граница интерпретации'],[['Организации RFSD','15лет2011–2025; собственные коды НП','Юрлица, не филиалы/занятые/реальные поездки'],['Школы','7267 strictPOI в397НП;386НП от10тыс. жителей','Крайне неполная сельская сеть; нет полной мощности'],['Медицина','CoarsePlaces385НП; с документным пилотом433НП','Нет полного национального разделения ФАП/стационаров'],['WB21октября','91043 исходных записи;74763 baseline;16172НП','Не все типы — ПВЗ; версии каталога не равны открытиям'],['WBтипы8/34','7991НП; базовая масса закреплена на19октября','Снимки разделены; свежий вариант — чувствительность'],['Аграрный контекст','Муниципальные площади/профили/производство','Повторённые поНП гектары не суммируются'],['1989','19863 известных значения в исходной оси','Частичный разнородный исторический слой'] ])))
 rows.append(('p','Официальный справочник подразделений ФРМО найден по описанию, но архив с физическими адресами и мощностями получить не удалось:503/502/403 и тайм-ауты задокументированы. Национальные образовательные порталы и ограниченные Overpass-попытки также не дали новых массовых строк. Это предел получения в данном проходе, не доказательство отсутствия данных. Полный Ozon-каталог не получен; медицинские центры «ДокторОЗОН» не превращены вПВЗ.'))
 rows.append(('h2','3. Модельный эксперимент и иерархия'))
 rows.append(('p','Внутри каждой функции источниковые счётчики объединяются максимумом, затемlog(1+n), с ограничением верхнего1% только в индексе. Масса—сумма11 преобразованных компонент с заранее равными весами. Такой индекс измеряет модельную представленность функций; не мощность, качество или незаменимость обслуживания. Кандидат имеет≥2 наблюдаемые группы либо физическую школу/больницу. Все кандидаты сохраняют собственное ядро по конструктивному правилу.'))
 rows.append(('p','Низовой потенциалGij=Mj×exp(−dij/15км); рассматриваются все кандидаты в75км без приближения ближайшими64. Шаги иерархии используют50/90/98квантили собственной массы и радиусы100/200/400км. Системы объединяются целиком; для каждой ступени все три суммы населения сохранены. Квантили—сценарные ранги, не доказанные уровни специализированной помощи. Дорожная матрица и изохроны не построены: используется корректное сферическое расстояние и сохраняется неопределённость барьеров.'))
 rows.append(('p','Из4890 обычных ядер2880 не имеют внешнего обычногоНП. У2795 ядер свободная гравитационная модель выбрала бы другой центр; собственный состав сохранён принудительно. Поэтому число4891 включает одиночные контейнеры и контекстный узел, а не4891 доказанныйСГРТ. Результат чувствителен к правилу ядра; модельная неоднозначность и свободный победитель сохранены отдельно.'))
 for l in [1,2,3,4]:rows.append(('image',(f'01_national_level{l}',f'Уровень{l}: современная модельная география. КонтурNaturalEarth—контекст, не переписные границы.')))
 rows.append(('h2','4. Демографические траектории'))
 rows.append(('p','Население центра, всего внешнего окружения, сельского окружения и системы рассчитано для2002/2010/2021 на одном современном фиксированном составе. Это ретроспективный аналитический разрез: не утверждается, что нынешние функции и границы существовали в2002. Три класса темпа системы(±0,25%/год) и три класса изменения доли центра(±2п.п.) дают девять демографических сочетаний; соседние пороги проверены отдельно.'))
 rows.append(('image',('02_typology','Числа и доли относятся только к2006 низовым системам с сельским окружением и обычным центром.')))
 rows.append(('image',('05_center_rural_dynamics','Ряды с нулевыми концами сохранены отдельно; они не дают определённого логарифмического темпа.')))
 rows.append(('p','Для1989 выполнено сравнение на одном и том же наборе известных НП, с долей их населения2002 в системе. Неизвестные значения не превращены в нули; национальная историческая полнота и сопоставимость границ не объявляются. Дистанционные градиенты вычислены для периферийных сельскихНП; они описывают ассоциации с расстоянием, не транспортную причинность.'))
 rows.append(('h2','5. Устойчивость, альтернативы и временная структура'))
 rows.append(('p','По всей оси проверены11 альтернатив: радиусы50/100км, масштабы затухания10/25км, степенная модель, ближайший центр, условные связки, только физические источники, только регистрационная масса2021 и варианты сP2002. Это сценарная неопределённость, не доверительный интервал истинной границы. Только21265НП сохраняют один и тот же центр во всех11 различающихся сценариях; критерий строгий и включает принципиально другие наборы источников.'))
 rows.append(('image',('03_sensitivity','Сравнение с основной моделью на одной полной оси. Незначительное изменение свежегоWBне устраняет неопределённость формы модели.')))
 rows.append(('p','Отдельное обновлениеWB19→21октября изменило6 назначений в сопоставлении только обычных центров. Временная организационная история сохранена за все15лет; модели по юридическим массам2011/2016/2021/2025 показывают чувствительность сетки к наблюдаемым регистрациям. Смены юридических записей и кодов не названы закрытиями услуг или реальным перехватом функций.'))
 rows.append(('p','Документная медицина:49 записей2025/2026, в двух обследованных сетях19 подразделений с работником в обоих срезах; прикреплённое население6422→6286. Это небольшой сопоставимый пример, не национальная динамика. Образование:98 внешних нормативных утверждений,74 исходныхНП,16 поставщиков; сохранены конкретные смены назначения, включаяНовуюТуру/Платину. Это проверки после разработки, не слепая отложенная выборка.'))
 rows.append(('image',('06_regional_insets','Региональные увеличения выбраны по географическим условиям, не по совпадению с известными ответами.')))
 rows.append(('h2','6. Аграрная среда и сравнениеМ1/М2/М3'))
 rows.append(('p','Для каждого низового района построено двустороннее соответствие с муниципальной сеткой по населению2002.603 пересечения проходят≥80% с обеих сторон; после требований к сельской периферии и полным ранним признакам в регрессиях остаются304 системы48регионов. Гектары и выпуск сохраняют муниципальный смысл; источники не разрезаются пропорционально населению деревень. Этот отбор ограничивает обобщение.'))
 rows.append(('p','М1: размер центра/сельского окружения2002, доля центра и среднее расстояние. М2 добавляет муниципальные площади и производство2010. М3 добавляет регистрационные свидетельства организаций2011. Все модели сравниваются на одинаковых304 строках; стандартизация обучается внутри пяти разнесённых по регионам блоков. Использован заранее фиксированныйRidgeα=1. Сетка современная, муниципальная принадлежность по справочнику2021 и RFSDретроспективного выпуска: причинный эффект и независимый исторический прогноз не установлены.'))
 rr=[]
 for _,a in reg.iterrows():rr.append([a.period,a.model,str(a.n),f'{100*a.rmse_annual_log:.3f}',f'{100*a.mae_annual_log:.3f}'])
 rows.append(('table',(['Период','Модель','N','RMSE×100','MAE×100'],rr)))
 rows.append(('p','В2002–2021 добавление аграрного контекста численно снизилоRMSEна3,1%, а весьМ3—на5,1%; для2010–2021—на2,3% и4,4%. Но все95%-интервалы региональногоbootstrapвключают ноль. Устойчивое улучшение не подтверждено; результата недостаточно для доказательства гипотезы. Агроклимат, полная конфигурационная спецификация, жильё и собственные миграционные компоненты ещё не включены.'))
 rows.append(('image',('04_model_comparison','1000парныхbootstrap-перевыборок регионов; интервалы не покрывают ошибку источников и границ.')))
 rows.append(('h2','7. Литературная опора и изменения методики'))
 rows.append(('p','Свежепрочитаны тематические оригинальные фрагментыПоляна,Трейвиша,Родомана,Покшишевского,Лаппо иХорева. Матрица литературы содержит12позиций сURL, страницами, состоянием чтения ихэшами. Маергойз,Колосовский иБаранский представлены явно обозначенными вторичными изложениями; оригиналы не объявлены прочитанными. Полян обосновывает разделение собственной массы/внешнего потенциала и ограниченность опорности по людности; Трейвиш—различие масштаба и ранга; Хорев—различные циклы поездок. Ни одному автору не приписываются наши численные квантили или радиусы.'))
 rows.append(('p','ВWordмасса включала демографическую динамику, калибровалась на собственное черновое окружение, а типы частично описывались заранее. Здесь круговой контур разорван, демографический исход исключён из основной массы, девять содержательных интерпретаций не назначаются автоматически. Сохранены цель, многоуровневая конструкция, сравнение частей района и последовательные блоки анализа. Подробное требование→выход→ограничение:reviews/word_completion_assessment.csv.'))
 rows.append(('h2','8. Что можно показать научному руководителю'))
 rows.append(('p','На доступных данных уже можно показать национальную вычислимую постановкуСГРТ: четыре масштаба, полный состав исходныхНП, демографическую типологию и различие результатов при изменении функций/массы/доступности. Предварительный содержательный тезис—часто концентрация населения вцентре происходит на фоне сжатия системы. Второй тезис—одна совокупная зона не заменяет раздельные контуры образования, медицины и доступа к товарам. Третий—нынешние каталоги дают слишком смещённое наблюдение сельских услуг, и это эмпирически видно в отрицательной документной проверке.'))
 rows.append(('p','Следующий научно значимый шаг—получить массовые физические реестры школ/медицинских подразделений с уровнями и мощностями, добавить дорожную доступность и проверить модель на новых территориях, не использованных для её настройки. Затем можно обсуждать окончательные функциональные границы, полицентричность и адресную политику. На этом этапе карта—предварительная исследовательская модель, не готовая схема территориальной поддержки.'))
 rows.append(('h2','9. Проверки и воспроизводимость'))
 rows.append(('p','Работали пять запрошенных профильных агентов и дополнительный независимый вычислительный проверяющий. Последний выполнил186 независимых условий: массы, UID/нули, сферические победители, все родительские назначения, балансы, типология и регрессии по другому решателю, полный свежийWB-сценарий и демографическую альтернативу. Агент соответствияWordвыполнил отдельные проверки оси, слоёв, геометрии и содержания. Эти проверки подтверждают реализацию, а не истинность реальных границ. Позднее приложение содержит результаты внешней и scaffold-проверки.'))
 rows.append(('p','Основные файлы:results/memberships.parquet—все142570НП;systems_all_levels.parquet—системы всех уровней;hierarchy.csv—ребра;municipal_overlap.parquet—соответствие сеток;sensitivity.csv—альтернативы;regression_comparison.csv иpaired_error_intervals—сравнение моделей. Все команды и состав поставки—вREADME.md. Исходные слои иmainне изменены; результаты новой волны имеют отдельный namespace.'))
 rows.append(('h2','10. Прочитанные литературные источники'))
 bib=pd.read_csv(ROOT/'literature/bibliography.csv').fillna('')
 for _,b in bib.iterrows():
  state='оригинальный тематический фрагмент' if b.current_read_status.startswith('original') else 'вторичное изложение / оригинал не прочитан'
  rows.append(('ref',f'[{b.source_id}] {b.author}. {b.title}. {b.year}. Прочитано: {b.read_locator}. Статус: {state}. {b.url}'))
 return rows

def clean_prose(text):
 text=re.sub(r'(?<=[а-яА-ЯёЁ])(?=\d)|(?<=\d)(?=[а-яА-ЯёЁ])',' ',str(text))
 text=re.sub(r'(?<=[а-яё])(?=[A-Z])',' ',text)
 text=re.sub(r'(?<=[A-Za-z])(?=[А-Яа-яёЁ])',' ',text)
 text=re.sub(r'М ([123])',r'М\1',text)
 for a,b in [('142570','142 570'),('139287','139 287'),('21246','21 246'),('141796','141 796'),('16172','16 172'),('91043','91 043'),('74763','74 763'),('19863','19 863'),('21265','21 265')]:text=text.replace(a,b)
 return text

def render(rows,basename,brief=False):
 rows=[(typ,clean_prose(c) if typ in ['h','h2','p','ref'] else ((c[0],clean_prose(c[1])) if typ=='image' else c)) for typ,c in rows]
 pdfmetrics.registerFont(TTFont('DejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'));pdfmetrics.registerFont(TTFont('DejaVuBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
 styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyRU',fontName='DejaVu',fontSize=9.5,leading=14,spaceAfter=8));styles.add(ParagraphStyle(name='TitleRU',fontName='DejaVuBold',fontSize=20,leading=26,spaceAfter=16));styles.add(ParagraphStyle(name='HeadRU',fontName='DejaVuBold',fontSize=13,leading=18,spaceBefore=13,spaceAfter=8));styles.add(ParagraphStyle(name='CaptionRU',fontName='DejaVu',fontSize=8,leading=11,textColor=colors.HexColor('#54636b'),spaceAfter=12))
 story=[];doc=Document();normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
 for typ,content in rows:
  if typ in ['h','h2','p','ref']:
   style={'h':'TitleRU','h2':'HeadRU','p':'BodyRU','ref':'CaptionRU'}[typ];story.append(Paragraph(escape(content),styles[style]))
   if typ=='h':doc.add_heading(content,0)
   elif typ=='h2':doc.add_heading(content,1)
   else:doc.add_paragraph(content)
  elif typ=='image':
   name,caption=content;path=FIG/f'{name}.png'
   if not path.exists():continue
   from PIL import Image as PILImage
   w,h=PILImage.open(path).size;width=480;height=width*h/w
   if height>440:height=440;width=height*w/h
   story.append(Image(str(path),width=width,height=height));story.append(Paragraph(escape(caption),styles['CaptionRU']))
   doc.add_picture(str(path),width=Inches(6));doc.add_paragraph(caption,style='Caption')
  elif typ=='table':
   header,data=content;grid=[[Paragraph(escape(str(x)),styles['CaptionRU']) for x in row] for row in [header]+data]
   widths=[480/len(header)]*len(header);table=Table(grid,colWidths=widths,repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf0f3')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#a9bac4')),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]));story.append(table);story.append(Spacer(1,10))
   t=doc.add_table(rows=1,cols=len(header));t.style='Light Shading Accent 1'
   for cell,text in zip(t.rows[0].cells,header):cell.text=str(text)
   for row in data:
    for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
 def page(canvas,document):
  canvas.setFont('DejaVu',7);canvas.setFillColor(colors.HexColor('#687780'));canvas.drawString(54,28,'СГРТ · предварительный национальный проход ·21.10.2026');canvas.drawRightString(A4[0]-50,28,str(document.page))
 SimpleDocTemplate(str(DELIVERY/f'{basename}.pdf'),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=48,bottomMargin=48,title='Динамическая типология СГРТ: предварительный национальный проход',author='Исследовательский проект СГРТ').build(story,onFirstPage=page,onLaterPages=page)
 doc.save(DELIVERY/f'{basename}.docx')
 md=[]
 for typ,c in rows:
  if typ=='h':md.append('# '+c)
  elif typ=='h2':md.append('## '+c)
  elif typ in ['p','ref']:md.append(c)
  elif typ=='image':md.append(f'![{c[1]}](../figures/{c[0]}.png)')
  elif typ=='table':
   header,data=c;md.append('| '+' | '.join(header)+' |\n| '+' | '.join(['---']*len(header))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in data))
 (DELIVERY/f'{basename}.md').write_text('\n\n'.join(md)+'\n')

def run():
 DELIVERY.mkdir(exist_ok=True)
 mem=pd.read_parquet(OUT/'memberships.parquet');sys=pd.read_parquet(OUT/'systems_all_levels.parquet');sensitivity=pd.read_csv(OUT/'sensitivity.csv');reg=pd.read_csv(OUT/'regression_comparison.csv');intervals=pd.read_csv(OUT/'regression_paired_error_intervals.csv')
 build_figures(mem,sys,sensitivity,reg,intervals)
 r=json.loads((OUT/'model_receipt.json').read_text());ext=json.loads((OUT/'extended_check_receipt.json').read_text());scaffold=json.loads((OUT/'demographic_scaffold/receipt.json').read_text())
 rows=build_text(r,ext,scaffold,reg,intervals);render(rows,'SGRT_PRELIMINARY_REPORT_20261021')
 end=next(i for i,(t,c) in enumerate(rows) if t=='h2' and c.startswith('1.'))
 brief=rows[:end]+[('image',('02_typology','Предварительная типология фиксированной современной сетки.'))]+[('h2','Три вопроса для обсуждения'),('p','Какую функцию и уровень обслуживания принять за основу низового района? Какие массовые физические реестры реально доступны? Как организовать независимую проверку и дорожную доступность до окончательной делимитации?')]
 render(brief,'SGRT_SUPERVISOR_BRIEF_20261021',True)
 import fitz
 checks=[]
 for p in DELIVERY.glob('*.pdf'):
  pdf=fitz.open(p);text='\n'.join(page.get_text() for page in pdf);checks.append({'file':p.name,'pages':len(pdf),'text_characters':len(text),'has_cyrillic':'СГРТ' in text,'no_empty_pages':all(len(page.get_text().strip())>10 or len(page.get_images()) for page in pdf)})
 (DELIVERY/'render_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2));print(json.dumps(checks,ensure_ascii=False,indent=2))

if __name__=='__main__':run()
