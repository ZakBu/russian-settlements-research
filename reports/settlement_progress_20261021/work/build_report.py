from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
BASE=Path.cwd();R=BASE/'work/sources/research/sgrt_word_completion_20261021';F=BASE/'work/figures';OUT=BASE/'outputs';OUT.mkdir(exist_ok=True)
s=Path('work/draft.txt').read_text();parts=re.split(r'\n(\d+) ([^\n]+)\n',s);sections={}
for i in range(1,len(parts),3):sections[int(parts[i])]=(parts[i+1],parts[i+2].strip().split('\n'))
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7);sec.header_distance=Cm(.65);sec.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.06;normal.paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(8)

for st in doc.styles:
 for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.styles['Subtitle'].font.size=Pt(11);doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)
doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(11.5);doc.styles['Heading 2'].font.bold=True
doc.styles['Caption'].font.size=Pt(8.5);doc.styles['Caption'].paragraph_format.space_after=Pt(5)
h=sec.header.paragraphs[0];h.text='РАССЕЛЕНИЕ РОССИИ  ·  ДАННЫЕ И ХОД ИССЛЕДОВАНИЯ';h.runs[0].font.size=Pt(8);h.runs[0].font.color.rgb=RGBColor.from_string('606C76')
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run('21 октября 2026  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
for run in p.runs:run.font.size=Pt(8)
metadata=[]
def para(t,style=None):return doc.add_paragraph(t,style)
def heading(t):doc.add_heading(t,1)
def page(t):doc.add_page_break();heading(t)
def body(n,skip=()):
 for i,p in enumerate(sections[n][1]):
  if p and i not in skip:para(p)
def caption(t):para(t,'Caption')
def picture(path,maxh=17,width=17.6):
 im=Image.open(path);w,h=im.size;cw=min(width,maxh*w/h);p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();r.add_picture(str(path),width=Cm(cw));
 inline=r._r.xpath('.//wp:docPr');
 if inline:inline[0].set('descr',Path(path).stem.replace('_',' '))
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 pr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243C50' if ri==0 else ('F0F4F6' if ri%2==0 else 'FFFFFF'));tcpr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   tcpr.append(borders);m=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');m.append(e)
   tcpr.append(m)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
    for r in p.runs:r.font.size=Pt(9.2);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
 return t
# 1
para('Трансформация территориальной структуры расселения России в 2002–2021 гг.','Title')
para('Замысел исследования, собранные данные и предварительные результаты\nКраткий иллюстрированный отчет · 21 октября 2026 года','Subtitle')
heading('Что исследуем и что уже получилось')
body(1,skip=(2,3))
table(['Перепись','Учтено населения','Контрольный итог','Покрытие'],[['2002','144 637 473','145 166 731','99,64%'],['2010','142 172 049','142 856 536','99,52%'],['2021','144 329 319','144 699 673','99,74%']],[2,5.1,5.1,3.4])
caption('Таблица 1. Широкий слой базы [1]. Покрытие — доля населения на общей территории сравнения; итог 2021 приведен без Крыма и Севастополя. Это не доля поселений и не точность каждого числа.')
para('Для карт динамики доступны 144 527 историй поселений с принятыми координатами. История — связанные записи об одном населенном пункте за разные годы. Значения 2002 и 2021 годов есть у 143 446 историй; три переписных значения — у 142 765. Москва и Санкт-Петербург ведутся отдельными территориальными итогами.')
#2
page('Какие данные собраны')
body(2)
rows=[
['Переписи и координаты','2002, 2010, 2021','Поселение','144 527 историй на демографических картах; полные и неполные ряды'],
['Образование','Снимок 01.01.2021; дополнения 2022–2026','Учреждение → поселение','Школы, детсады, колледжи, вузы; адресные и приблизительные связи раздельны'],
['Медицина','Снимок 01.01.2021; современные проверки','Учреждение → поселение','ФАПы, амбулатории, поликлиники, больницы, аптеки; полнота сети не доказана'],
['Торговля и услуги','01.01.2021; отдельные современные каталоги','Объект → поселение','Магазины, рынки, банки, банкоматы, почта, администрация'],
['Доставка и общепит','WB 21.10.2026; Places; юрлица 2011–2025','Объект или код поселения','WB типы 8/34 — 7 991 НП; Ozon — только непроверенные сигналы названий'],
['Экономический контекст','2021; отраслевые ряды 2011–2025','Код поселения','Регистрационные записи; не рабочие места и не физические предприятия'],
['Демография и жилье','Преимущественно 2006–2022; миграция с 2011','Муниципалитет','Рождения, смерти, прирост, ввод жилья; не собственные числа деревни'],
['Сельское хозяйство','Ряды 2006–2022; доли 2008 и 2021','Муниципалитет / регион','Посевы, урожайность, выпуск; доля посевов 2021 по 84 субъектам'],
['Климат и почвы','Нормы 1970–2000 и 1981–2010; почвенные карты','Растр → координата','Тепло, осадки, увлажнение, почвенные классы; бонитет — эксперимент'],
['Транспорт','OSM 2026; пересборка по снимку 2021','Дорожный граф','Модельные времена поездки; не наблюденные маршруты людей'],
['Дополнительный контекст','1979–1989; 2002; 2021; 2022–2024; разные даты','Поселение / город / территория','Пол, национальность, история населения, рекреация, моногорода, добыча']]
table(['Блок','Даты','Уровень данных','Что можно использовать'],rows,[3.15,3.3,2.7,8.45])
caption('Таблица 2. Основные блоки зафиксированных срезов [1–7]. Объем рабочего перечня для функциональных слоев — 173 862 единицы; он шире демографической картографической выборки по правилам включения.')
#3
page('Как изменялось население')
para(sections[3][1][0]);picture(F/'dynamics.png',maxh=16)
for p in sections[3][1][1:]:para(p)
caption('Рисунок 1. Относительное изменение населения в трех периодах [1]. Пропуски и начальные нули не превращены в проценты. Все точки одного размера; число точек не равно населению.')
#4
page('Как меняется сеть поселений')
picture(F/'structure.png',maxh=13.2);body(4)
table(['Сравнение 2002 и 2021','Число историй','Доля из 143 446'],[['Убыль','107 921','75,2%'],['Рост','26 437','18,4%'],['Без изменения','9 088','6,3%'],['Переход от >0 к 0 (часть убыли)','14 110','9,8% всех пар; входит в убыль']],[7,4,6.6])
caption('Рисунок 2 и таблица 3. Разные способы прочитать одну демографическую основу [1]. На карте численности нули показаны нижним цветом; на карте абсолютного изменения крайние значения насыщены за пределами ±10 тыс. человек. Округление долей дает 99,9%.')
#5-7
for n,key,cap in [(5,'education','Рисунок 3. Образовательные записи 2021, связанные с поселениями [2]. На панелях указано число поселений, а не число учреждений.'),(6,'medical','Рисунок 4. Медицинские записи 2021, связанные с поселениями [2]. ФАП — фельдшерско-акушерский пункт. НП — населенный пункт. Категории могут пересекаться.'),(7,'commerce','Рисунок 5. Торговля, финансы, почта и административные объекты 2021 [2]. Количество символов зависит от полноты источника.')]:
 page(sections[n][0]);picture(F/(key+'.png'),maxh=16);body(n);caption(cap)
#8 modern
page('Современная коммерция и доставка товаров')
picture(F/'modern_commerce.png',maxh=16)
para('Отдельный современный блок содержит Wildberries, каталог Places и ряды зарегистрированных организаций розницы и общепита. В снимке WB на 21 октября 2026 года опубликованы 91 043 записи раздела ru; 74 763 адресно сопоставлены с 16 172 поселениями. Это все типы записей доставки. Выбранные типы точек доставки WB (коды источника 8 и 34) представлены в 7 991 поселении; более строгая привязка до 5 км — в 7 973.')
para('Ozon показан только как лексический сигнал: 50 принятых записей названий в 31 поселении, а не проверенная сеть пунктов выдачи. Places сохраняется как самостоятельный каталог со своими правилами; регистрационная розница и общепит — записи организаций по кодам деятельности, а не физические магазины или рестораны.')
caption('Рисунок 6. Присутствие положительных записей в отдельных коммерческих источниках [3]. Точки — поселения; наборы имеют разные даты и не складываются в единый перечень уникальных предприятий. Современный слой не реконструирует сеть 2021 года.')
#9
page(sections[8][0]);picture(F/'municipal.png',maxh=15.7);body(8)
caption('Рисунок 7. Муниципальные контексты 2021 [4]. n — число муниципалитетов с числом и географической опорой для данного показателя. Цвета за границами шкал насыщены. Повторенные по деревням муниципальные значения предварительно сведены к одному муниципалитету.')
#10 nature + economy contextual map move to separate later
page(sections[9][0]);picture(F/'nature.png',maxh=12.2);body(9)
picture(F/'economy.png',maxh=5.1)
para('Экономический контекст 2021 доступен для 142 570 записей поселений; повторный подсчет регистрационных сведений совпал с сохраненными числами. Показатель не равен занятости или числу действующих предприятий. На карте показаны положительные значения с координатами.')
caption('Рисунок 8. Природный фон [4] и регистрационные экономические сведения [7]. Проекция новых карт — Альберса, параллели 50° и 70° с.ш. Основа Natural Earth; она служит для ориентации, а не доказательства границ обслуживания.')
#11 agriculture
page(sections[10][0]);picture(R/'maps/14_municipal2008_vs_regional2021_land.png',maxh=12.2);body(10)
picture(R/'maps/12_experimental_bonitet_intervals.png',maxh=5.8)
caption('Рисунок 9. Разные масштабы аграрных данных и пробный слой бонитета [5]. Две верхние карты не показывают временную динамику; нижняя карта не принята как точная национальная поверхность.')
#12 extra
page(sections[11][0]);picture(F/'extra_demography.png',maxh=11.8);body(11)
picture(R/'maps/06_ethnic_composition_context.png',maxh=6.3)
caption('Рисунок 10. Дополнительные источники демографии [1] и этнокультурный контекст [5]. n — число отображаемых рядов. Исторические серии наличного населения не смешиваются с постоянным населением; оценки после 2021 года округлены.')
#13 historical function context
page('Исторические и современные функции территорий')
picture(R/'maps/02_functional_evidence.png',maxh=10.1)
para(sections[12][1][1]);picture(R/'maps/09_current_monotowns_2026.png',maxh=7.5)
caption('Рисунок 11. Рекреационно-дачный, монопрофильный и добывающий контекст [5]. Слой OSM относится к датированному источнику; исторические и правовые перечни разделены. Текущий перечень моногородов — 2026 год, а не 2021.')
#14 model
page('От поселений к локальным системам')
for p in sections[13][1][:3]:para(p)
picture(F/'transport_extract.png',maxh=5)
picture(R/'maps/11_regional_sgrt_output_word_ci_2021_baseline.png',maxh=6.4)
para(sections[13][1][3].replace('карта ниже','карта выше'));caption('Рисунок 12. Фрагмент транспортного листа и примеры предшествующего расчета [5]. Дорожная сеть 19.10.2026; расчет WordCi. НП без нижнего назначения может быть связан с центром более высокого уровня. Новые циклы имеют отдельный статус.')
#15 progress references
page('Что сделано и что проверяем дальше')
para(sections[14][1][0]);para(sections[14][1][1])
table(['Результат','Что уже сохранено','Что еще нужно'],[
['База и карты','Переписные цепочки, координаты, происхождение значений; тематические слои','Первичная сверка 2010 и исторических границ'],
['Модель систем','Масса центров, дорожные времена, назначения, четыре уровня; сценарии','Повторный допуск сети 2021, исправление конфигураций и пересчет'],
['Содержательные типы','Проверяемые признаки девяти гипотез из Word','Работа и мощности учреждений, реальные связи и устойчивость границ']],[3.1,7.2,7.3])
para('Типы из исходного Word: урбанизирующиеся пригородные и аграрно-сервисные; стабильные аграрные и периферийные; рекреационно-дачные и этнокультурные; депрессивные аграрные, старопромышленные и ресурсно-вахтовые. Это девять исследовательских гипотез, а не готовое деление всей России.')
para('Дальше сравниваем динамику центра, сельского окружения и всей системы на согласованной сетке, проверяем чувствительность к источникам и параметрам, затем выделяем устойчивые содержательные типы. Уже создана основа для общероссийского анализа; окончательные границы и механизмы требуют независимого подтверждения.')
heading('Источники и версии')
base='https://github.com/ZakBu/russian-settlements-research/blob/'
refs=[('1','Широкая переписная база и методические ограничения','90fe15a2aebc7327b7d525d29c5d70057e320e2e/docs/CURRENT_REVIEW_RU_20261013.md'),('2','Образование медицина коммерция и правила привязки','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_completion_20261021/model_inputs_cycle05/relinked_v4/README.md'),('3','Современный коммерческий слой Wildberries и Places','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_replication_20261021/commerce/README.md'),('4','Муниципальные климатические и почвенные контексты','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_completion_20261021/context/README.md'),('5','Карты и ограничения предшествующей модели','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_completion_20261021/maps/README.md'),('6','Исходная концепция Word и статус циклов 05–06','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_completion_20261021/coordination/PROGRESS_cycle05_06.md'),('7','Повторный подсчет экономических записей','d86b2f8b0144537c1bb8348b7f6d320fe844e835/research/sgrt_word_completion_20261021/economy_source_recheck/README.md')]
for num,label,path in refs:
 p=para('['+num+'] ');p.paragraph_format.space_after=Pt(3)
 hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),doc.part.relate_to(base+path,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);hyperlink.append(r);p._p.append(hyperlink)
caption('Основа исследования — «Динамическая типология сельско-городских районов тяготения», текст исходного Word сохранен в reviews/original_word_source.txt исследовательского среза. OSM © OpenStreetMap contributors, ODbL; поставщик выгрузок Geofabrik. Первичные статистические источники — Росстат, региональные переписные публикации и Росреестр; климат и почвы — FAO GAEZ, WorldClim, GlobalAI и Национальный атлас почв. Данные имеют разные условия использования.')
# Update fields and language
settings=doc.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
for p in doc.paragraphs:
 for r in p.runs:
  lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');r._r.get_or_add_rPr().append(lang)
doc.core_properties.title='Трансформация территориальной структуры расселения России в 2002–2021 гг.';doc.core_properties.subject='Данные и ход исследования';doc.core_properties.author='';doc.core_properties.keywords='расселение, Россия, переписи, 2002, 2010, 2021'
path=OUT/'Трансформация_расселения_России_2002–2021.docx';doc.save(path);print(path)
