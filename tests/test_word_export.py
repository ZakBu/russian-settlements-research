"""Exercise real Pandoc DOCX output: semantics, citations, images and editable structure."""
import json
from pathlib import Path
import shutil
import struct
import tempfile
import unittest
from zipfile import ZipFile
import zlib

from docx import Document
from docx.shared import Mm, Pt
from lxml import etree

from scripts import export_word


def tiny_png(path):
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xffffffff)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">2I5B", 2, 2, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"\0\xff\xff\xff\xff\xff\xff" * 2)) + chunk(b"IEND", b""))


@unittest.skipUnless(shutil.which("pandoc"), "Pandoc is required for Word export")
class WordExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.picture = self.root / "map.png"
        tiny_png(self.picture)
        self.reference = self.root / "reference.docx"
        export_word.create_reference(self.reference)

    def publication(self):
        documents = []
        for index in (1, 2):
            documents.append({"id": f"chapter-{index}", "title": f"Глава {index}", "html": f'''
                <nav>Ненужное меню</nav>
                <h2 id="method">Метод и ограничения</h2>
                <p>Численность 142 765 человек <a href="#reference-1">[1]</a>.</p>
                <figure><img src="map.png" alt="Карта поселений"/><figcaption>
                <p>Рисунок {index} — Населённые пункты</p>
                <p>Источник: построено автором по данным переписей <a href="#reference-1">[1]</a>.</p>
                </figcaption></figure>
                <details><summary>Метод расчёта</summary><p>Скрытые в браузере подробности сохранены.</p></details>
                <p>Уточнение<sup><a href="#note" id="note-ref" data-footnote-ref>1</a></sup>.</p>
                <section class="footnotes"><ol><li id="note"><p>Пояснение к главе {index}. <a href="#note-ref" data-doc-footnote-backref>↩</a></p></li></ol></section>
                <table><tr><th>Год</th><th>Население</th></tr><tr><td>2021</td><td>142 765</td></tr></table>
                <p><a href="https://example.org/data-v1">Версия данных</a></p>'''} )
        return {"documents": documents, "bibliographyHtml": '<h2>Список использованных источников</h2><ol><li id="reference-1"><p>Переписи. URL: <a href="https://example.org/data-v1">https://example.org/data-v1</a>.</p></li></ol>'}

    def test_real_docx_preserves_figures_one_bibliography_links_tables_and_native_notes(self):
        output = self.root / "research.docx"
        export_word.export_job(self.publication(), output, self.root, self.reference)
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Глава 1", text)
        self.assertIn("Глава 2", text)
        self.assertIn("Рисунок 1 — Населённые пункты", text)
        self.assertIn("Рисунок 2 — Населённые пункты", text)
        self.assertEqual(text.count("Источник: построено автором по данным переписей [1]."), 2)
        self.assertEqual(text.count("Список использованных источников"), 1)
        self.assertEqual(text.count("Скрытые в браузере подробности сохранены."), 2)
        self.assertNotIn("Ненужное меню", text)
        self.assertEqual(len(document.tables), 2)
        self.assertEqual(len(document.inline_shapes), 2)
        with ZipFile(output) as archive:
            tree = etree.fromstring(archive.read("word/document.xml"))
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            self.assertTrue(tree.xpath('//w:hyperlink[@w:anchor="reference-1"]', namespaces=namespaces))
            self.assertTrue(tree.xpath('//w:bookmarkStart[@w:name="reference-1"]', namespaces=namespaces))
            self.assertTrue(tree.xpath('//w:instrText[contains(text(), "TOC")]', namespaces=namespaces))
            self.assertEqual(len(tree.xpath('//w:footnoteReference', namespaces=namespaces)), 2)
            notes = archive.read("word/footnotes.xml").decode()
            self.assertIn("Пояснение к главе 1", notes)
            self.assertIn("Пояснение к главе 2", notes)
            self.assertNotIn("↩", notes)
            self.assertTrue(tree.xpath('//w:sdtContent//w:hyperlink[starts-with(@w:anchor,"AtlasHeading")]', namespaces=namespaces))
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertIn('Target="https://example.org/data-v1"', relationships)
            self.assertIn('TargetMode="External"', relationships)
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210, delta=.03)
        self.assertAlmostEqual(document.sections[0].left_margin.mm, 30, delta=.03)
        self.assertAlmostEqual(document.sections[0].right_margin.mm, 15, delta=.03)
        self.assertEqual(document.styles["Normal"].font.name, "Times New Roman")
        self.assertEqual(document.styles["Normal"].font.size, Pt(14))
        self.assertEqual(str(document.styles["Heading 1"].font.color.rgb), "000000")
        heading_fonts = document.styles["Heading 1"].element.rPr.rFonts.attrib
        self.assertFalse(any("theme" in key.lower() for key in heading_fonts))
        image_paragraph = next(p for p in document.paragraphs if p._p.xpath(".//w:drawing"))
        self.assertEqual(image_paragraph.paragraph_format.first_line_indent, 0)
        captions = [p for p in document.paragraphs if p.text.startswith("Рисунок")]
        self.assertEqual([p.style.name for p in captions], ["Figure Caption", "Figure Caption"])
        self.assertEqual(document.styles["Figure Caption"].font.size, Pt(12))
        self.assertEqual(document.styles["Figure Caption"].paragraph_format.line_spacing, 1)
        self.assertEqual(document.paragraphs[-1].style.name, "Bibliography")

    def test_batch_cli_writes_downloadable_files_and_respects_reference_override(self):
        university = Document(self.reference)
        university.sections[0].left_margin = Mm(25)
        custom = self.root / "university.docx"
        university.save(custom)
        manifest = self.root / "jobs.json"
        manifest.write_text(json.dumps({"jobs": [dict(self.publication(), output=str(self.root / "first.docx")), dict(self.publication(), output=str(self.root / "second.docx"), toc=False)]}, ensure_ascii=False))
        export_word.main(["--input", str(manifest), "--reference-doc", str(custom)])
        self.assertAlmostEqual(Document(self.root / "first.docx").sections[0].left_margin.mm, 25, delta=.03)
        self.assertTrue((self.root / "second.docx").is_file())
        with ZipFile(self.root / "second.docx") as archive:
            self.assertNotIn(" TOC ", archive.read("word/document.xml").decode())

    def test_collection_requires_unique_chapters_and_preserves_order(self):
        job = self.publication()
        job["documents"].append(job["documents"][0])
        with self.assertRaisesRegex(ValueError, "unique"):
            export_word.export_job(job, self.root / "duplicate.docx", self.root, self.reference)

    def test_missing_or_remote_images_fail_before_writing_a_misleading_export(self):
        for source in ("missing.png", "https://example.org/current-map.png"):
            with self.subTest(source=source):
                with self.assertRaises(ValueError):
                    export_word.export_job({"documents": [{"id": "x", "title": "Карта", "html": f'<img src="{source}"/>'}]}, self.root / "missing.docx", self.root, self.reference)


if __name__ == "__main__":
    unittest.main()
