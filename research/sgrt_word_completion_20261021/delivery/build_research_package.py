#!/usr/bin/env python3
"""Render an inspectable scientific package from completed computations, never invented results."""
from pathlib import Path
from xml.sax.saxutils import escape
import argparse
import hashlib
import json
import re
import fitz
import pandas as pd
import numpy as np
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches,Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image

from report_enrichment import ensure_fresh,atlas_sources,enrich,supervisor_brief,configuration_for_model,typology_source

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'delivery'
def number(x):return f'{int(x):,}'.replace(',',' ')
def read_json(path):return json.loads(path.read_text())
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def readable_text(value):
    value=str(value).replace('Росстатallcategories','Росстат, хозяйства всех категорий').replace('угодьяРосреестра','угодья Росреестра').replace('national-контроль','общероссийская контрольная сумма')
    value=value.replace('RETURN_TO_REPAIR_LOOP','Требуется дальнейшая источниковая и методическая доработка; полное соответствие концепции пока не принято')
    value=re.sub(r'(?<=[А-Яа-яёЁ])(?=\d)|(?<=\d)(?=[А-Яа-яёЁ])',' ',value)
    value=re.sub(r'(?<=[А-Яа-яёЁ])(?=[A-Za-z])|(?<=[A-Za-z])(?=[А-Яа-яёЁ])',' ',value)
    value=value.replace('logMi','log Mi').replace('уровнеНП','уровне НП')
    value=re.sub(r'(?<=%)(?=[А-Яа-яёЁ])',' ',value)
    return value

def render(blocks,name):
    pdfmetrics.registerFont(TTFont('SGRT','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
    pdfmetrics.registerFont(TTFont('SGRTBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
    st=getSampleStyleSheet()
    for key,size,leading in [('TitleRU',19,25),('HeadRU',12,17),('BodyRU',9.4,13.8),('SmallRU',7.4,10.5)]:
        st.add(ParagraphStyle(name=key,fontName='SGRTBold' if key in ['TitleRU','HeadRU'] else 'SGRT',fontSize=size,leading=leading,spaceAfter=9,spaceBefore=7 if key=='HeadRU' else 0,keepWithNext=key in ['TitleRU','HeadRU']))
    story=[];doc=Document();doc.styles['Normal'].font.name='DejaVu Sans';doc.styles['Normal'].font.size=Pt(10)
    md=[]
    for kind,body in blocks:
        if kind in ['title','head','text','small']:
            body=readable_text(body)
            style={'title':'TitleRU','head':'HeadRU','text':'BodyRU','small':'SmallRU'}[kind]
            story.append(Paragraph(escape(body),st[style]))
            if kind=='title':doc.add_heading(body,0);md.append('# '+body)
            elif kind=='head':doc.add_heading(body,1);md.append('## '+body)
            else:doc.add_paragraph(body);md.append(body)
        elif kind=='table':
            header,rows=body
            header=[readable_text(v)for v in header];rows=[[readable_text(v)for v in row]for row in rows]
            grid=[[Paragraph(escape(str(v)),st['SmallRU']) for v in row] for row in [header]+rows]
            widths=[480/len(header)]*len(header)
            if header[:2]==['Ур.','Тип исходной концепции']:widths=[28,230,74,74,74]
            elif header[:2]==['Блок','Фактический массив']:widths=[100,125,255]
            elif header[:2]==['Центр и регион','Ур.']:widths=[145,28,28,62,62,155]
            tab=Table(grid,colWidths=widths,repeatRows=1)
            tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eff2')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#99aabb')),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story.extend([tab,Spacer(1,10)]);t=doc.add_table(rows=1,cols=len(header));t.style='Light Shading Accent 1'
            for cell,v in zip(t.rows[0].cells,header):cell.text=str(v)
            for row in rows:
                for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
            md.append('| '+' | '.join(map(str,header))+' |\n| '+' | '.join(['---']*len(header))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows))
        elif kind=='image':
            p,caption=body
            if not p.exists():raise FileNotFoundError(p)
            w,h=PILImage.open(p).size;width=480;height=width*h/w
            if height>390:height=390;width=height*w/h
            story.extend([Image(str(p),width=width,height=height),Paragraph(escape(caption),st['SmallRU'])]);doc.add_picture(str(p),width=Inches(6));doc.add_paragraph(caption)
            md.append(f'![{caption}](../{p.relative_to(ROOT)})')
    def footer(c,d):
        c.setFont('SGRT',7);c.drawString(54,26,'СГРТ России · расчёты и источниковая проверка · 21.10.2026');c.drawRightString(A4[0]-54,26,str(d.page))
    SimpleDocTemplate(str(OUT/f'{name}.pdf'),pagesize=A4,leftMargin=54,rightMargin=54,topMargin=45,bottomMargin=45,title='Динамическая типология СГРТ России').build(story,onFirstPage=footer,onLaterPages=footer)
    doc.save(OUT/f'{name}.docx');(OUT/f'{name}.md').write_text('\n\n'.join(md)+'\n')

def build(model):
    model=model.resolve()
    ensure_fresh(ROOT,model)
    receipt=read_json(model/'run_receipt.json')
    typepath=typology_source(ROOT,model)
    if not receipt['status'].startswith('COMPUTED'):raise ValueError('Full routed calculation required, not PREPARED output')
    inputs=[model/'run_receipt.json',model/'systems.parquet',model/'memberships.parquet',model/'word_type_hypotheses.parquet',model/'nodes.parquet',model/'hierarchy.parquet']
    inputs.extend(p for p in [ROOT/'analysis/run_receipt.json',ROOT/'analysis/model_metrics.csv',ROOT/'analysis/sample_gates.csv',ROOT/'analysis/analysis_inputs.json',ROOT/'maps/map_manifest.json',ROOT/'reviews/strict_gate_receipt.json',ROOT/'reviews/word_core_typology_operationalization.csv',ROOT/'typology_commuting_cases/source_manifest.json',ROOT/'typology_commuting_cases/explicit_mobile2015_named_shares.parquet',ROOT/'land2021_recovery/compatible_pair_receipt.json',ROOT/'soil_bonitet_recovery/recovery_receipt.json',Path(__file__),Path(__file__).with_name('report_enrichment.py'),model/'independent_medical_validation_receipt.json',model/'independent_medical_relation_audit.parquet'] if p.exists())
    narrative_sources=[ROOT/'demography2010_recovery/root_four_regions_admission_receipt.json',ROOT/'land2021_recovery/regional_sown_share_farmland2021.parquet',ROOT/'education/history_build_receipt.json',ROOT/'medical/validation_receipt.json',ROOT/'context/climate_settlement_features.parquet',ROOT/'context/soil_settlement_features.parquet',ROOT/'literature/industrial_case_evidence.parquet',ROOT/'literature/new_primary_bibliography.csv',ROOT.parent/'sgrt_word_replication_20261021/literature/bibliography.csv',ROOT/'validation'/model.name/'mass_block_ablation_summary.csv',ROOT/'validation'/model.name/'mass_block_ablation_receipt.json']
    inputs.extend(p for p in narrative_sources if p.exists())
    inputs.extend(p for p in (ROOT/'validation').glob('holdout2021_*')if p.suffix in ['.json','.csv','.parquet'])
    inputs.extend(sorted((ROOT/'model/scenario_comparison').glob('*.csv')))
    scenario_receipt=ROOT/'model/scenario_comparison/receipt.json'
    if scenario_receipt.exists():inputs.append(scenario_receipt)
    if typepath.parent!=model:inputs.extend(sorted(typepath.parent.glob('*.json'))+sorted(typepath.parent.glob('*.parquet')))
    frozen_before={str(p):sha(p) for p in inputs}
    configuration,configuration_path=configuration_for_model(model,receipt)
    inputs.append(configuration_path);frozen_before[str(configuration_path)]=sha(configuration_path)
    if configuration.get('rank_provisional_receipt'):
        rp=ROOT.parents[1]/configuration['rank_provisional_receipt'];inputs.append(rp);frozen_before[str(rp)]=sha(rp)
    if not (ROOT/'analysis/run_receipt.json').exists():raise FileNotFoundError('Actual regression/sample-gate receipt required')
    for stem in ['03_road_hierarchy_','04_nine_word_types_','11_regional_sgrt_']:
        if not (ROOT/f'maps/{stem}{model.name}.pdf').exists():raise FileNotFoundError('Actual model atlas/region figure required: '+stem+model.name)
    s=pd.read_parquet(model/'systems.parquet');m=pd.read_parquet(model/'memberships.parquet');types=pd.read_parquet(typepath)
    low=s[(s.level==1)&s.n_rural_external_settlements.gt(0)].copy();axis=pd.read_parquet(ROOT/'axis/full_settlement_axis.parquet')
    audit=read_json(ROOT/'reviews/strict_gate_receipt.json');road=read_json(ROOT/'transport/graph_receipt.json')
    blocks=[('title','Динамическая типология сельско-городских районов тяготения России'),('small','Вычислительная реализация концепции исходной рукописи. Материалы для научного обсуждения, 21 октября 2026 года.'),('head','1. Постановка задачи и статус результата'),('text','Единица исследования — центр и связанное с ним сельско-городское окружение, определяемое транспортной доступностью и собственной демографической, экономической и сервисной массой. Административная сетка используется для источникового контекста и независимой проверки, а не как готовая граница СГРТ.'),('text',f'Рассчитаны назначения, иерархия, раздельные сервисные сети и динамика центра, сельского окружения и всей системы. Исходная ось содержит {number(len(axis))} источниковых единиц, включая неполные временные цепочки; {number(axis.lat.notna().sum())} имеют принятые координаты. Модель содержит {number(len(low))} низовых систем с сельским окружением. Содержательные типы отделены от механических демографических сочетаний; недоказанные необходимые признаки остаются неизвестными.'),('text','Основная масса использует данные среза 2021 года; динамика 2010–2021 рассматривается в фиксированной аналитической сетке. Дороги характеризуют доступную картографическую сеть октября 2026 года и сценарные скорости. Такая конструкция допускает ретроспективное описание и анализ ассоциаций, но не является исторической реконструкцией транспортной сети 2010 года или причинным прогнозом.'),('small',f'Строгая независимая научная приёмка: {audit.get("decision",audit.get("state",audit.get("verdict",audit.get("status","см. reviews/strict_gate_receipt.json"))))}. Отсутствие допуска нельзя заменить успешным выполнением программных проверок.')]
    coverage=[['Переписная ось',number(len(axis)),'Пропуски и территориальные единицы сохранены отдельно'],['Образование2021','30 094 школьных объекта','Явный адрес и географический кандидат раздельны'],['Медицина2021','64 976 картографических объектов','Аптеки, стоматология и общая медицина разделены'],['Торговля/финансы/администрация2021','456 966 объектов','Количество записей не равно мощности'],['САТ >10°C FAO','154 968 НП','Суточный температурный показатель; норма1981–2010'],['Почвенные рейтинги FAO','152 337 НП','Агрономические классы; не российский бонитет'],['Посевы/сельхозугодья2008','68 560 контекстов НП','Ретроспективно гармонизированный МО-контекст; границы не доказаны']]
    protected_count=int(axis.population_value_quality_2010.eq('secondary_confidentiality_protected_value_exact_scope_unverified').sum())
    blocks.extend([('head','2. Информационная база и измеренные ограничения'),('table',(['Блок','Фактический массив','Смысл и ограничение'],coverage)),('text','Наблюдения услуг получены из полных российских срезов OSM2021 и2026 с сохранением исходных объектов, тегов, дат и правил связывания. Явно названное местоположение отделено от ближайшего пространственного кандидата. Ненаблюдаемость не кодируется как доказанное отсутствие услуги. Изменение картографических записей между датами не приравнивается к открытию или закрытию учреждений.'),('text',f'Отдельный существенный риск демографии: для {number(protected_count)} единиц численность2010 унаследована из вторичного источника с защитой конфиденциальности и неподтверждённой точностью охвата. Численно полная пара не равна паре независимо подтверждённых точных переписных значений. Качество источника сохраняется в оси и агрегатах; национальные демографические результаты зависят от этой исходной базы.'),('text','Независимая проверка по официальному перечню473 однозначно связанных НП с ФАП Удмуртии выявила лишь27 НП с любым координатным медицинским свидетельством в OSM. Поэтому модель использует измеренную наблюдаемость сети, а полная функциональная обеспеченность не установлена. Муниципальные естественное движение, миграция и ввод жилья сохраняют муниципальное зерно и не становятся собственными показателями каждой деревни. Рождения, смерти и прирост согласованы, численные конфликты и неоднозначные исходные коды исключены.'),('head','3. Масса, дорожная доступность и иерархия'),('text','Базовая положительная масса строится экспоненцированием среднего стандартизированных компонентных блоков. В неё входят собственная людность, наблюдаемые образование, здравоохранение, торговля и финансы, регистрационный экономический прокси, муниципальные жильё и демографическое движение. Веса не обучаются на сельском населении, которое эта же масса сама притянула. Тем самым исходная круговая калибровка устранена.'),('text','Модель выбора использует Gij=Mj·exp(−Tij/τ), Pij=Gij/(ΣkGik+Moutside). Время Tij рассчитывается по направленному дорожному графу; вне дорожной сети отдельно учитывается соединительный участок со сценарной скоростью15км/ч. Параметры τ=30минут, предел низового обслуживания180минут и порог выбора0,35 заданы как сценарий. Они не называются установленными законами поведения.'),('text','Ранги выделяются по распределению собственной функциональной силы, а не по одному порогу людности. Низовые системы объединяются целиком; родитель имеет более высокий ранг, состав не разрезается, суммы известных переписных наблюдений сохраняются. Изолированные, негеолоцированные и неопределённые единицы имеют отдельные статусы. Несколько сопоставимых центров обозначают модельный полицентрический потенциал; для подтверждения фактической полицентричности нужны независимые связи.')])
    levelrows=[]
    for level,g in s.groupby('level'):
        levelrows.append([str(level),number(len(g)),number(g.n_rural_external_settlements.gt(0).sum()),number(g.system_population_complete_2010.sum())])
    blocks.append(('table',(['Уровень','Системы','С сельским окружением','Полное население2010'],levelrows)))
    calpath=ROOT/'model/calibration_exploratory_2021/receipt.json'
    if calpath.exists():
        c=read_json(calpath)
        blocks.append(('text',f'Дополнительная нециркулярная административная калибровка выполнена на {c["training_units"]} полных единицах из {c["regions"]} регионов, независимо связанных с центром источниковыми административными отношениями. Ошибка межрегиональной проверки log-населения: {c["spatial_heldout_rmse"]:.3f}, базового среднего: {c["mean_baseline_rmse"]:.3f}. Это калибровка административного охвата, а не наблюдаемых поездок; текущие OSM-коды не доказывают историческую эквивалентность границ. Калиброванная масса автоматически основной не объявлена.'))
    blocks.extend([('head','4. Динамика центра, села и системы'),('text','Для каждого уровня сохранены известные суммы, число неизвестных значений, реальные нули, собственное население центра и население внешнего сельского окружения. Полным темпом района называется только изменение сопоставимой полной пары. Изменение состава территории и ошибка исторических связок остаются самостоятельным источником неопределённости.')])
    complete=low.system_population_complete_2010&low.system_population_complete_2021&low.rural_external_population_complete_2010&low.rural_external_population_complete_2021&low.center_population_2010.gt(0)&low.center_population_2021.ge(0)&low.rural_external_population_known_sum_2010.gt(0)
    d=low[complete].copy()
    if len(d):
        center=d.center_population_2021/d.center_population_2010-1;rural=d.rural_external_population_known_sum_2021/d.rural_external_population_known_sum_2010-1
        blocks.append(('text',f'На полной низовой подвыборке {number(len(d))} систем центр и сельское окружение одновременно сокращаются в {number(((center<0)&(rural<0)).sum())} случаях; центр растёт при сокращении окружения — в {number(((center>0)&(rural<0)).sum())}; центр сокращается при росте окружения — в {number(((center<0)&(rural>0)).sum())}. Это знаки совокупного изменения, без причинного объяснения миграцией или хозяйственной специализацией.'))
    blocks.append(('head','5. Девять содержательных типов исходной концепции'))
    rows=[]
    type_names=types.groupby('word_type').word_type_name.first().sort_index().to_dict()
    for typ,name in type_names.items():
        t=types[(types.level==1)&types.word_type.eq(typ)]
        rows.append([str(typ),name,number(t.rule_status.eq('satisfied_observed_indicators').sum()),number(t.operational_hypothesis.sum()),number(t.rule_status.eq('unknown_missing_required_indicators').sum())])
    blocks.append(('table',(['№','Тип','Все необходимые признаки','Операциональные гипотезы','Неизвестные признаки'],rows)))
    blocks.append(('text','Строгая матрица TRUE/FALSE/UNKNOWN проверяет необходимые признаки каждого типа. Отдельный слой операциональных типов-гипотез допускает содержательную классификацию по измеренным демографическим и функциональным сочетаниям, сохраняя недоказанные механизмы и возможные пересечения. Картографирование гипотез не превращает их в окончательно подтверждённые типы. Место проживания, национальные ответы и динамика населения не устанавливают индивидуальные причины поведения; моногород не тождествен промышленному спаду, добывающая площадка — переходу на вахту.'))
    validation=model/'independent_medical_validation_receipt.json'
    if validation.exists():
        v=read_json(validation);blocks.extend([('head','6. Независимая проверка модельных назначений'),('text',f'Официальные территории первичной медицинской помощи Ленинградской области2024 дают {v["unique_external_resolved_pairs"]} уникальных внешних отношений с явно связанным поставщиком. Точное совпадение низового назначения на всех таких парах: {100*v["exact_lowest_agreement_all_resolved_pairs"]:.1f}%; среди пар с поставщиком в наборе модельных кандидатов: {100*v["exact_lowest_agreement_eligible_pairs"]:.1f}%. Доля источниковых поставщиков, попавших в набор кандидатов, оценивается отдельно от точности назначения. Нормативное прикрепление не является наблюдаемой поездкой и не служит универсальной истиной для всех услуг. Этот корпус не используется для обучения и не назван слепой заранее зарегистрированной проверкой.')])
    blocks.extend([('head','7. Проверка гипотез и границы интерпретации'),('text','М1, М2 и М3 сравниваются на одной и той же полной подвыборке: сначала масса и доступность, затем сельский аграрный и природный контекст, затем конфигурация и пригородность. Стандартизация выполняется внутри обучающей части, регионы разделяют обучение и проверку. Неизвестная площадь сельхозугодий2021 не заменяется площадью2008; российский бонитет не заменяется рейтингом FAO. Отбор полной подвыборки и пространственная зависимость остатков проверяются отдельно.'),('text','Атлас содержит исходное покрытие, климат и почвы, вложенную иерархию, содержательные типы-гипотезы и фактические дорожные изохроны. Точки населённых пунктов и сеть не изображаются доказанными непрерывными территориальными полигонами. Для основного варианта и варианта с географическими кандидатами сохраняется отдельный состав, чтобы масштаб неопределённости связывания был видим.'),('head','8. Воспроизводимость и литературная опора'),('text','Все слои привязаны к устойчивым идентификаторам и снабжены датами, URL, SHA256, исходным зерном, единицами и правилами пропусков. Проверяющий агент независимо контролирует соответствие исходной рукописи; отказ по критерию возвращает расчёт на исправление. Программный цикл сохраняет команды, журналы и состояние приёмки. Прохождение технических проверок не объявляется научным завершением.'),('text','Методическая интерпретация опирается на выборочное первичное чтение Поляна, Трейвиша, Родомана, Покшишевского, Лаппо и Хорева, дополненное первичными текстами Баранского1946, Колосовского1947 и архивными выступлениями Маергойза в публикации2012. Перечень прочитанных страниц и ограничения чтения сохранены в литературной матрице; чужим авторам не приписываются наши сценарные пороги.'),('small','Проверяемые выходы: model/*/memberships.parquet, systems.parquet, hierarchy.parquet, function_specific_memberships.parquet, word_type_hypotheses.parquet, operational_type_candidates.parquet; analysis/; maps/; transport/; reviews/. Требования и научные ограничения читаются вместе с картами.')])
    metrics_path=ROOT/'analysis/model_metrics.csv'
    analysis_receipt=read_json(ROOT/'analysis/run_receipt.json')
    extra=[('head','7а. Фактически рассчитанное сравнение моделей')]
    if metrics_path.exists() and metrics_path.stat().st_size>2:
        metrics=pd.read_csv(metrics_path);primary=metrics[metrics.scenario.eq('primary') & metrics.outcome.eq('outcome_system_change_pct')]
        metric_rows=[[str(int(r.level)),r.model,number(r.n),str(int(r.regions)),f'{r.group_oof_rmse:.3f}',f'{r.group_oof_mae:.3f}'] for r in primary.itertuples()]
        extra.append(('table',(['Уровень','Модель','Общая N','Регионов','RMSE по новым регионам','MAE'],metric_rows)))
        extra.append(('text','Ошибки относятся к совокупному процентному изменению2010–2021, а не к ежегодным процентам. Общая выборка и совпадающие региональные разбиения сохранены для всех трёх спецификаций. Масса2021 включает население конечного года и потому механически связана с исходом; даже контроль исходного населения и отдельный сельский исход не устраняют эндогенность. Коэффициенты и интервалы имеют описательный смысл.'))
    else:
        extra.append(('text','Ни один уровень не прошёл заранее заданные требования общей полной выборки и числа регионов. Недостающие наблюдения не заменялись нулями; результат проверки состава выборки опубликован в analysis/sample_gates.csv. Отсутствующий результат не объявляется подтверждением гипотезы.'))
    gates=ROOT/'analysis/sample_gates.csv'
    if gates.exists():
        g=pd.read_csv(gates);pure=g[g.scenario.eq('pure_primary_population_source_sensitivity')]
        if len(pure):extra.append(('text','Проверка только на полностью первично подтверждённых переписных значениях: '+ '; '.join(f'уровень{int(r.level)}: N={int(r.common_sample)}, регионов={int(r.regions)}, {r.status}' for r in pure.itertuples())+'.'))
    location=next(i for i,(kind,body) in enumerate(blocks) if kind=='head' and body.startswith('8.'))
    blocks[location:location]=extra
    pars=configuration['model']
    position=next(i for i,(kind,body) in enumerate(blocks) if kind=='text' and body.startswith('Модель выбора'))
    blocks[position]=('text',f'Модель выбора использует Gij=Mj·exp(−Tij/τ), Pij=Gij/(ΣkGik+Moutside). Tij получено по направленному дорожному графу; соединительный внедорожный участок со скоростью15км/ч хранится отдельно. В основной конфигурации τ={pars["time_decay_minutes"]:g}минут, предел={pars["max_travel_minutes"]:g}минут, порог выбора={pars["minimum_choice_probability"]:g}, Moutside={pars["outside_mass"]:g}. При Moutside=0 вероятность условна на наблюдаемом наборе доступных поставщиков. Консервативный вариант с внешней альтернативой сохраняется отдельно. Параметры сценарные и не подбирались по точности на проверочных медицинских отношениях.')
    source_summary,type_summary,case_summary=enrich(ROOT,model,blocks,s,types,axis)
    render(blocks,'SGRT_RESEARCH_REPORT_20261021')
    render(supervisor_brief(ROOT,model,s,types,axis,source_summary,case_summary),'SGRT_SUPERVISOR_BRIEF_20261021')
    atlas=fitz.open()
    atlasfiles,atlas_skipped=atlas_sources(ROOT,model)
    for p in atlasfiles:
        with fitz.open(p) as src:atlas.insert_pdf(src)
    for p in sorted((ROOT/'transport/isochrones').glob('*.pdf')):
        with fitz.open(p) as src:atlas.insert_pdf(src)
    if len(atlas):atlas.save(OUT/'SGRT_ATLAS_20261021.pdf')
    assert frozen_before=={str(p):sha(p) for p in inputs},'Model changed during report rendering'
    checks=[]
    for p in OUT.glob('*.pdf'):
        with fitz.open(p) as pdf:checks.append({'file':p.name,'pages':len(pdf),'no_empty_pages':all(page.get_text().strip() or page.get_images() for page in pdf),'sha256':sha(p)})
    (OUT/'render_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
    (OUT/'package_inputs.json').write_text(json.dumps({'model_directory':str(model.relative_to(ROOT)),'configuration_path':str(configuration_path.relative_to(ROOT)),'model_inputs':frozen_before,'atlas_source_files':[str(p.relative_to(ROOT)) for p in atlasfiles],'scientific_acceptance_from':'reviews/strict_gate_receipt.json','stale_or_unselected_atlas_exclusions':atlas_skipped,'final_population_source_summary':source_summary},ensure_ascii=False,indent=2))
    print(json.dumps(checks,ensure_ascii=False,indent=2))
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--model',type=Path,required=True);a=p.parse_args();OUT.mkdir(exist_ok=True);build(a.model)
