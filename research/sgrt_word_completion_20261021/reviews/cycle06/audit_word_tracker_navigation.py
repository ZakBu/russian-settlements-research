from pathlib import Path
import json,hashlib,zipfile
from docx import Document
from lxml import etree
B=Path('/workspace/russian-settlements-research');O=B/'research/sgrt_word_completion_20261021/reviews/cycle06';P=O/'word_requirement_tracking';checks=[]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def ck(n,v):checks.append({'check':n,'pass':bool(v)})
x=json.loads((P/'requirement_tracking.json').read_text());ix=json.loads((P/'word_source_index.json').read_text());catalog=json.loads((P/'evidence_catalog.json').read_text())['evidence'];freeze=json.loads((P/'freeze_manifest.json').read_text());doc=Document(P/'original_user_concept.docx');ck('99_actual_DOCX_paragraphs',len(doc.paragraphs)==len(ix['paragraphs'])==99);ck('original_attachment_exact',sha(P/'original_user_concept.docx')==ix['original_attachment_sha256']==x['original_word_sha256'])
for actual,record in zip(doc.paragraphs,ix['paragraphs']):
 no=record['paragraph_1based'];ck(f'actual_word_paragraph_{no}',actual.text==record['visible_text']);math=actual._p.xpath('.//m:t');ck(f'actual_OMML_paragraph_{no}',[t.text or '' for t in math]==record['omml_math_texts'])
for fn,info in freeze['files'].items():ck('frozen_tracker_bytes_'+fn,sha(P/fn)==info['sha256'])
stale=[]
for eid,record in catalog.items():
 f=B/record['path'];same=f.exists() and sha(f)==record['sha256'];ck('evidence_pointer_exists_'+eid,f.exists())
 if not same:stale.append(dict(evidence_id=eid,path=record['path'],snapshot_sha256=record['sha256'],current_sha256=sha(f) if f.exists() else None))
ck('all32_unique_requirements',len(x['requirements'])==32 and {r['requirement_id'] for r in x['requirements']}=={f'W{i:02d}' for i in range(1,33)})
for r in x['requirements']:
 k=r['requirement_id'];ck(k+'_no_self_acceptance',r['current_requirement_acceptance']=='NOT_REASSESSED_NO_SELF_ACCEPTANCE');ck(k+'_evidence_ids_exist',set(r['evidence_ids'])<=set(catalog));ck(k+'_scope_and_next_work_explicit',bool(r['known_open_gap']) and bool(r['next_falsifiable_step']) and r['technical_receipt_check_counts_are_not_requirement_approvals'] is True)
 for q in r['word_source_quotes']:
  src=ix['paragraphs'][q['paragraph_1based']-1];ck(k+'_literal_quote_'+str(q['paragraph_1based']),q['literal_visible_text']==src['visible_text'] and q['omml_math_texts']==src['omml_math_texts'])
ck('no_new_acceptedcount_or_stopdecision',x['current_cycle06_accepted_requirements_count'] is None and x['not_a_new_stop_decision'] is True)
r=dict(status='ACCEPTED_HISTORICAL_NAVIGATION_SNAPSHOT_NOT_CURRENT_REQUIREMENT_ACCEPTANCE' if all(c['pass'] for c in checks) else 'REJECTED',all_checks_passed=all(c['pass'] for c in checks),n_checks=len(checks),n_failed=sum(not c['pass'] for c in checks),checks=checks,tracker_manifest_sha256=sha(P/'freeze_manifest.json'),snapshot_created_utc=x['created_utc'],current_changed_evidence_pointers=stale,scope='Original99DOCXparagraphs/OMML and32requirement attribution/scope/navigation independently checked. Evidence hashes describe their historical snapshot, not current files or new national output acceptance. No scientific requirements closed and no stop decision.',required_next_step='Versioned new evidence catalog for actual cycle06 baseline/configuration/typology/regression/review results; retain old snapshot immutable; classify operational completion separately from missing causal mechanisms.')
(O/'word_requirement_tracking_independent_review.json').write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n');print(r['status'],r['n_checks'],r['n_failed'],'changed_snapshot_pointers',len(stale));print([c for c in checks if not c['pass']])
