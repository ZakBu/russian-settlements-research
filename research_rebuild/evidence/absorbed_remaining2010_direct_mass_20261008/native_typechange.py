"""Separate genuine ownNP rural/PGT/rural history; never a direct absorption edge."""
import gzip,importlib.util,json,sys
from pathlib import Path
import pandas as pd
import openpyxl
import xlrd
ROOT=Path(__file__).resolve().parents[3];E=ROOT/'research_rebuild/evidence';OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'research_rebuild/mass_linkage'))
from working_state_20261007 import load
from current_chain_state_20261007 import sha,distance_km,normalize
spec=importlib.util.spec_from_file_location('typechange_finite',E/'working_full_chain_20261007/replay_additional_native_20261008.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
def main():
    source=OUT/'own_wikipedia_batch.json.gz';page=next(p for p in json.load(gzip.open(source,'rt'))['query']['pages'].values() if p['title']=='Донское (Светлогорский городской округ)')
    text=page['revisions'][0]['slots']['main']['*'];assert 'утратил статус посёлка городского типа' in text and 'в декабре 2018 года' in text
    assert 'Светлогорск' in text
    old='2002:023_38d853be37_02c_Kaliningrad.xls:Sheet1:13';ten='KAL2010:TOM1:S4:R1071';current='2021:data_allsettlements_anon_156_v20251217.parquet:parquet:26777'
    source02=Path('/workspace/settlements-raw/data/raw/2002/023_38d853be37_02c_Kaliningrad.xls');s02=xlrd.open_workbook(str(source02)).sheet_by_name('Sheet1')
    assert 'Светлогорска' in s02.cell_value(11,0) and 'Донское' in s02.cell_value(12,0) and float(s02.cell_value(12,1))==3170
    source10=Path('/workspace/settlements-work/sources/r2-missing/kaliningrad_tom1.xlsx');book=openpyxl.load_workbook(source10,read_only=True,data_only=True);sheet=book['4']
    values=list(next(sheet.iter_rows(min_row=1071,max_row=1071,values_only=True)));assert 'Донское' in str(values) and any(v==2924 for v in values)
    preceding=list(sheet.iter_rows(min_row=1040,max_row=1070,values_only=True))
    assert any('Донское' in str(row) or 'Светлогорск' in str(row) for row in preceding)
    state=load(51);before=helper.finite(state);beforemetrics=state.metrics();pins={str(p):sha(p) for p in state.inputs}
    for p in [source,source02,source10]:pins[str(p)]=sha(p)
    a,b,c=[state.by_id.loc[i] for i in [old,ten,current]]
    assert [a.population,b.population,c.population]==[3170,2924,2919]
    assert [a.settlement_type,b.settlement_type,c.settlement_type]==['посёлок','пгт','посёлок']
    assert all(r.region_norm=='калининградская' and normalize(r.settlement_name)=='донское' for r in [a,b,c])
    assert 'Светлогорск' in c.district_raw and str(c.oktmo)=='27734000106' and str(c.okato)=='27420000001'
    rivals=state.obs[state.obs.region_norm.eq('калининградская')&state.obs.settlement_name.map(normalize).eq('донское')].copy()
    assert len(rivals[rivals.census_year.eq(2021)])==2
    carrier=state.point_rows[current];co=page['coordinates'][0]
    assert distance_km((carrier['latitude'],carrier['longitude']),(co['lat'],co['lon']))<1
    roots={state.uf.find(i) for i in [old,ten,current]};years=set()
    for root in roots:assert not (state.years[root]&years);years|=state.years[root]
    protected=state.obs[['source_record_id','population','population_value_quality','settlement_name','settlement_type']].copy()
    edges=[];points=[]
    rule='Literal own Don NP with raw2002 Svetlogorsk source-parent and raw2010 own PGT row; exact current own Svetlogorsk county/code and independent own article coordinate; explicit2018 rural status conversion, all regional homonyms retained'
    for first,second in [(old,ten),(ten,current)]:
        if state.uf.find(first)!=state.uf.find(second):
            edges.append(dict(from_source_record_id=first,to_source_record_id=second,relation='same_place',decision_status='checked_rule_accepted',admission_rule=rule,boundary_comparability_asserted=False))
            state.union(first,second)
    for sid in [old,ten]:
        if sid not in state.point_rows:
            p=dict(carrier);p.update(target_source_record_id=sid,coordinate_source_record_id=current,coordinate_admission_status='reviewed_extension_rule_accepted',
                admission_allowed=True,admission_rule=rule,point_temporal_interpretation='Current own NP representative retrospectively reused by explicit source/county/status continuity; censusday measurement not asserted',
                source_sha256=state.by_id.loc[sid,'source_sha256'],source_locator=state.by_id.loc[sid,'source_locator'],boundary_comparability_asserted=False)
            points.append(p)
    pd.DataFrame(edges).to_csv(OUT/'accepted_separate_native_identity_edge_delta.csv',index=False)
    pd.DataFrame(points).to_csv(OUT/'accepted_separate_native_point_use_delta.csv',index=False)
    rivals.to_csv(OUT/'native_typechange_all_region_homonyms.csv',index=False)
    pd.DataFrame([dict(source_record_id=old,source_file=str(source02),source_sha256=sha(source02),raw_locator='Sheet1!row13;sourceparentrow12',raw_row_json=json.dumps(s02.row_values(12),ensure_ascii=False)),
        dict(source_record_id=ten,source_file=str(source10),source_sha256=sha(source10),raw_locator='sheet4!row1071',raw_row_json=json.dumps(values,ensure_ascii=False,default=str))]).to_csv(OUT/'native_typechange_actual_raw_source_checks.csv',index=False)
    state.add_deltas(point_paths=[OUT/'accepted_separate_native_point_use_delta.csv'])
    state.obs['root']=state.obs.source_record_id.map(state.uf.find);after=helper.finite(state);aftermetrics=state.metrics()
    assert protected.equals(state.obs[protected.columns])
    assert after['histories']-before['histories']==1
    result=dict(status='separate_actual51_native_Don_status_history_application_root_separate_stage_pending',baseline_stage=51,
        before_finite=before,after_finite=after,before=beforemetrics,after=aftermetrics,accepted_cases=1,accepted_edges=len(edges),accepted_points=len(points),
        finite_native_gain={'histories':1,'population_by_year':{'2002':3170,'2010':2924,'2021':2919}},
        direct_included_in_gain={'2002':0,'2010':0,'2021':0},all_regional_homonyms_retained=True,
        earlier_rural_to2010_PGT_status_date_unverified=True,explicit2018_PGT_to_rural_status_change=True,
        own_current_source_code_coordinate_binding=True,population_values_quality_unchanged=True,
        input_pins=pins,output_pins={p.name:sha(p) for p in OUT.glob('*native*csv')},code_sha256=sha(Path(__file__)))
    (OUT/'separate_native_application_receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result['finite_native_gain']))
if __name__=='__main__':main()
