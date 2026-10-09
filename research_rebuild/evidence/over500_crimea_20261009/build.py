from pathlib import Path
import json,hashlib,re
import pandas as pd
from openpyxl import load_workbook
ROOT=Path('/workspace/russian-settlements-research')
OUT=ROOT/'research_rebuild/evidence/over500_crimea_20261009'
BASE=ROOT/'research_rebuild/evidence/crimea_actual_available_census_2014_2021_20261008'
pins={}
def pin(p):
 p=Path(p); pins[str(p)]={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size};return p
r=pd.read_csv(pin('/dev/shm/over500-20261009/residual.csv'),dtype=str,keep_default_na=False).query("region_norm=='крым'").copy()
assert len(r)==483 and r.source_record_id.nunique()==483 and set(r.has_ownpoint)=={'True'} and set(r.census_year)=={'2021'}
a=pd.read_csv(pin(BASE/'actually_available_census_observations.csv.gz'),dtype=str,keep_default_na=False)
a=a[a.place_current_source_record_id.isin(r.source_record_id)].copy()
d=pd.read_csv(pin(BASE/'residual4_source_join/accepted_scoped_2014_observation_delta.csv.gz'),dtype=str,keep_default_na=False)
e=pd.read_csv(pin(BASE/'residual4_source_join/accepted_scoped_identity_delta.csv.gz'),dtype=str,keep_default_na=False)
for x in d.to_dict('records'):
 if x['current_2021_source_record_id'] not in set(r.source_record_id):continue
 assert len(e[(e.from_source_record_id==x['source_record_id'])&(e.to_source_record_id==x['current_2021_source_record_id'])&(e.decision_status=='reviewed_available_calendar_identity_accepted')])==1
 a=pd.concat([a,pd.DataFrame([dict(place_current_source_record_id=x['current_2021_source_record_id'],settlement_name=x['name_current_native'],grain='physical_settlement',Russian_2002_status='outside_scope',Russian_2010_status='outside_scope',strict_NP3_eligible='False',boundary_comparability='unknown_not_asserted',coordinate_basis='existing_accepted_current_ownpoint; retrospective_continuity_only',point_reference_target_id=x['source_record_id'],existing_identity_review_status='reviewed_available_calendar_identity_accepted',year='2014',source_record_id=x['source_record_id'],population=x['population'],population_raw=x['population_raw'],population_status=x['population_quality'],source_name_raw=x['source2014_caption_literal'],source_path=x['source_file'],source_sha256=x['source_sha256'],source_locator=x['source_locator'],native_code='',source_quality_original=x['population_quality'])])],ignore_index=True)
# Replay the exact archived literal worksheet once.
wpath=Path(a[a.year=='2014'].source_path.iloc[0]);pin(wpath)
assert pins[str(wpath)]['sha256']=='35e6acf1e5ecb66c23355591a0ddf630a70eefbedfca0ba470d6b0004baba06e'
wb=load_workbook(wpath,read_only=True,data_only=True)
cells={i:row for i,row in enumerate(wb['pub-01-03'].iter_rows(values_only=True),1)};wb.close()
for x in a[a.year=='2014'].to_dict('records'):
 row=int(re.search(r'row=(\d+)',x['source_locator']).group(1));c=cells[row]
 assert int(c[1])==int(x['population'])
 assert ' '.join(str(c[0]).split())==' '.join(x['source_name_raw'].split())
 assert int(c[2])+int(c[3])==int(c[1])
p2021=Path('/workspace/settlements-raw/data/interim/2021_tochno/data_allsettlements_anon_156_v20251217.parquet');pin(p2021)
assert pins[str(p2021)]['sha256']=='86c197cd522e0b63669e9c6e7f43fd3d82b3704c6a126c800a9968ecd16cae14'
n=pd.read_parquet(p2021,columns=['object_level','object_name','population','oktmo','mun_upper','mun_lower'])
for x in r.to_dict('records'):
 native=n.iloc[int(x['source_record_id'].rsplit(':',1)[1])-1]
 assert int(native.population)==int(float(x['population'])) and native.object_level=='Населенный пункт' and str(native.oktmo)==x['oktmo']
 if x['source_record_id'] not in set(a[a.year=='2021'].place_current_source_record_id):
  a=pd.concat([a,pd.DataFrame([dict(place_current_source_record_id=x['source_record_id'],settlement_name=x['settlement_name'],grain='physical_settlement',Russian_2002_status='outside_scope',Russian_2010_status='outside_scope',strict_NP3_eligible='False',boundary_comparability='unknown_not_asserted',coordinate_basis='existing_accepted_current_ownpoint',point_reference_target_id=x['source_record_id'],existing_identity_review_status='reviewed_available_calendar_identity_accepted',year='2021',source_record_id=x['source_record_id'],population=str(int(native.population)),population_raw=str(int(native.population)),population_status='actual_selected_native_2021',source_name_raw=native.object_name,source_path=str(p2021),source_sha256=pins[str(p2021)]['sha256'],source_locator='parquet:row_1based='+x['source_record_id'].rsplit(':',1)[1],native_code=x['oktmo'],source_quality_original=x['population_value_quality'])])],ignore_index=True)
assert len(a)==966 and not a.duplicated(['place_current_source_record_id','year']).any()
assert set(a.groupby('place_current_source_record_id').year.apply(lambda v:','.join(sorted(v))))=={'2014,2021'}
a.to_csv(OUT/'available_observations.csv.gz',index=False)
o=a[a.year=='2014'].set_index('place_current_source_record_id')
dis=[]
for x in r.to_dict('records'):
 v=o.loc[x['source_record_id']]
 dis.append(dict(source_record_id=x['source_record_id'],root=x['root'],settlement_name=x['settlement_name'],region_norm='крым',scope_disposition='outside_common_scope',Russian_2002_status='outside_common_scope',Russian_2010_status='outside_common_scope',available_years='2014,2021',available_year_route_status='accepted_existing_scoped_official_observations',existing_2014_route_id=v.source_record_id,route_decision_status=v.existing_identity_review_status,source2014_sha256=v.source_sha256,source2014_locator=v.source_locator,population_2014=v.population,population_2021=x['population'],ownpoint_status='existing_accepted_ownpoint',ownpoint_delta_required=False,strict_NP3_eligible=False,common_national_coverage_gain=0,boundary_comparability_asserted=False,historical_coordinate_measurement_asserted=False))
pd.DataFrame(dis).to_csv(OUT/'scope_dispositions.csv',index=False)
r.to_csv(OUT/'exact483_membership.csv.gz',index=False)
for p in [BASE/'verification_receipt.json',BASE/'application_receipt.json',BASE/'residual4_source_join/application_receipt.json',BASE/'residual4_source_join/verification_receipt.json']:pin(p)
(OUT/'source_pins.json').write_text(json.dumps(pins,ensure_ascii=False,indent=2)+'\n')
receipt=dict(targets=483,existing_ownpoints=483,new_point_uses=0,primary2014_rows_replayed=483,native2021_rows_replayed=483,available_calendar_numeric_pairs=483,base_existing_pairs=480,supplemental_previously_reviewed_pairs=3,scope='outside_common_scope',national2002_2010_coverage_gain=0,strictNP3_gain=0,unknown_values_allocated_or_zero_filled=False,populations_modified=False,graph_modified=False,source_note='2014 original Rosstat archived worksheet literal cells; 2021 native physical NP rows. Existing scoped decision accepted independently; supplemental packet retained exact typed alias/date limitations. No inference of Russian 2002/2010 observations.',outputs={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir() if p.is_file() and p.name!='receipt.json'})
(OUT/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
