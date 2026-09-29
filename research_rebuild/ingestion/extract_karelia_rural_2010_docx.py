"""Extract the official 2010 Karelia rural locality list from Rosstat DOCX.

Uses the DOCX table structure rather than the legacy XLS parser. The full
source row ledger preserves all rows in the locality-list tables, including
repeated headers and administrative aggregates that must not be added.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "research_rebuild/evidence/ingestion/sources/karelia_2010_rural_settlements.docx"
OUT = ROOT / "research_rebuild/evidence/ingestion"
LANDING_URL = "https://10.rosstat.gov.ru/folder/60059?print=1"
DOWNLOAD_URL = "https://10.rosstat.gov.ru/storage/mediabank/2_Сельские+населенные+пункты+РК.docx"
LOCALITY_PREFIX = re.compile(
    r"^(деревня|село|пос[её]лок|станция|разъезд|мест\.|хутор|погост|остров|"
    r"заимка|кишлак|аул|слобода|местечко)(?=\s|$)",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold().replace("ё", "е")
    value = re.sub(r"[^0-9a-zа-я]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


_SKIP_XML_TAGS = {qn("w:del"), qn("w:moveFrom"), qn("w:instrText"), qn("w:delText")}

def visible_xml_text(element) -> str:
    """Extract displayed Word text from XML, including smartTag/hyperlink runs.

    Deleted/moved-from content and field instructions are not displayed text.
    Tabs and line breaks are retained as separators before whitespace cleanup.
    """
    pieces: list[str] = []
    def visit(node):
        if node.tag in _SKIP_XML_TAGS:
            return
        if node.tag == qn("w:p") and pieces:
            pieces.append("\n")
        if node.tag == qn("w:t"):
            pieces.append(node.text or "")
            return
        if node.tag == qn("w:tab"):
            pieces.append("\t")
            return
        if node.tag in {qn("w:br"), qn("w:cr")} :
            pieces.append("\n")
            return
        for child in node:
            visit(child)
    visit(element)
    return re.sub(r"\s+", " ", "".join(pieces).replace("\xa0", " ")).strip()

def cell_text(cell) -> str:
    return visible_xml_text(cell._tc)

def document_cell_audit(document):
    """Compare visible XML extraction with python-docx across every document cell."""
    audit = []
    for ti, table in enumerate(document.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                xml_text = cell_text(cell)
                python_text = re.sub(r"\s+", " ", cell.text.replace("\xa0", " ")).strip()
                audit.append({
                    "source_locator": f"table[{ti}].row[{ri}].cell[{ci}]",
                    "table_index_zero_based": ti, "row_index_zero_based": ri,
                    "cell_index_zero_based": ci, "display_text_xml": xml_text,
                    "python_docx_text": python_text,
                    "text_differs_from_python_docx": xml_text != python_text,
                    "raw_cell_xml": cell._tc.xml,
                })
    return audit


def count_cell(raw: str) -> int:
    value = raw.strip().replace("\xa0", "").replace(" ", "")
    if value in {"-", "–", "—"}:
        return 0
    if value.isdigit():
        return int(value)
    raise ValueError(f"Expected integer or dash in a locality count cell: {raw!r}")


def xml_merge_facts(row) -> tuple[int, int, int]:
    grid_spans = vertical_merges = horizontal_merges = 0
    for tc in row._tr.tc_lst:
        tc_pr = tc.tcPr
        if tc_pr is None:
            continue
        if tc_pr.gridSpan is not None:
            grid_spans += 1
            horizontal_merges += 1
        if tc_pr.vMerge is not None:
            vertical_merges += 1
    return len(row._tr.tc_lst), horizontal_merges, vertical_merges


def extract(source_path: Path = SOURCE, output_dir: Path = OUT, source_label: str | None = None):
    source_path = source_path.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    if not source_path.exists():
        raise FileNotFoundError(f"Missing downloaded primary source: {source_path}")
    document = Document(source_path)
    digest = sha256(source_path)
    if source_label:
        # Prefer a caller-supplied portable locator in both repository and
        # staged clean-root runs so row values do not depend on checkout path.
        source_locator_path = source_label
    else:
        try:
            source_locator_path = source_path.relative_to(ROOT).as_posix()
        except ValueError:
            source_locator_path = source_path.as_posix()
    all_rows = []
    observations = []
    cell_audit = document_cell_audit(document)
    district_context = None
    municipality_context = None

    # In the contents, Table 1.8 starts on printed page 26; Table 1.9 starts
    # on page 46. Tables 22–41 are the 20 successive page-sized pieces of 1.8.
    for table_index in range(22, 42):
        table = document.tables[table_index]
        for row_index, row in enumerate(table.rows):
            cells = [cell_text(cell) for cell in row.cells]
            cells += [""] * max(0, len(table.columns) - len(cells))
            label = cells[0]
            label_norm = norm(label)
            if "муниципальный" in label_norm and "район" in label_norm:
                district_context = label
                municipality_context = None
            elif "городской округ" in label_norm:
                district_context = label
                municipality_context = None
            elif label_norm.endswith(" сельское поселение") or label_norm.endswith(" городское поселение"):
                municipality_context = label

            match = LOCALITY_PREFIX.match(label)
            raw_tc_count, hmerge_count, vmerge_count = xml_merge_facts(row)
            if match:
                raw_type = match.group(1)
                name = label[match.end():].strip()
                population, men, women = [count_cell(cells[i]) for i in (1, 2, 3)]
                if population != men + women:
                    raise AssertionError(
                        f"Sex counts do not sum to population at table {table_index}, row {row_index}"
                    )
                rec = {
                    "source_record_id": f"ROSSTAT2010KARELIA:RURAL:DOCX:t{table_index}:r{row_index}",
                    "census_year": 2010,
                    "population_scope": "permanent_residents_at_2010_census",
                    "region_raw": "Республика Карелия",
                    "region_norm": "карелия",
                    "settlement_name_raw": name,
                    "settlement_name_norm": norm(name),
                    "settlement_type_raw": raw_type,
                    "settlement_type_norm": norm(raw_type.replace(".", "")),
                    "district_context_raw": district_context,
                    "municipality_context_raw": municipality_context,
                    "population": population,
                    "men": men,
                    "women": women,
                    "population_raw_cell": cells[1],
                    "men_raw_cell": cells[2],
                    "women_raw_cell": cells[3],
                    "source_path": source_locator_path,
                    "source_sha256": digest,
                    "source_table_index_zero_based": table_index,
                    "source_table_row_index_zero_based": row_index,
                    "source_locator": f"table[{table_index}].row[{row_index}]",
                    "merged_cell_status": "present" if hmerge_count or vmerge_count else "none",
                    "source_cells_json": json.dumps(cells[:len(table.columns)], ensure_ascii=False),
                }
                observations.append(rec)
                disposition = "rural_locality_observation"
                reason = "official table 1.8 row with locality type and population/sex cells"
            elif not any(cells):
                disposition, reason = "empty_row", "no text in source table row"
                name = None
            elif row_index < 2 and any("Мужчины" in c or "женщины" in c for c in cells):
                disposition, reason = "repeated_table_header", "repeated header on a continued page of table 1.8"
                name = label or None
            else:
                disposition, reason = "administrative_or_control_row", "hierarchy or subtotal; never added to locality observations"
                name = label or None
            all_rows.append({
                "source_path": source_locator_path,
                "source_sha256": digest,
                "census_year": 2010,
                "source_table_index_zero_based": table_index,
                "source_table_row_index_zero_based": row_index,
                "source_locator": f"table[{table_index}].row[{row_index}]",
                "source_name_or_row_text": name,
                "row_disposition": disposition,
                "disposition_reason": reason,
                "district_context_raw": district_context,
                "municipality_context_raw": municipality_context,
                "raw_table_cell_count": raw_tc_count,
                "horizontal_merge_cell_count": hmerge_count,
                "vertical_merge_cell_count": vmerge_count,
                "source_cells_json": json.dumps(cells[:len(table.columns)], ensure_ascii=False),
                "source_row_xml": row._tr.xml,
            })

    # Official controls: Table 1.4 (index 4) gives 2010 totals and Table 1.6
    # (index 5) the district breakdown. Both are published official tables;
    # locality counts are checked against these controls but not force-fitted.
    t4 = document.tables[4]
    total_row = [cell_text(c) for c in t4.rows[3].cells]
    table4_control = {
        "source_locator": "table[4].row[3]",
        "settlement_count_2010": int(total_row[2]),
        "population_2010": int(total_row[4]),
    }
    t5 = document.tables[5]
    table5_controls = []
    for row_index, row in enumerate(t5.rows[4:], start=4):
        cells = [cell_text(c) for c in row.cells]
        if len(cells) < 5 or not cells[0]:
            continue
        label = cells[0]
        try:
            count_2010 = count_cell(cells[2])
            population_2010 = count_cell(cells[4])
        except ValueError:
            continue
        table5_controls.append({
            "district_or_city_okrug_raw": label,
            "settlement_count_2002": count_cell(cells[1]),
            "settlement_count_2010": count_2010,
            "population_2002": count_cell(cells[3]),
            "population_2010": population_2010,
            "source_locator": f"table[5].row[{row_index}]",
        })

    # Derive district totals exclusively from the 776 extracted locality rows.
    derived: dict[str, dict[str, int]] = {}
    for row in observations:
        district = row["district_context_raw"] or "UNRESOLVED_DISTRICT_CONTEXT"
        item = derived.setdefault(district, {"settlement_count": 0, "inhabited_settlement_count": 0,
                                            "population": 0, "men": 0, "women": 0})
        item["settlement_count"] += 1
        item["inhabited_settlement_count"] += int(row["population"] > 0)
        item["population"] += row["population"]
        item["men"] += row["men"]
        item["women"] += row["women"]

    controls = []
    official_keyed = {norm(x["district_or_city_okrug_raw"]): x for x in table5_controls}
    for district, values in sorted(derived.items()):
        district_key = norm(district)
        if district_key == norm("Городской округ г. Костомукша"):
            matches = [x for key, x in official_keyed.items() if key == norm("Костомукшский городской округ")]
        else:
            matches = [
                x for key, x in official_keyed.items()
                if key in district_key or district_key in key
            ]
        official = matches[0] if len(matches) == 1 else None
        controls.append({
            "district_context_raw": district,
            **values,
            "sexes_sum_to_population": values["men"] + values["women"] == values["population"],
            "table1_6_control_match_count": len(matches),
            "table1_6_settlement_count_2010": None if official is None else official["settlement_count_2010"],
            "table1_6_population_2010": None if official is None else official["population_2010"],
            "derived_count_matches_published_inhabited_count": (
                None if official is None else values["inhabited_settlement_count"] == official["settlement_count_2010"]
            ),
            "derived_population_matches_published_total": (
                None if official is None else values["population"] == official["population_2010"]
            ),
            "official_control_locator": None if official is None else official["source_locator"],
        })

    total_population = sum(row["population"] for row in observations)
    total_men = sum(row["men"] for row in observations)
    total_women = sum(row["women"] for row in observations)
    source_facts = {
        "source_title": "Сельские населенные пункты Республики Карелия. Итоги Всероссийской переписи населения 2010 года. Том 2",
        "issuing_body": "Территориальный орган Федеральной службы государственной статистики по Республике Карелия",
        "edition": "Официальное издание",
        "publication_place_year": "Петрозаводск, 2012",
        "official_site_card_publication_date": "2019-09-19",
        "observation_year": 2010,
        "census_reference_date": "2010-10-14",
        "boundary_reference": "Administrative-territorial boundaries as of 2010-10-14",
        "population_definition_excerpt": "All settlements not legally designated as cities or urban-type settlements are rural.",
        "population_definition_source_paragraph_index_zero_based": 517,
        "source_landing_url": LANDING_URL,
        "source_download_url": DOWNLOAD_URL,
        "source_path": source_locator_path,
        "source_sha256": digest,
        "source_bytes": source_path.stat().st_size,
        "publisher_listing_status": "Official Rosstat Karelia page lists the matching collection title, 3.63 MiB, and 2019-09-19 site publication date.",
        "download_transport_status": (
            "TLS certificate verification failed on initial retrieval; bytes and SHA-256 are preserved. "
            "The file is usable for provisional extraction and independent content/control checks; "
            "a reviewer should verify transport or an archived copy before final publication."
        ),
        "observation_date_is_not_site_publication_date": True,
        "locality_list_source_table_indices_zero_based": list(range(22, 42)),
        "source_table_title": "1.8 Численность населения сельских населенных пунктов",
        "extractor_version": "karelia-rural-docx-xml-visible-text-v2",
        "extractor_code_sha256": sha256(Path(__file__).resolve()),
        "extractor_config_sha256": hashlib.sha256(b"tables=22:42;visible_text=w:t,w:tab,w:br,w:cr;exclude=w:del,w:moveFrom,w:delText,w:instrText;locality-prefix-v1").hexdigest(),
        "source_table_title_contents_page": 26,
        "regional_summary_control": table4_control,
        "regional_population_derived_from_localities": total_population,
        "regional_men_derived_from_localities": total_men,
        "regional_women_derived_from_localities": total_women,
        "regional_sexes_sum_to_population": total_men + total_women == total_population,
        "locality_row_count": len(observations),
        "inhabited_locality_count": sum(row["population"] > 0 for row in observations),
        "zero_population_locality_count": sum(row["population"] == 0 for row in observations),
        "source_row_ledger_count": len(all_rows),
        "document_cell_audit_count": len(cell_audit),
        "document_cell_text_difference_count_vs_python_docx": sum(r["text_differs_from_python_docx"] for r in cell_audit),
        "document_cell_raw_xml_preserved": True,
        "visible_xml_policy": "include w:t through smartTag/hyperlink/runs; represent tabs/br/cr as separators; exclude w:del, w:moveFrom, w:delText, w:instrText",
        "district_controls": table5_controls,
        "district_locality_reconciliation": controls,
        "table_structure": {
            "table_count_in_document": len(document.tables),
            "document_paragraph_count": len(document.paragraphs),
            "locality_table_count": 20,
            "repeated_header_rows": 2,
            "merged_cell_rows": sum(bool(r["horizontal_merge_cell_count"] or r["vertical_merge_cell_count"]) for r in all_rows),
            "note": "Merged-cell structure and complete row XML are preserved; cell audit covers every cell and compares visible XML text to python-docx text."
        },
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    with (output_dir / "karelia_2010_rural_docx_source_rows.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(all_rows[0]))
        writer.writeheader()
        writer.writerows(all_rows)
    with (output_dir / "karelia_2010_rural_official_observations.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(observations[0]))
        writer.writeheader()
        writer.writerows(observations)
    with (output_dir / "karelia_2010_docx_all_cell_text_audit.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(cell_audit[0]))
        writer.writeheader()
        writer.writerows(cell_audit)
    with (output_dir / "karelia_2010_rural_district_controls.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(controls[0]))
        writer.writeheader()
        writer.writerows(controls)
    (output_dir / "karelia_2010_rural_docx_source_facts.json").write_text(
        json.dumps(source_facts, ensure_ascii=False, indent=2) + "\n"
    )
    print(json.dumps({
        "rural_locality_rows": len(observations),
        "population": total_population,
        "men": total_men,
        "women": total_women,
        "inhabited": source_facts["inhabited_locality_count"],
        "zero_population": source_facts["zero_population_locality_count"],
        "district_controls_all_match": all(
            row["derived_count_matches_published_inhabited_count"]
            and row["derived_population_matches_published_total"]
            for row in controls
        ),
    }, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-docx", type=Path, default=SOURCE,
                        help="Path to official rural-settlement DOCX source")
    parser.add_argument("--output-dir", "--output-root", dest="output_dir", type=Path, default=OUT,
                        help="Empty or writable directory for fresh extraction outputs")
    parser.add_argument("--source-label", default=None,
                        help="Portable source-relative locator label recorded in outputs")
    args = parser.parse_args()
    extract(args.source_docx, args.output_dir, args.source_label)

if __name__ == "__main__":
    main()
