#!/usr/bin/env python3
"""Apply the independently reviewed 2014 Crimea rows as a separate observation layer.

The 2014 rows are not added to the canonical Russian 2002/2010/2021 identity graph.
They are associated with reviewed current 2021 physical-place components, preserve
their own official archived source cells, and carry explicit scope/date limits.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from openpyxl import load_workbook

try:  # package import for tests; script import for direct CLI execution
    from .build_long_table import ACCEPTED_COORDINATE_STATUSES
except ImportError:  # pragma: no cover - exercised by direct script invocation
    from build_long_table import ACCEPTED_COORDINATE_STATUSES


EXPECTED_REVIEW_SHA256 = "356e93e3bed9ec86e5f421cd659987e6306315d4818fbf5ff0f7d3ad2723e951"
EXPECTED_WORKBOOK_SHA256 = "35e6acf1e5ecb66c23355591a0ddf630a70eefbedfca0ba470d6b0004baba06e"
EXPECTED_OBSERVATIONS_SHA256 = "b66bc22e0d174314a0ee9c7e7c76ba09239cb43d3127eaa5d96f7e196e68caab"
EXPECTED_EDGES_SHA256 = "268f7f9e7958a34e63cac1816bb53cfed965cf2eeda4c002db97a9060a38706e"
EXPECTED_RETRO_POINTS_SHA256 = "b1f8aeb1e13613659234862ef17c59fb0d943610050cb1339356e3997451fb0e"
EXPECTED_SOURCE_FETCH_RECEIPT_SHA256 = "d5e79b77b06147d936a8e34e8a57255fd39d736817ea90ae36b63c4344c6dbca"
SOURCE_ORIGINAL_URL = "http://www.gks.ru/free_doc/new_site/population/demo/perepis_krim/tab-krim/pub-01-03.xlsx"
SOURCE_ARCHIVE_URL = "https://web.archive.org/web/20150924122536id_/http://www.gks.ru/free_doc/new_site/population/demo/perepis_krim/tab-krim/pub-01-03.xlsx"
EXPECTED_REVIEW_STATUS = "independent_scoped_2014_crimea_current_place_review_complete_candidate_only"
EXPECTED_EDGE_STATUS = "independently_eligible_for_scoped_application_candidate_only"


def sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _need(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def _read_csv(path: str | Path) -> pd.DataFrame:
    # pandas' C parser has segfaulted on the wide Russian source packets in the
    # project environment; the slower Python parser is stable for these small files.
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", engine="python")


def _literal_int(value: Any) -> int | None:
    """Accept actual integer-like cells only; dashes/blanks remain unknown."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None
        return int(value)
    raw = str(value).strip()
    if not raw or raw in {"-", "—", "–", "−"}:
        return None
    try:
        number = Decimal(raw.replace(" ", "").replace("\u00a0", ""))
    except InvalidOperation:
        return None
    if number != number.to_integral_value():
        return None
    return int(number)


def verify_review_packet(manifest: dict[str, Any]) -> tuple[dict[str, Path], dict[str, Any]]:
    required = {
        "full_long", "selected_population", "identity_graph", "point_ledger",
        "review_receipt", "observations_csv", "identity_edges_csv",
        "retrospective_points_csv", "source_workbook", "source_fetch_receipt",
    }
    _need(required <= set(manifest.get("inputs", {})), f"manifest missing input pins: {sorted(required - set(manifest.get('inputs', {})))}")
    paths: dict[str, Path] = {}
    for name in sorted(required):
        entry = manifest["inputs"][name]
        path = Path(entry["path"])
        expected = str(entry["sha256"])
        _need(path.is_file(), f"missing input {name}: {path}")
        actual = sha256(path)
        _need(actual == expected, f"{name} hash mismatch: expected {expected}, got {actual}")
        paths[name] = path

    _need(sha256(paths["review_receipt"]) == EXPECTED_REVIEW_SHA256,
          "independent review receipt is not the approved frozen receipt")
    _need(sha256(paths["source_workbook"]) == EXPECTED_WORKBOOK_SHA256,
          "official archived source workbook hash differs from reviewed bytes")
    _need(sha256(paths["observations_csv"]) == EXPECTED_OBSERVATIONS_SHA256,
          "reviewed observation CSV differs from frozen input")
    _need(sha256(paths["identity_edges_csv"]) == EXPECTED_EDGES_SHA256,
          "reviewed identity-edge CSV differs from frozen input")
    _need(sha256(paths["retrospective_points_csv"]) == EXPECTED_RETRO_POINTS_SHA256,
          "reviewed retrospective-point CSV differs from frozen input")
    _need(sha256(paths["source_fetch_receipt"]) == EXPECTED_SOURCE_FETCH_RECEIPT_SHA256,
          "official source retrieval receipt differs from frozen input")
    fetch = json.loads(paths["source_fetch_receipt"].read_text(encoding="utf-8"))
    _need(fetch.get("url") == SOURCE_ARCHIVE_URL and fetch.get("final_url") == SOURCE_ARCHIVE_URL,
          "archival source URL differs from verified retrieval receipt")
    _need(fetch.get("response_sha256") == EXPECTED_WORKBOOK_SHA256 and
          fetch.get("response_bytes") == paths["source_workbook"].stat().st_size,
          "retrieval receipt does not bind the exact workbook bytes")
    receipt = json.loads(paths["review_receipt"].read_text(encoding="utf-8"))
    _need(receipt.get("status") == EXPECTED_REVIEW_STATUS, "unexpected review receipt status")
    _need(receipt.get("candidate_only") is True and receipt.get("not_an_admission_or_population_overlay_receipt") is True,
          "receipt must remain candidate-only; application must be explicitly scoped")
    _need(receipt.get("checks", {}).get("eligible_identity_observations") == 994,
          "receipt eligible observation count mismatch")
    _need(receipt.get("checks", {}).get("eligible_identity_edges") == 994,
          "receipt eligible identity edge count mismatch")
    _need(receipt.get("checks", {}).get("eligible_retrospective_point_uses") == 993,
          "receipt eligible point-use count mismatch")
    _need(receipt.get("checks", {}).get("strict_Russian_2002_2010_2021_chain_contribution") == 0,
          "2014 packet must not contribute to strict Russian three-census chain")
    return paths, receipt


def replay_workbook_rows(observations: pd.DataFrame, workbook_path: str | Path) -> dict[str, dict[str, Any]]:
    """Replay reviewed literal cells against the hash-pinned archived workbook."""
    required = {"historical_source_record_id", "observation_year_actual", "population",
                "source_workbook_sha256", "source_sheet", "source_row", "source_row_caption_literal",
                "source_row_type", "source_workbook_path", "raw_row_cells_json", "population_scope", "population_value_quality"}
    _need(required <= set(observations.columns), f"observations missing columns {sorted(required - set(observations.columns))}")
    _need(not observations.historical_source_record_id.duplicated().any(), "duplicate 2014 observation IDs")
    _need(set(observations.observation_year_actual) == {"2014"}, "unsupported observation year; only 2014 is authorized")
    _need(observations.historical_source_record_id.str.startswith("2014:").all(), "non-2014 source record ID in observation packet")
    _need(set(observations.review_status) == {EXPECTED_EDGE_STATUS}, "observation row is not in the independently reviewed eligible set")
    _need(observations.source_row_type.str.strip().ne("").all(), "physical observation is missing its literal source type")
    _need((observations.source_workbook_sha256 == EXPECTED_WORKBOOK_SHA256).all(), "observation workbook hash is not the reviewed official archive copy")
    _need((observations.population_scope == "single settlement row; current type and source row type match").all(),
          "non-physical or aggregate observation in reviewed population rows")

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        # `ReadOnlyWorksheet.cell(row=N)` rescans XML for every random row access.
        # Stream the one 1,378-row source sheet once and retain only reviewed rows.
        target_rows: dict[int, list[Any]] = {}
        for row in observations.itertuples(index=False):
            target_rows.setdefault(int(row.source_row), None)
        needed = set(target_rows)
        max_row = max(needed) if needed else 0
        ws = wb["pub-01-03"]
        for row_number, cells in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_number in needed:
                target_rows[row_number] = list(cells)
            if row_number >= max_row:
                break
        parsed: dict[str, dict[str, Any]] = {}
        for row in observations.itertuples(index=False):
            sheet_name = str(row.source_sheet)
            _need(sheet_name == "pub-01-03", f"unsupported source sheet: {sheet_name}")
            _need(sheet_name in wb.sheetnames, f"source sheet missing from workbook: {sheet_name}")
            sheet_row = int(row.source_row)
            _need(sheet_row > 0, f"invalid source row: {row.source_row}")
            raw_cells = json.loads(row.raw_row_cells_json)
            _need(isinstance(raw_cells, list) and len(raw_cells) >= 4, f"malformed raw cell vector at row {sheet_row}")
            streamed = target_rows.get(sheet_row)
            _need(streamed is not None, f"source row not found in workbook at {sheet_name}:{sheet_row}")
            actual_cells = list(streamed[:len(raw_cells)])
            _need(actual_cells == raw_cells,
                  f"official source row cells differ from reviewed CSV at {sheet_name}:{sheet_row}")
            normalized_caption = " ".join(str(actual_cells[0]).split())
            _need(normalized_caption == " ".join(str(row.source_row_caption_literal).split()),
                  f"source caption differs after whitespace-only normalization at {sheet_name}:{sheet_row}")
            total = _literal_int(actual_cells[1])
            _need(total is not None, f"non-numeric official total at {sheet_name}:{sheet_row}")
            _need(total == _literal_int(row.population), f"population mismatch against official workbook at {sheet_name}:{sheet_row}")
            male = _literal_int(actual_cells[2])
            female = _literal_int(actual_cells[3])
            if male is not None and female is not None:
                _need(male + female == total, f"male/female totals do not equal total at {sheet_name}:{sheet_row}")
            parsed[str(row.historical_source_record_id)] = {
                "total": total,
                "male": male,
                "female": female,
                "male_raw": actual_cells[2],
                "female_raw": actual_cells[3],
                "raw_cells": actual_cells,
            }
        return parsed
    finally:
        wb.close()


def validate_review_tables(
    observations: pd.DataFrame,
    edges: pd.DataFrame,
    points: pd.DataFrame,
    current_rows: pd.DataFrame,
    current_points: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Validate the exact reviewed 2014→2021 pair and coordinate-use packet."""
    _need(len(observations) == 994 and len(edges) == 994 and len(points) == 993,
          f"review packet row counts changed: observations={len(observations)}, edges={len(edges)}, points={len(points)}")
    for frame, key, label in ((edges, "from_source_record_id", "edge sources"),
                              (edges, "to_source_record_id", "edge targets"),
                              (points, "target_source_record_id", "point targets")):
        _need(not frame[key].duplicated().any(), f"duplicate {label}")
    _need(set(edges.review_status) == {EXPECTED_EDGE_STATUS}, "unexpected edge review status")
    _need(set(edges.relation) == {"same_place"}, "unexpected relation type")
    _need(set(edges.from_year) == {"2014"} and set(edges.to_year) == {"2021"}, "wrong endpoint years")
    _need((edges.boundary_comparability_asserted.str.lower() == "false").all(), "edge claims boundary comparability")
    _need((edges.strict_Russian_2002_2010_2021_chain_eligible.str.lower() == "false").all(),
          "2014 endpoint incorrectly marked strict-three-census eligible")

    obs_ids = set(observations.historical_source_record_id)
    _need(set(edges.from_source_record_id) == obs_ids, "identity edges do not cover exactly the official 2014 rows")
    _need((edges.source_population_row_id == edges.from_source_record_id).all(), "population row/source identity mismatch")
    obs_by_id = observations.set_index("historical_source_record_id")
    edge_by_from = edges.set_index("from_source_record_id")
    for sid, edge in edge_by_from.iterrows():
        obs = obs_by_id.loc[sid]
        _need(str(edge.source_population) == str(obs.population), f"edge population mismatch for {sid}")
        _need(edge.source_workbook_sha256 == EXPECTED_WORKBOOK_SHA256, f"edge source workbook mismatch for {sid}")

    _need(not current_rows.source_record_id.duplicated().any(), "current selected rows contain duplicate target IDs")
    current = current_rows.set_index("source_record_id", drop=False)
    targets = set(edges.to_source_record_id)
    _need(targets <= set(current.index), "identity edge targets a source row absent from the pinned 2021 core")
    for edge in edges.itertuples(index=False):
        row = current.loc[edge.to_source_record_id]
        _need(str(row.record_type) == "census" and int(row.observation_year) == 2021,
              f"edge target is not a current 2021 census row: {edge.to_source_record_id}")
        _need(str(row.entity_category) == "settlement", f"edge target is not a physical settlement: {edge.to_source_record_id}")
        _need(pd.notna(row.entity_id) and str(row.entity_id), f"edge target lacks core entity: {edge.to_source_record_id}")
        obs = obs_by_id.loc[edge.from_source_record_id]
        _need(str(obs.source_row_type) == str(row.settlement_type),
              f"reviewed source type disagrees with current physical type for {edge.from_source_record_id}")
        _need(str(edge.to_source_record_id) == str(obs.current_2021_source_record_id),
              f"wrong identity target for {edge.from_source_record_id}")

    _need(set(points.target_source_record_id) <= obs_ids, "point-use row has no reviewed 2014 observation")
    _need(len(set(points.target_source_record_id)) == len(points), "duplicate historical point-use target")
    point_ids = set(points.target_source_record_id)
    _need(len(obs_ids - point_ids) == 1, "expected exactly one identity observation with no accepted current-point carrier")
    current_points = current_points.copy()
    current_points["target_source_record_id"] = current_points.target_source_record_id.astype(str)
    _need(not current_points.target_source_record_id.duplicated().any(), "current point ledger has ambiguous 2021 target rows")
    # Retain the key column as well as the index: downstream record construction
    # needs a stable row key even when it receives this validated frame.
    cp = current_points.set_index("target_source_record_id", drop=False)
    edge_target = edges.set_index("from_source_record_id").to_source_record_id
    edge_qid = edges.set_index("from_source_record_id").qid
    for point in points.itertuples(index=False):
        _need(str(point.target_year) == "2014", f"unsupported point-use year for {point.target_source_record_id}")
        _need(point.coordinate_use_review_status == EXPECTED_EDGE_STATUS, "point use lacks independent review status")
        _need(point.coordinate_admission_status == "candidate_retrospective_current_point_use",
              "unexpected reviewed point-use admission marker")
        sid = str(point.target_source_record_id)
        current_id = edge_target.loc[sid]
        reviewed_qid = str(edge_qid.loc[sid])
        # The reviewer packet preserves the source field as either the QID or the
        # exact current source-row ID, depending on the raw P625 replay route. Both
        # are permitted; the explicit carrier must always be the reviewed target.
        _need(str(point.coordinate_source_record_id) in {reviewed_qid, current_id} and
              str(point.source_2021_carrier_target_source_record_id) == current_id,
              f"retrospective point is attached to wrong 2021 target for {sid}")
        _need(current_id in cp.index, f"no canonical 2021 point for reviewed point-use target {current_id}")
        carrier = cp.loc[current_id]
        _need(str(carrier.coordinate_admission_status) in ACCEPTED_COORDINATE_STATUSES,
              f"current target point is not canonically accepted: {current_id}")
        _need(math.isclose(float(point.latitude), float(carrier.latitude), rel_tol=0, abs_tol=1e-10) and
              math.isclose(float(point.longitude), float(carrier.longitude), rel_tol=0, abs_tol=1e-10),
              f"retrospective coordinate differs from current accepted point for {sid}")
        _need(str(carrier.coordinate_source).strip() and str(carrier.coordinate_provider).strip(),
              f"current accepted point lacks its actual source/provider for {current_id}")
        _need(str(point.coordinate_measurement_date_unknown).lower() == "true" and
              str(point.direct_historical_coordinate_measurement).lower() == "false" and
              str(point.boundary_comparability_asserted).lower() == "false" and
              str(point.population_scope_comparability_asserted).lower() == "false",
              f"retrospective point makes unsupported date/comparability claims for {sid}")
    return current, cp


def create_long_observations(
    observations: pd.DataFrame,
    edges: pd.DataFrame,
    points: pd.DataFrame,
    current_rows: pd.DataFrame,
    current_points: pd.DataFrame,
    replayed_cells: dict[str, dict[str, Any]],
    workbook_path: str | Path,
    workbook_sha: str,
) -> pd.DataFrame:
    current = current_rows.set_index("source_record_id", drop=False)
    point_carriers = current_points.set_index("target_source_record_id", drop=False)
    edge_by_from = edges.set_index("from_source_record_id")
    point_by_from = points.set_index("target_source_record_id")
    records = []
    for obs in observations.itertuples(index=False):
        sid = str(obs.historical_source_record_id)
        edge = edge_by_from.loc[sid]
        target_id = str(edge.to_source_record_id)
        target = current.loc[target_id]
        cells = replayed_cells[sid]
        point = point_by_from.loc[sid] if sid in point_by_from.index else None
        loc = {"sheet": obs.source_sheet, "row": int(obs.source_row), "caption": obs.source_row_caption_literal}
        record = {
            "observation_id": f"official2014:{sid}",
            "record_type": "scoped_official_observation",
            "entity_id": target.entity_id,
            "associated_census_entity_id": target.entity_id,
            "current_place_entity_id": target.entity_id,
            "association_status": "reviewed_scoped_same_place_to_2021",
            "spatial_identity_status": "reviewed_scoped_same_place_to_2021",
            "identity_quality": "reviewed_scoped_same_place_to_2021",
            "observation_year": 2014,
            "observation_time_precision": obs.observation_time_precision,
            "reference_date": None,
            "reference_date_basis": "source_year_only_no_exact_census_date_claimed",
            "actual_census_date_claimed": False,
            "source_record_id": sid,
            "source_publication_row_id": sid,
            "source_name_raw": obs.source_row_caption_literal,
            "settlement_name": target.settlement_name,
            "settlement_type": obs.source_row_type,
            "region_raw": target.region_raw,
            "source_region_raw": obs.source_section,
            "district_raw": None,
            "municipality_raw": " | ".join(json.loads(obs.admin_path_raw_json)),
            "population_value": cells["total"],
            "population_raw": str(cells["raw_cells"][1]),
            "population_male": cells["male"],
            "population_female": cells["female"],
            # These convenience raw-cell fields are text so mixed integer/dash
            # source cells serialize consistently. Exact original Excel scalar
            # types remain preserved in population_raw_cells_json.
            "population_male_raw": None if cells["male_raw"] is None else str(cells["male_raw"]),
            "population_female_raw": None if cells["female_raw"] is None else str(cells["female_raw"]),
            "population_raw_cells_json": json.dumps(cells["raw_cells"], ensure_ascii=False),
            "population_value_quality": "official_archived_copy_exact_np_row",
            "population_quality_limitation": "archived Rosstat source copy; not independently measured a second time",
            "population_scope": "one literal 2014 official settlement row; no 2021 boundary/scope equivalence asserted",
            "population_scope_source_literal": obs.population_scope,
            "population_scope_comparability_asserted": False,
            "population_comparability_to_2021": False,
            "administrative_boundary_comparability_to_2021": "unknown_not_asserted",
            "boundary_comparability_asserted": False,
            "source_path": str(workbook_path),
            "source_file": str(workbook_path),
            "source_workbook_path_reviewed": obs.source_workbook_path,
            "source_sha256": workbook_sha,
            "source_sheet": obs.source_sheet,
            "source_row": int(obs.source_row),
            "source_locator": json.dumps(loc, ensure_ascii=False, sort_keys=True),
            "source_row_caption_literal": obs.source_row_caption_literal,
            "source_row_caption_raw": cells["raw_cells"][0],
            "source_row_type_literal": obs.source_row_type,
            "source_admin_heading_raw": obs.admin_heading_raw,
            "source_admin_path_raw_json": obs.admin_path_raw_json,
            "source_section": obs.source_section,
            "source_item_qid": obs.source_item_qid,
            "source_item_title": obs.source_item_title,
            "source_workbook_retrieval_context": "official Rosstat workbook from TLS-validated Internet Archive capture; exact archived copy hash pinned",
            "source_publisher": "Federal State Statistics Service (Rosstat)",
            "source_original_url": SOURCE_ORIGINAL_URL,
            "source_archive_url": SOURCE_ARCHIVE_URL,
            "current_qid": obs.current_qid,
            "current_2021_source_record_id": target_id,
            "current_2021_source_native_id": target.source_native_id,
            "current_2021_population_context": target.population_value,
            "current_place_region_raw": target.region_raw,
            "native_2014_okato_raw": None,
            "native_2014_oktmo_raw": None,
            "native_2014_identifier_binding_asserted": False,
            "strict_Russian_2002_2010_2021_chain_eligible": False,
            "census_full_chain": False,
            "census_2002_status": "outside_russian_census_scope",
            "census_2010_status": "outside_russian_census_scope",
            "census_2021_status": "observed",
            "historical_identity_admitted": True,
            "historical_coordinate_asserted": False,
        }
        if point is None:
            record.update({"latitude": None, "longitude": None, "coordinate_quality": None,
                           "coordinate_admission_status": None, "coordinate_temporal_basis": None,
                           "coordinate_measurement_date_unknown": None, "coordinate_source": None,
                           "coordinate_source_record_id": None, "point_source_file": None,
                           "point_source_sha256": None, "point_source_locator": None})
        else:
            carrier = point_carriers.loc[target_id]
            record.update({"latitude": float(point.latitude), "longitude": float(point.longitude),
                           "coordinate_quality": "reviewed_modern_representative_point_retrospective",
                           "coordinate_admission_status": "reviewed_rule_accepted",
                           "coordinate_temporal_basis": "modern_rep_point_spatial_continuity_inference_dateunknown",
                           "coordinate_use_application_status": "accepted_scoped_retrospective_point_use",
                           "coordinate_measurement_date_unknown": True,
                           "coordinate_source": carrier.coordinate_source,
                           "coordinate_source_record_id": point.coordinate_source_record_id,
                           "coordinate_carrier_target_source_record_id": target_id,
                           "coordinate_provider": carrier.coordinate_provider, "coordinate_provider_id": carrier.get("coordinate_provider_id"),
                           "coordinate_admission_rule": "reviewed 2014 to 2021 same-physical-place relation; representative point reused retrospectively; no 2014 measurement or boundary claim",
                           "coordinate_provenance": carrier.get("coordinate_provenance"),
                           "coordinate_quality_source_ledger": carrier.get("coordinate_quality"),
                           "coordinate_provider_family": carrier.get("coordinate_provider_family"),
                           "point_source_file": carrier.get("source_file"),
                           "point_source_sha256": carrier.get("source_sha256"),
                           "point_source_locator": carrier.get("source_locator"),
                           "retrospective_point_reviewed_source": point.coordinate_source,
                           "retrospective_point_reviewed_source_record_id": point.coordinate_source_record_id,
                           "retrospective_point_reviewed_origin_file": point.point_origin_file,
                           "retrospective_point_reviewed_origin_sha256": point.point_origin_sha256,
                           "retrospective_point_reviewed_origin_locator": point.point_origin_locator})
        records.append(record)
    out = pd.DataFrame.from_records(records)
    _need(len(out) == 994, "long observation construction dropped or duplicated 2014 rows")
    _need(not out.observation_id.duplicated().any(), "duplicate 2014 observation IDs")
    return out


def append_observations_by_name(
    full_long_path: str | Path,
    observations: pd.DataFrame,
    output_path: str | Path,
) -> dict[str, Any]:
    """Write a new long snapshot without rebuilding or projecting its old rows."""
    output_path = Path(output_path)
    _need(not output_path.exists(), f"refusing to overwrite full long output: {output_path}")
    _need("observation_id" in observations and "record_type" in observations,
          "new observations require observation_id and record_type")
    _need(not observations.observation_id.duplicated().any(), "new long observations have duplicate IDs")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(":memory:")
    # DuckDB needs working memory while unioning the wide 864k-row canonical
    # table. Keep that bounded and deterministic, but leave sufficient headroom
    # for Parquet encoding in this environment.
    con.execute("SET memory_limit='1GB'")
    con.execute("SET threads=1")
    con.execute("SET preserve_insertion_order=false")
    input_path = str(full_long_path)
    before = con.execute("SELECT count(*) FROM read_parquet(?)", [input_path]).fetchone()[0]
    before_types = dict(con.execute("SELECT record_type, count(*) FROM read_parquet(?) GROUP BY 1", [input_path]).fetchall())
    new_ids = observations.observation_id.astype(str).tolist()
    collisions = con.execute(
        "SELECT count(*) FROM read_parquet(?) WHERE observation_id IN (SELECT unnest(?))",
        [input_path, new_ids],
    ).fetchone()[0]
    _need(collisions == 0, f"{collisions} new observation IDs already exist in the current long table")
    input_sql = "'" + input_path.replace("'", "''") + "'"
    con.execute(f"CREATE TEMP VIEW base_long AS SELECT * FROM read_parquet({input_sql})")
    con.register("new_observation_rows", observations)
    con.execute("COPY (SELECT * FROM base_long UNION ALL BY NAME SELECT * FROM new_observation_rows) TO ? (FORMAT PARQUET, COMPRESSION ZSTD)",
                [str(output_path)])
    try:
        after = con.execute("SELECT count(*) FROM read_parquet(?)", [str(output_path)]).fetchone()[0]
        after_types = dict(con.execute("SELECT record_type, count(*) FROM read_parquet(?) GROUP BY 1", [str(output_path)]).fetchall())
        _need(after == before + len(observations), f"long union row count mismatch: before={before}, added={len(observations)}, after={after}")
        expected_types = before_types.copy()
        for record_type, n in observations.record_type.value_counts(dropna=False).items():
            expected_types[record_type] = expected_types.get(record_type, 0) + int(n)
        _need(after_types == expected_types, "long union changed existing record-type counts")
        _verify_original_rows_unchanged(con, input_path, str(output_path))
    except Exception:
        con.close()
        output_path.unlink(missing_ok=True)
        raise
    con.close()
    return {"input_rows": before, "output_rows": after, "input_record_type_counts": before_types,
            "output_record_type_counts": after_types, "original_columns_null_safe_equal": True}


def _verify_original_rows_unchanged(con: duckdb.DuckDBPyConnection, input_path: str, output_path: str) -> None:
    """Check every original column/value survives the union, in bounded batches."""
    base_schema = con.execute("DESCRIBE SELECT * FROM read_parquet(?)", [input_path]).fetchall()
    output_schema = con.execute("DESCRIBE SELECT * FROM read_parquet(?)", [output_path]).fetchall()
    base_types = {str(row[0]): str(row[1]).upper() for row in base_schema}
    output_types = {str(row[0]): str(row[1]).upper() for row in output_schema}
    _need("observation_id" in base_types and "observation_id" in output_types,
          "both long tables must have observation_id")
    missing = sorted(set(base_types) - set(output_types))
    _need(not missing, f"long union dropped original columns: {missing}")
    type_mismatches = {c: (base_types[c], output_types[c]) for c in base_types if base_types[c] != output_types[c]}
    _need(not type_mismatches, f"long union changed original column types: {type_mismatches}")
    for label, path in (("input", input_path), ("output", output_path)):
        count, unique, nulls = con.execute(
            "SELECT count(*), count(DISTINCT observation_id), count(*) FILTER (WHERE observation_id IS NULL) FROM read_parquet(?)",
            [path],
        ).fetchone()
        _need(count == unique and nulls == 0, f"{label} long observation_id is not unique/non-null")
    joined = con.execute(
        "SELECT count(*) FROM read_parquet(?) b JOIN read_parquet(?) o USING(observation_id)",
        [input_path, output_path],
    ).fetchone()[0]
    base_count = con.execute("SELECT count(*) FROM read_parquet(?)", [input_path]).fetchone()[0]
    _need(joined == base_count, f"only {joined} of {base_count} original rows matched by observation_id")
    columns = [c for c in base_types if c != "observation_id"]
    for start in range(0, len(columns), 24):
        batch = columns[start:start + 24]
        predicates = " OR ".join(f"b.{_qident(c)} IS DISTINCT FROM o.{_qident(c)}" for c in batch)
        if not predicates:
            continue
        changed = con.execute(
            f"SELECT count(*) FROM read_parquet(?) b JOIN read_parquet(?) o USING(observation_id) WHERE {predicates}",
            [input_path, output_path],
        ).fetchone()[0]
        _need(changed == 0, f"union changed {changed} original row(s) in columns {batch}")


def _qident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def apply_manifest(manifest_path: str | Path) -> dict[str, Any]:
    manifest_path = Path(manifest_path)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    paths, review = verify_review_packet(manifest)
    outdir = Path(manifest["output_dir"])
    outdir.mkdir(parents=True, exist_ok=True)
    expected_outputs = ["accepted_scoped_observations_2014.parquet", "accepted_scoped_observations_2014.csv",
                        "accepted_scoped_identity_edges_2014_2021.csv", "accepted_scoped_retrospective_point_uses_2014.csv",
                        "application_receipt.json"]
    long_output = manifest.get("long_output")
    if long_output:
        long_output_path = Path(long_output)
        _need(not long_output_path.exists(), f"refusing to overwrite full long output: {long_output_path}")
        _need(long_output_path.resolve() not in { (outdir/name).resolve() for name in expected_outputs },
              "full long output path conflicts with another scoped application output")
        expected_outputs.append(Path(long_output).name)
    existing = [str(outdir / name) for name in expected_outputs if (outdir / name).exists()]
    _need(not existing, f"refusing to overwrite output(s): {existing}")

    observations = _read_csv(paths["observations_csv"])
    edges = _read_csv(paths["identity_edges_csv"])
    point_uses = _read_csv(paths["retrospective_points_csv"])
    _need(set(observations.source_workbook_path) == {str(paths["source_workbook"])},
          "observation source paths differ from pinned archived workbook path")
    cells = replay_workbook_rows(observations, paths["source_workbook"])

    # Pull only the reviewed target rows through DuckDB instead of materializing
    # the 864k-row long or 400k-row point ledger in pandas.
    target_ids = sorted(set(edges.to_source_record_id))
    longcols = ["record_type", "observation_year", "source_record_id", "entity_id", "entity_category",
                "settlement_name", "settlement_type", "region_raw", "population_value", "source_native_id"]
    con = duckdb.connect(":memory:")
    con.execute("SET memory_limit='384MB'")
    projected = ", ".join(f'"{name}"' for name in longcols)
    current_rows = con.execute(
        f"SELECT {projected} FROM read_parquet(?) WHERE record_type='census' AND source_record_id IN (SELECT unnest(?))",
        [str(paths["full_long"]), target_ids],
    ).fetchdf()
    _need(len(current_rows) == len(target_ids), "core long has missing or duplicate target observation rows")
    # The full selected 2021 population pins the same current target rows and is independently compared.
    selected_cols = ["census_year", "source_record_id", "settlement_name", "settlement_type", "region_raw",
                     "population", "source_native_id"]
    projected = ", ".join(f'"{name}"' for name in selected_cols)
    selected = con.execute(
        f"SELECT {projected} FROM read_parquet(?) WHERE census_year=2021 AND source_record_id IN (SELECT unnest(?))",
        [str(paths["selected_population"]), target_ids],
    ).fetchdf()
    _need(not selected.source_record_id.duplicated().any(), "selected 2021 target IDs are duplicated")
    selected_map = selected.set_index(selected.source_record_id.astype(str))
    for row in current_rows.itertuples(index=False):
        sid = str(row.source_record_id)
        _need(sid in selected_map.index, f"core target missing from selected population: {sid}")
        src = selected_map.loc[sid]
        for core_name, selected_name in (("settlement_name", "settlement_name"), ("settlement_type", "settlement_type"),
                                         ("region_raw", "region_raw"), ("population_value", "population"),
                                         ("source_native_id", "source_native_id")):
            a, b = getattr(row, core_name), src[selected_name]
            _need((pd.isna(a) and pd.isna(b)) or str(a) == str(b), f"7th long/current selected value mismatch {sid}:{core_name}")

    point_cols = ["target_source_record_id", "target_year", "latitude", "longitude", "coordinate_admission_status"]
    projected = ", ".join(f'"{name}"' for name in point_cols)
    point_cols = ["target_source_record_id", "target_year", "latitude", "longitude", "coordinate_admission_status",
                  "coordinate_source", "coordinate_provider", "coordinate_provider_id", "coordinate_quality",
                  "coordinate_provenance", "coordinate_provider_family", "source_file", "source_sha256", "source_locator"]
    projected = ", ".join(f'"{name}"' for name in point_cols)
    cp = con.execute(
        f"SELECT {projected} FROM read_parquet(?) WHERE regexp_replace(CAST(target_year AS VARCHAR), '\\.0$', '')='2021' AND target_source_record_id IN (SELECT unnest(?))",
        [str(paths["point_ledger"]), target_ids],
    ).fetchdf()
    con.close()
    current, cp = validate_review_tables(observations, edges, point_uses, current_rows, cp)
    long_obs = create_long_observations(observations, edges, point_uses, current_rows, cp, cells,
                                        paths["source_workbook"], EXPECTED_WORKBOOK_SHA256)

    con = duckdb.connect(":memory:")
    con.execute("SET memory_limit='384MB'")
    observation_ids = long_obs.observation_id.astype(str).tolist()
    collision_count = con.execute(
        "SELECT count(*) FROM read_parquet(?) WHERE observation_id IN (SELECT unnest(?))",
        [str(paths["full_long"]), observation_ids],
    ).fetchone()[0]
    con.close()
    _need(collision_count == 0, f"{collision_count} scoped 2014 observation IDs already exist in the 7th long table")

    # Enrich exact reviewed ledgers. The 2014 edges/points remain a separate layer;
    # the canonical 2002/2010/2021 graph and point ledger are read-only inputs.
    out_edges = edges.copy()
    out_edges["decision_status"] = "accepted_scoped_same_place_observation_only"
    out_edges["population_boundary_comparability_asserted"] = False
    out_edges["strict_Russian_2002_2010_2021_chain_eligible"] = False
    out_edges["canonical_three_census_graph_modified"] = False
    out_points = point_uses.copy()
    out_points["coordinate_admission_status"] = "reviewed_rule_accepted"
    out_points["point_use_application_status"] = "accepted_scoped_retrospective_point_use"
    out_points["coordinate_temporal_basis"] = "modern_rep_point_spatial_continuity_inference_dateunknown"
    out_points["provider_identifier_binding_claimed"] = False
    out_points["historical_measurement_claimed"] = False
    out_points["boundary_comparability_asserted"] = False
    out_points["population_scope_comparability_asserted"] = False
    out_points["current_seventh_point_ledger_sha256"] = manifest["inputs"]["point_ledger"]["sha256"]

    obs_parquet = outdir / "accepted_scoped_observations_2014.parquet"
    obs_csv = outdir / "accepted_scoped_observations_2014.csv"
    edge_csv = outdir / "accepted_scoped_identity_edges_2014_2021.csv"
    point_csv = outdir / "accepted_scoped_retrospective_point_uses_2014.csv"
    long_obs.to_parquet(obs_parquet, index=False)
    long_obs.to_csv(obs_csv, index=False, quoting=csv.QUOTE_MINIMAL)
    out_edges.to_csv(edge_csv, index=False, quoting=csv.QUOTE_MINIMAL)
    out_points.to_csv(point_csv, index=False, quoting=csv.QUOTE_MINIMAL)

    long_summary = None
    if long_output:
        long_output = Path(long_output)
        long_summary = append_observations_by_name(paths["full_long"], long_obs, long_output)

    outpaths = [obs_parquet, obs_csv, edge_csv, point_csv] + ([Path(long_output)] if long_output else [])
    output_pins = {p.name: {"path": str(p), "sha256": sha256(p), "bytes": p.stat().st_size} for p in outpaths}
    receipt = {
        "status": "applied_independently_reviewed_scoped_2014_observations_candidate_population_layer_only",
        "review_receipt_sha256": sha256(paths["review_receipt"]),
        "official_source_workbook_sha256": EXPECTED_WORKBOOK_SHA256,
        "rows": {"observations": len(long_obs), "identity_edges": len(out_edges), "retrospective_point_uses": len(out_points),
                 "observations_without_point": len(long_obs) - len(out_points), "population_sum": int(long_obs.population_value.sum()),
                 "male_female_numeric_rows_checked": int(sum(x["male"] is not None and x["female"] is not None for x in cells.values())),
                 "male_female_numeric_sum_matches": True},
        "semantics": {
            "same_physical_place_identity_to_reviewed_2021_target": True,
            "canonical_2002_2010_2021_graph_modified": False,
            "strict_Russian_2002_2010_2021_chain_contribution": 0,
            "exact_census_date_asserted": False,
            "2014_native_OKATO_or_OKTMO_binding_asserted": False,
            "2014_population_boundary_comparability_to_2021_asserted": False,
            "modern_point_is_2014_measurement": False,
            "population_value_source": "literal total cell from hash-pinned archived official Rosstat workbook row",
            "unresolved_gender_cells": "retained raw; never coerced from dash/blank to zero",
            "held_identity_rows_not_applied": 15,
            "held_missing_point_identity_row_preserved_without_coordinate": "2014:Q127387785:pub-01-03:733",
        },
        "long_append": long_summary,
        "inputs": {name: {"path": str(path), "sha256": manifest["inputs"][name]["sha256"]} for name, path in paths.items()},
        "outputs": output_pins,
    }
    receipt_path = outdir / "application_receipt.json"
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return receipt


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, help="JSON manifest with byte-pinned inputs, output_dir and optional long_output")
    args = parser.parse_args()
    receipt = apply_manifest(args.manifest)
    print(json.dumps({"status": receipt["status"], "rows": receipt["rows"], "long_append": receipt["long_append"],
                      "outputs": receipt["outputs"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
