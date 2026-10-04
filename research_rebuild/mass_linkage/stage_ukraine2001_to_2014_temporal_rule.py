#!/usr/bin/env python3
"""Build an independently reviewable 2001→2014→2021 scoped continuity packet.

No old identity is admitted. The packet tests source names, types, hierarchy,
whole-table uniqueness, current native settlement rows, existing 2014→2021
associations, accepted contemporary point context, and event holds. It writes
candidate edges and retrospective point-use candidates for independent review.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
import xlrd
from openpyxl import load_workbook

REVIEW = Path("/workspace/settlements-work/continuation_20261004/independent_review/ukraine2001_27_primary_review")
STAGED_2001 = Path("/workspace/settlements-work/continuation_20261004/independent_review/ukraine2001_crimea_primary_staged_v2")
OUT_DEFAULT = Path("/workspace/settlements-work/continuation_20261004/independent_review/ukraine2001_to_2014_temporal_rule_candidate")
T1_XLS = Path("/workspace/settlements-work/continuation_20261004/federal_and_history/crimea_2001_source_claim_review_20261004/official_tls_source/extracted/5.xls")
PREFACE = Path("/workspace/settlements-work/continuation_20261004/federal_and_history/ukraine_2001_official_archive_inspection/doc_text/Передмова 1.txt")
XLS14 = Path("/workspace/settlements-work/continuation_20261004/federal_and_history/crimea_2001_source_claim_review_20261004/official_tls_source/gks_2014_perepis_krim_pub-01-03_wayback_20150924.xlsx")
BASE_LONG = Path("/workspace/settlements-work/continuation_20261004/root/scoped2014_applied/seventh_long_with_scoped2014.parquet")
SELECTED = Path("/workspace/settlements-delivery/continuation-consolidated-20261003/selected_observations.parquet")
CURRENT_RAW = Path("/workspace/settlements-raw/data/interim/2021_tochno/data_allsettlements_anon_156_v20251217.parquet")
POINTS = Path("/workspace/settlements-work/continuation_20261004/accepted_mass_seventh_reviewed/accepted_point_uses.parquet")
EDGES14 = Path("/workspace/settlements-work/continuation_20261004/independent_review/crimea2014_scoped/eligible_2014_to_2021_identity_edges.csv")
POINTS14 = Path("/workspace/settlements-work/continuation_20261004/independent_review/crimea2014_scoped/eligible_2014_retrospective_point_uses.csv")

PINS = {
    "ukraine2001_review_receipt": (REVIEW / "review_receipt.json", "7c32557cb5af6b2eeaaae4b00589a61d4035c392d85ae465b3c8ee95ae23f806"),
    "ukraine2001_reviewed_source_rows": (REVIEW / "eligible_2001_official_observations.csv", "4eb8f0ac200f63d2916b269e023f592e08b61f020042581787a2a42dbd8526e8"),
    "ukraine2001_subject_context": (REVIEW / "candidate_2001_observation_subject_contexts.csv", "3123b74ab742afc7f01919927959be396b924096d4e9a545859bb31ee46c220b"),
    "ukraine2001_point_context": (REVIEW / "candidate_2001_current_point_context_uses.csv", "8940826082945832c05a61d6378b53335c23f28ff6ae81376a1f2758955a5560"),
    "ukraine2001_staged_receipt": (STAGED_2001 / "application_receipt.json", "ee9f5b985cec3b6c1219c395c87d980600a5169e08055b8c13412a9b709668ef"),
    "ukraine2001_staged_observations": (STAGED_2001 / "official_2001_source_observations.parquet", "ac8de2ad79657b9ae7a1f89f7985a85f476b2001e16aeb06872a94ebabf2e8f7"),
    "source_2001_workbook": (T1_XLS, "5088ed43450aff371eb8b7382222d3c08ebc591e676368f5e00a76ffa4a56b33"),
    "source_2001_preface": (PREFACE, "6f2a5380e62b44898741de218f49f7017b91cff96145f6ea24e35bd8b56ab3f1"),
    "source_2014_workbook": (XLS14, "35e6acf1e5ecb66c23355591a0ddf630a70eefbedfca0ba470d6b0004baba06e"),
    "selected_2021": (SELECTED, "4ff918ae07715e98a37aa5dc77546f3d7b7ac9c241c7c01a041c8f72a6f8c657"),
    "current_2021_raw": (CURRENT_RAW, "86c197cd522e0b63669e9c6e7f43fd3d82b3704c6a126c800a9968ecd16cae14"),
    "base_long": (BASE_LONG, "be04c25f2cdcc1614e2c17865ca8a96593e4308f36ad344281062c37c7bd6360"),
    "accepted_current_points": (POINTS, "2215702e772417bb512da8c0a6c0d31ad9783c6a5345d4c0bbfe18403e11ad22"),
    "independent_2014_to_2021_edges": (EDGES14, "268f7f9e7958a34e63cac1816bb53cfed965cf2eeda4c002db97a9060a38706e"),
    "independent_2014_point_candidates": (POINTS14, "b1f8aeb1e13613659234862ef17c59fb0d943610050cb1339356e3997451fb0e"),
}

# These two official Ukrainian/Russian spellings are explicit source-language
# aliases, not fuzzy matches. Other names use only deterministic orthographic
# character normalization and punctuation/spacing normalization.
EXPLICIT_NAME_ALIASES = {
    "старийкрим": "старыйкрым",
    "курпати": "курпаты",
}

RULE_TEXT = (
    "Candidate 2001→2014 physical-place continuity requires an individual official 2001 city/urban-type row; exact official 2014 city/PGT row; same observed settlement class; deterministic Ukrainian/Russian spelling match (including only the two explicit aliases listed in this packet); unique name-within-observed-type competitors in the whole 2001 ARK urban-locality table, whole 2014 pub-01-03 locality table, and current 2021 selected Crimea settlement frame; actual source hierarchy retained on both rows; current native 2021 NP row with accepted point; existing scoped 2014→2021 same-place association; no unresolved move/successor/merge event. 2001 native code is not required. The current representative point is optional retrospective physical-place context only; it is not a 2001 coordinate measurement or native-code binding. The relation does not assert boundary or population comparability, exact legal-status dates, or Russian 2002/2010 census coverage."
)


def sha256(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def need(ok: bool, message: str) -> None:
    if not ok:
        raise ValueError(message)


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", engine="python")


def normalize_name(name: str) -> tuple[str, str]:
    """Return explicit canonical key plus the normalization class used."""
    raw = unicodedata.normalize("NFKC", str(name)).strip().casefold()
    raw = re.sub(r"^(?:м\.|смт|г\.|пгт)\s*", "", raw, flags=re.I)
    key = re.sub(r"[^\wа-яёієїґ]+", "", raw, flags=re.I)
    explicit = EXPLICIT_NAME_ALIASES.get(key)
    if explicit:
        return explicit, "explicit_Ukrainian_Russian_place_name_alias"
    translit = (raw.replace("і", "и").replace("ї", "и").replace("є", "е").replace("ґ", "г")
                .replace("’", "").replace("'", "").replace("ь", ""))
    key = re.sub(r"[^\wа-яё]+", "", translit, flags=re.I)
    rule = "Ukrainian_Russian_orthography" if translit != raw else "case_space_punctuation_only"
    return key, rule


def normalize_source_caption(value: str) -> str:
    """Normalize only layout whitespace in a copied spreadsheet caption."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def type_key_2001(raw_type: str) -> str:
    return "city" if str(raw_type) == "city" else "urban_type_settlement" if str(raw_type) == "urban_type_settlement" else "unknown"


def type_key_2014(raw: str) -> str:
    value = str(raw).strip().casefold()
    if value in {"город", "городской", "м."}:
        return "city"
    if value in {"пгт", "поселок городского типа", "посёлок городского типа"}:
        return "urban_type_settlement"
    return "unknown"


def type_key_2021(raw: str) -> str:
    value = str(raw).strip().casefold()
    if value in {"город", "город", "м."}:
        return "city"
    if value in {"пгт", "поселок городского типа", "посёлок городского типа"}:
        return "urban_type_settlement"
    return "unknown"


def parse_2001_locality(label: str) -> tuple[str, str] | None:
    m = re.match(r"^\s*(м\.|смт)\s+(.+?)\s*$", str(label), flags=re.I)
    if not m:
        return None
    return ("city" if m.group(1).casefold() == "м." else "urban_type_settlement", m.group(2).strip())


def parse_2014_locality(label: str) -> tuple[str, str] | None:
    text = str(label).strip()
    pgt = re.search(r"(?i)(?:^|\s)пгт\s+([^()\n]+)", text)
    if pgt:
        return "urban_type_settlement", pgt.group(1).strip().rstrip(" -–")
    city = re.search(r"(?i)г\.\s*([^()\n]+)", text)
    if city:
        return "city", city.group(1).strip().rstrip(" -–")
    return None


def evaluate_rule(predicates: dict[str, bool]) -> tuple[str, list[str]]:
    hard = [k for k, v in predicates.items() if not v]
    if hard:
        return "held_not_rule_admissible", hard
    return "eligible_for_fixed_independent_review", []


def distance_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 6371.0088
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2-lat1), math.radians(lon2-lon1)
    a = math.sin(dp/2)**2 + math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 2*r*math.asin(math.sqrt(a))


def verify_inputs() -> dict[str, str]:
    actuals = {}
    for name, (path, expected) in PINS.items():
        need(path.is_file(), f"missing pinned input {name}: {path}")
        actual = sha256(path)
        need(actual == expected, f"hash mismatch for {name}: {actual}")
        actuals[name] = actual
    r = json.loads(PINS["ukraine2001_review_receipt"][0].read_text(encoding="utf-8"))
    need(r.get("independently_verified_source_rows") == 27 and r.get("candidate_population_total") == 845627 and
         int(r.get("identity_admissions_2001_to_2014", 0)) == 0 and int(r.get("identity_admissions_2001_to_2021", 0)) == 0,
         "independent 2001 source review receipt changed")
    staged = json.loads(PINS["ukraine2001_staged_receipt"][0].read_text(encoding="utf-8"))
    need(staged.get("official_2001_source_observations") == 27 and staged.get("2001_identity_edges_created") == 0,
         "root-approved 2001 source assertion layer changed")
    return actuals


def read_2001_universe() -> list[dict[str, Any]]:
    wb = xlrd.open_workbook(str(T1_XLS), on_demand=True)
    try:
        sh = wb.sheet_by_name("АРК")
        out = []
        for i in range(sh.nrows):
            label = sh.cell_value(i, 0)
            parsed = parse_2001_locality(str(label))
            if parsed:
                typ, name = parsed
                # Table 5 columns 1 and 2 are language counts; column 3 is total present population.
                out.append({"row": i+1, "type": typ, "name": name, "key": normalize_name(name)[0], "label": str(label), "population": sh.cell_value(i, 3)})
        return out
    finally:
        wb.release_resources()


def read_2014_universe() -> tuple[list[dict[str, Any]], dict[int, tuple[str, Any]]]:
    wb = load_workbook(XLS14, read_only=True, data_only=True)
    try:
        need("pub-01-03" in wb.sheetnames, "2014 Crimea source worksheet missing")
        sh = wb["pub-01-03"]
        out = []
        rows: dict[int, tuple[str, Any]] = {}
        for i, row in enumerate(sh.iter_rows(min_col=1, max_col=2, values_only=True), 1):
            label, pop = row
            if isinstance(label, str):
                rows[i] = (label, pop)
                parsed = parse_2014_locality(label)
                if parsed:
                    typ, name = parsed
                    out.append({"row": i, "type": typ, "name": name, "key": normalize_name(name)[0], "label": label, "population": pop})
        return out, rows
    finally:
        wb.close()


def fixed_risk_sample(candidates: pd.DataFrame) -> pd.DataFrame:
    selected: dict[str, set[str]] = {}
    candidates = candidates.copy()
    candidates["population_2001_int"] = pd.to_numeric(candidates.population_2001_present_persons, errors="coerce")
    def take(frame: pd.DataFrame, label: str, n: int | None = None, ascending: bool = False) -> None:
        order = frame.sort_values("population_2001_int", ascending=ascending)
        if n is not None:
            order = order.head(n)
        for row in order.itertuples(index=False):
            selected.setdefault(str(row.source_record_id_2001), set()).add(label)
    take(candidates[candidates.name_normalization_class.eq("explicit_Ukrainian_Russian_place_name_alias")], "explicit_language_alias_all")
    take(candidates[candidates.current_event_flags_raw.eq("P571")], "nonblocking_P571_all")
    take(candidates[candidates.admin_2001_status.str.contains("standalone", case=False)], "standalone_2001_hierarchy_all")
    take(candidates.nlargest(4, "population_2001_int"), "top_four_2001_population")
    take(candidates[candidates.source_type_observed_2001.eq("urban_type_settlement")].nsmallest(2, "population_2001_int"), "smallest_two_urban_type")
    wanted = set(selected)
    sample = candidates[candidates.source_record_id_2001.isin(wanted)].copy()
    sample["fixed_risk_strata"] = sample.source_record_id_2001.map(lambda x: "|".join(sorted(selected[x])))
    sample["sample_method"] = "deterministic_union_of_alias_P571_standalone_population_tail_and_top_population_strata"
    return sample.drop(columns=["population_2001_int"])


def build_packet() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    obs = read_csv(PINS["ukraine2001_reviewed_source_rows"][0])
    subject = read_csv(PINS["ukraine2001_subject_context"][0])
    points_review = read_csv(PINS["ukraine2001_point_context"][0])
    edge14 = read_csv(PINS["independent_2014_to_2021_edges"][0])
    point14 = read_csv(PINS["independent_2014_point_candidates"][0])
    stage_obs = pd.read_parquet(PINS["ukraine2001_staged_observations"][0])
    need(len(obs) == len(subject) == len(points_review) == len(stage_obs) == 27, "27-row source packet changed")
    universe01 = read_2001_universe()
    universe14, raw14 = read_2014_universe()
    counts01: dict[str, int] = {}
    counts14: dict[str, int] = {}
    for row in universe01:
        counts01[row["key"]] = counts01.get(row["key"], 0) + 1
    for row in universe14:
        counts14[row["key"]] = counts14.get(row["key"], 0) + 1
    current_ids = sorted(set(obs.current_source_record_id_2021))
    ids14 = [str(k).split(" -> ")[1] for k in obs.historical_source_candidate_key]
    con = duckdb.connect(":memory:")
    con.execute("SET memory_limit='512MB'")
    con.execute("SET threads=1")
    selected = con.execute("SELECT census_year,source_record_id,settlement_name,settlement_type,region_raw,region_norm,population,source_native_id,analysis_population_additive,is_additive_settlement_record FROM read_parquet(?) WHERE census_year=2021 AND region_raw='Республика Крым'", [str(SELECTED)]).fetchdf()
    current = con.execute("SELECT census_year,source_record_id,settlement_name,settlement_type,region_raw,population,source_native_id FROM read_parquet(?) WHERE census_year=2021 AND source_record_id IN (SELECT unnest(?))", [str(SELECTED), current_ids]).fetchdf()
    rawrownums = sorted(int(x.rsplit(":", 1)[1]) for x in current_ids)
    current_raw = con.execute("WITH x AS (SELECT row_number() OVER() AS row_number,object_level,object_name,oktmo,region,population FROM read_parquet(?)) SELECT * FROM x WHERE row_number IN (SELECT unnest(?))", [str(CURRENT_RAW), rawrownums]).fetchdf()
    base = con.execute("SELECT observation_id,record_type,entity_id,associated_census_entity_id,association_status,spatial_identity_status,identity_quality,observation_year,source_record_id,settlement_name,settlement_type,region_raw,population_value,population_scope,source_name_raw,source_admin_heading_raw,municipality_raw,source_sheet,source_row,source_path,source_sha256,latitude,longitude,coordinate_admission_status,historical_identity_admitted,population_scope_comparability_asserted,boundary_comparability_asserted FROM read_parquet(?) WHERE source_record_id IN (SELECT unnest(?))", [str(BASE_LONG), current_ids+ids14]).fetchdf()
    points = con.execute("SELECT target_source_record_id,target_year,latitude,longitude,coordinate_source,coordinate_provider,coordinate_provider_id,coordinate_provenance,coordinate_admission_status,point_origin_file,point_origin_sha256,point_origin_locator,point_origin_kind FROM read_parquet(?) WHERE regexp_replace(target_year,'\\.0$','')='2021' AND target_source_record_id IN (SELECT unnest(?))", [str(POINTS), current_ids]).fetchdf()
    con.close()
    need(len(current) == len(current_raw) == len(points) == 27, "current row/point vectors incomplete")
    c_map = current.set_index("source_record_id", drop=False)
    raw_map = current_raw.set_index("row_number", drop=False)
    b_map = base.set_index("source_record_id", drop=False)
    p_map = points.set_index("target_source_record_id", drop=False)
    sr_map = stage_obs.set_index("source_record_id", drop=False)
    subj_map = subject.set_index("historical_2001_candidate_key", drop=False)
    review_point_map = points_review.set_index("current_source_record_id_2021", drop=False)
    edge14_map = edge14.set_index("from_source_record_id", drop=False)
    point14_map = point14.set_index("target_source_record_id", drop=False)

    current_names: dict[str, int] = {}
    for r in selected.to_dict("records"):
        key = normalize_name(str(r["settlement_name"]))[0]
        current_names[key] = current_names.get(key, 0) + 1
    rows: list[dict[str, Any]] = []
    point_uses: list[dict[str, Any]] = []
    paths: list[dict[str, Any]] = []
    origin_hashes: dict[str, str] = {}
    for r in obs.to_dict("records"):
        key_full = str(r["historical_source_candidate_key"])
        chain = key_full.split(" -> ")
        need(len(chain) == 3, "source key no longer contains 2001/2014/2021 endpoints")
        sid01, sid14, sid21 = chain
        cur01 = subj_map.loc[key_full]
        hobs = sr_map.loc[sid01]
        b14 = b_map.loc[sid14]
        b21 = b_map.loc[sid21]
        c21 = c_map.loc[sid21]
        p21 = p_map.loc[sid21]
        p_review = review_point_map.loc[sid21]
        e14 = edge14_map.loc[sid14]
        pt14 = point14_map.loc[sid14]
        rowno01 = int(r["official_row_1based"])
        rowno14 = int(sid14.rsplit(":", 1)[1])
        name01 = str(r["official_row_label_raw"]).strip().split(maxsplit=1)[-1]
        name14 = str(b14.settlement_name)
        name21 = str(c21.settlement_name)
        name01_key, name_rule = normalize_name(name01)
        name14_key, name_rule14 = normalize_name(name14)
        name21_key, name_rule21 = normalize_name(name21)
        expected_class = type_key_2001(str(r["official_row_type"]))
        class14 = type_key_2014(str(b14.settlement_type))
        class21 = type_key_2021(str(c21.settlement_type))
        # Raw official Table5 and the archived Table5 row are replayed at their physical locators.
        raw01_matches = [u for u in universe01 if u["row"] == rowno01 and u["key"] == name01_key and u["type"] == expected_class]
        need(len(raw01_matches) == 1 and str(raw01_matches[0]["label"]).strip() == str(r["official_row_label_raw"]).strip() and
             int(float(raw01_matches[0]["population"])) == int(r["population_raw_table5"]),
             f"2001 source row not found in complete urban frame: {sid01}")
        need(rowno14 in raw14, f"2014 physical source row not found: {sid14}")
        raw_label14, raw_pop14 = raw14[rowno14]
        parsed14 = parse_2014_locality(str(raw_label14))
        need(parsed14 is not None, f"2014 raw row is not an individual city/PGT source row: {sid14}")
        need(parsed14[0] == class14 and normalize_name(parsed14[1])[0] == name14_key,
             f"2014 raw typed name differs from baseline for {sid14}")
        need(int(float(raw_pop14)) == int(b14.population_value), f"2014 raw source population differs for {sid14}")
        # Check source origin raw row itself, not only a selected normalized label.
        need(str(b14.source_path) == str(XLS14) and str(b14.source_sha256) == PINS["source_2014_workbook"][1],
             f"2014 baseline row is not bound to the pinned official workbook: {sid14}")
        need(int(float(b14.source_row)) == rowno14 and normalize_source_caption(str(b14.source_name_raw)) == normalize_source_caption(str(raw_label14)),
             f"2014 baseline source locator/caption differs from official row: {sid14}")

        rawno21 = int(sid21.rsplit(":", 1)[1])
        raw21 = raw_map.loc[rawno21]
        need(str(raw21.object_level) == str(r["current_2021_raw_source_object_level"]) == "Населенный пункт" and
             str(raw21.region) == str(r["current_2021_raw_source_region"]) == "Республика Крым" and
             str(raw21.oktmo) == str(r["current_native_OKTMO_2021_only"]),
             f"current 2021 source row is not a proper Crimea NP with native code: {sid21}")
        need(int(raw21.population) == int(float(r["current_2021_source_population"])),
             f"current source population context changed: {sid21}")
        need(str(c21.region_raw) == "Республика Крым" and str(c21.source_native_id) == str(raw21.oktmo),
             f"selected 2021 native source row differs from raw current publisher: {sid21}")
        need(str(b21.entity_id).startswith("settlement:") and
             str(b14.associated_census_entity_id) == str(b21.entity_id) and
             str(b14.association_status) == "reviewed_scoped_same_place_to_2021" and
             str(b14.historical_identity_admitted).lower() == "true",
             f"2014-to-current baseline association changed for {sid14}")
        need(str(e14.to_source_record_id) == sid21 and
             str(e14.review_status) == "independently_eligible_for_scoped_application_candidate_only" and
             int(float(e14.source_population)) == int(b14.population_value),
             f"independent 2014→2021 source candidate differs for {sid14}")
        need(str(pt14.source_2021_carrier_target_source_record_id) == sid21 and
             str(pt14.coordinate_use_review_status) == "independently_eligible_for_scoped_application_candidate_only",
             f"2014 point-context candidate differs for {sid14}")
        need(math.isclose(float(p21.latitude), float(p_review.point_latitude_current_representative), abs_tol=1e-10) and
             math.isclose(float(p21.longitude), float(p_review.point_longitude_current_representative), abs_tol=1e-10),
             f"accepted current point differs from reviewed raw carrier: {sid21}")
        origin = Path(str(p21.point_origin_file))
        need(origin.is_file(), f"accepted current point origin missing: {sid21}")
        if str(origin) not in origin_hashes:
            origin_hashes[str(origin)] = sha256(origin)
        need(origin_hashes[str(origin)] == str(p21.point_origin_sha256) == str(p_review.point_origin_sha256),
             f"accepted current point origin hash mismatch: {sid21}")
        point_distance = distance_km(float(b14.latitude), float(b14.longitude), float(p21.latitude), float(p21.longitude))
        event_flag = str(r["current_event_or_date_claim_flags_raw"])
        hard_event = bool({p for p in event_flag.split("|") if p} - {"P571"})
        # Admin structures are retained verbatim. Standalone source rows carry an explicit note and are not forward-filled.
        admin01 = str(r["source_admin_parent_row_raw"])
        admin14 = str(b14.source_admin_heading_raw)
        admin_ok = bool(admin01.strip()) and bool(admin14.strip())
        key_count01 = counts01.get(name01_key, 0)
        key_count14 = counts14.get(name14_key, 0)
        key_count21 = current_names.get(name21_key, 0)
        key_count01_type = sum(1 for u in universe01 if u["key"] == name01_key and u["type"] == expected_class)
        key_count14_type = sum(1 for u in universe14 if u["key"] == name14_key and u["type"] == class14)
        key_count21_type = int(((selected.settlement_name.map(lambda x: normalize_name(str(x))[0]) == name21_key) &
                                selected.settlement_type.map(lambda x: type_key_2021(str(x)) == class21)).sum())
        predicates = {
            "official_2001_individual_urban_row_replayed": str(r["row_is_individual_urban_locality_not_aggregate"]).lower() == "true",
            "official_2014_individual_city_or_PGT_row_replayed": parsed14 is not None,
            "same_place_name_after_explicit_source_language_normalization": name01_key == name14_key == name21_key,
            "same_observed_city_or_urban_type_class": expected_class == class14 == class21,
            "whole_2001_typed_name_unique": key_count01_type == 1,
            "whole_2014_typed_name_unique": key_count14_type == 1,
            "current_selected_Crimea_typed_name_unique": key_count21_type == 1,
            "source_admin_hierarchies_preserved_no_forward_fill": admin_ok,
            "current_native_2021_NP_and_region_replayed": bool(str(raw21.oktmo).strip()) and str(raw21.object_level) == "Населенный пункт",
            "existing_scoped_2014_to_2021_identity_present": str(b14.association_status) == "reviewed_scoped_same_place_to_2021",
            "accepted_current_point_origin_replayed": origin_hashes[str(origin)] == str(p21.point_origin_sha256),
            "2014_to_current_point_context_within_5km": point_distance <= 5.0,
            "no_known_move_successor_or_merge_hold": not hard_event and str(r["unresolved_successor_or_merger_event_hold"]).lower() == "false",
        }
        status, holds = evaluate_rule(predicates)
        normalization_class = "explicit_Ukrainian_Russian_place_name_alias" if "explicit_Ukrainian_Russian_place_name_alias" in {name_rule, name_rule14, name_rule21} else ("Ukrainian_Russian_orthography" if "Ukrainian_Russian_orthography" in {name_rule, name_rule14, name_rule21} else "exact_name")
        risk = []
        if normalization_class == "explicit_Ukrainian_Russian_place_name_alias": risk.append("explicit cross-language alias")
        if event_flag == "P571": risk.append("P571-only date flag; not used as identity proof")
        if "standalone" in admin01.lower(): risk.append("2001 locality row is standalone; no parent forward-filled")
        if expected_class == "urban_type_settlement": risk.append("urban-type locality")
        if point_distance > 0.1: risk.append("2014 and current point context differ")
        rows.append({
            "source_record_id_2001": sid01, "year_2001": 2001, "year_2014": 2014, "year_2021": 2021,
            "source_record_id_2014": sid14, "current_source_record_id_2021": sid21,
            "current_2021_entity_id_context_only": str(b21.entity_id),
            "name_2001_raw": name01, "name_2014_raw": str(b14.settlement_name), "name_2021_selected_raw": str(c21.settlement_name),
            "name_2001_key": name01_key, "name_2014_key": name14_key, "name_2021_key": name21_key,
            "name_normalization_class": normalization_class,
            "official_name_2001_raw_row": str(r["official_row_label_raw"]),
            "official_name_2014_raw_row": str(raw_label14),
            "type_2001": str(r["official_row_type"]), "type_2014": str(b14.settlement_type), "type_2021": str(c21.settlement_type),
            "source_type_observed_2001": expected_class, "source_type_observed_2014": class14, "source_type_observed_2021": class21,
            "type_change_observed": expected_class != class14 or class14 != class21,
            "admin_2001_parent_raw": admin01, "admin_2001_grouping_raw": str(r["source_grouping_row_raw"]),
            "admin_2001_status": "standalone_individual_locality_row_no_parent_forward_fill" if "standalone" in admin01.lower() else "published_parent_row_retained",
            "admin_2014_heading_raw": admin14, "admin_2014_municipality_raw": str(b14.municipality_raw or ""),
            "admin_hierarchy_comparison": "separate source hierarchies retained; no administrative-boundary continuity inferred",
            "population_2001_present_persons": int(r["population_raw_table5"]), "population_2014_source_value": int(b14.population_value),
            "population_2014_scope_raw": str(b14.population_scope), "population_2021_selected_context": int(c21.population),
            "population_boundary_comparability_asserted": False,
            "2001_native_OKTMO_or_KOATUU_published": False, "old_code_missing_is_identity_veto": False,
            "current_native_2021_OKTMO_literal_only": str(c21.source_native_id),
            "current_QID_context_only": str(r["current_qid_context"]), "current_QID_used_as_2001_identity_key": False,
            "current_event_flags_raw": event_flag, "P571_only_nonblocking": event_flag == "P571",
            "move_successor_or_merge_event_hold": hard_event,
            "2014_to_2021_existing_association": str(b14.association_status),
            "accepted_current_latitude": float(p21.latitude), "accepted_current_longitude": float(p21.longitude),
            "accepted_current_point_provider": str(p21.coordinate_provider), "accepted_current_point_provider_id": str(p21.coordinate_provider_id),
            "current_point_origin_file": str(p21.point_origin_file), "current_point_origin_sha256": str(p21.point_origin_sha256),
            "current_point_origin_locator": str(p21.point_origin_locator), "current_point_origin_kind": str(p21.point_origin_kind),
            "2014_point_context_to_current_accepted_point_km": point_distance,
            "current_point_is_2001_measurement": False,
            "rule_predicates_json": json.dumps(predicates, ensure_ascii=False, sort_keys=True),
            "predicate_pass_count": sum(predicates.values()), "predicate_count": len(predicates),
            "whole_2001_name_count_all_types": key_count01, "whole_2014_name_count_all_types": key_count14,
            "whole_2021_name_count_all_types": key_count21,
            "candidate_status": status, "hold_reasons_json": json.dumps(holds, ensure_ascii=False),
            "risk_flags_json": json.dumps(risk, ensure_ascii=False),
            "administrative_boundary_comparability_asserted": False,
            "legal_type_or_event_effective_date": None, "historical_coordinate_measurement_claimed": False,
            "strict_Russian_2002_2010_2021_coverage": "outside_coverage_Ukraine_2001_is_not_Russian_census_2002_or_2010",
        })
        point_uses.append({
            "point_use_candidate_id": f"ukraine2001-retrospective-point:{len(point_uses)+1:02d}",
            "target_source_record_id": sid01, "target_year": 2001,
            "current_carrier_source_record_id": sid21, "intervening_source_record_id": sid14,
            "latitude": float(p21.latitude), "longitude": float(p21.longitude),
            "candidate_status": "candidate_retrospective_point_use_pending_identity_review" if status == "eligible_for_fixed_independent_review" else "held_point_use_identity_rule_hold",
            "coordinate_temporal_basis": "accepted_current_representative_point_inferred_retrospectively_if_2001_to_2014_physical_continuity_is_reviewed",
            "current_point_origin_file": str(p21.point_origin_file), "current_point_origin_sha256": str(p21.point_origin_sha256),
            "current_point_origin_locator": str(p21.point_origin_locator), "current_point_origin_kind": str(p21.point_origin_kind),
            "historical_measurement_claimed": False, "historical_provider_binding_asserted": False,
            "native_historical_code_binding_asserted": False, "administrative_boundary_comparability_asserted": False,
            "identity_review_gate_required": True,
        })
        paths.append({
            "path_candidate_id": f"ukraine2001-temporal-path:{len(paths)+1:02d}",
            "source_record_id_2001": sid01, "population_2001_present_persons": int(r["population_raw_table5"]),
            "source_record_id_2014": sid14, "population_2014_source_context": int(b14.population_value),
            "source_record_id_2021": sid21, "population_2021_selected_context": int(c21.population),
            "edge_2001_to_2014_status": status,
            "edge_2014_to_2021_status": "already_scoped_associated_in_frozen_baseline",
            "full_same_place_path_admitted": False,
            "path_status": "candidate 2001→2014 edge pending independent review; 2014→2021 association is already in baseline",
            "population_boundary_comparability_asserted": False,
            "population_measure_2001": "present_population", "2014_grain": str(b14.population_scope),
            "population_series_interpretation": "values shown side-by-side only; no direct comparability or sum",
        })

    rule_df = pd.DataFrame(rows)
    edge_df = pd.DataFrame([{**r,
        "edge_id": f"ukraine2001-to-2014-scoped-continuity:{i+1:02d}",
        "from_source_record_id": r["source_record_id_2001"], "from_year": 2001,
        "to_source_record_id": r["source_record_id_2014"], "to_year": 2014,
        "relation": "same_named_physical_urban_locality_continuity_candidate",
        "decision_status": "candidate_for_independent_fixed_risk_review_not_admitted",
        "native_code_binding_asserted": False, "current_QID_binding_asserted": False,
        "population_boundary_comparability_asserted": False, "legal_effective_date": None,
        "historical_coordinate_asserted": False, "strict_Russian_2002_2010_2021_chain": False,
    } for i, r in enumerate(rows)])
    point_df = pd.DataFrame(point_uses)
    path_df = pd.DataFrame(paths)
    risk_df = fixed_risk_sample(rule_df)
    need(len(rule_df) == len(edge_df) == len(point_df) == len(path_df) == 27, "candidate packet row counts changed")
    need(rule_df.candidate_status.eq("eligible_for_fixed_independent_review").sum() + rule_df.candidate_status.eq("held_not_rule_admissible").sum() == 27,
         "candidate/hold status vector is inconsistent")
    need(not rule_df.population_boundary_comparability_asserted.any() and not rule_df.historical_coordinate_measurement_claimed.any(),
         "out-of-scope boundary/coordinate claims present")
    need(edge_df.decision_status.eq("candidate_for_independent_fixed_risk_review_not_admitted").all(),
         "candidate packet must remain unapplied")
    return rule_df, edge_df, point_df, path_df, risk_df, {"origin_hashes": origin_hashes}


def write_packet(out: Path, input_hashes: dict[str, str], packet: tuple[pd.DataFrame, ...], auxiliary: dict[str, Any]) -> dict[str, Any]:
    need(not out.exists(), f"refusing to overwrite output directory: {out}")
    out.mkdir(parents=True)
    names = ("rule_evaluation", "candidate_2001_to_2014_edges", "candidate_retrospective_2001_points", "available_year_path_candidates", "fixed_independent_risk_sample")
    outputs = {}
    for name, frame in zip(names, packet):
        for ext in ("csv", "parquet"):
            p = out / f"{name}.{ext}"
            frame.to_csv(p, index=False, quoting=csv.QUOTE_MINIMAL) if ext == "csv" else frame.to_parquet(p, index=False)
            outputs[f"{name}_{ext}"] = {"path": str(p), "sha256": sha256(p), "rows": len(frame), "bytes": p.stat().st_size}
    test_path = Path(__file__).parent / "tests/test_stage_ukraine2001_to_2014_temporal_rule.py"
    manifest = {"status": "frozen_temporal_rule_packet_inputs", "inputs": {n: {"path": str(p), "sha256": h} for n, (p, h) in PINS.items()},
                "producer_code": {"path": str(Path(__file__).resolve()), "sha256": sha256(Path(__file__).resolve())},
                "producer_tests": {"path": str(test_path.resolve()), "sha256": sha256(test_path)},
                "candidate_rule_text": RULE_TEXT, "explicit_place_name_aliases": EXPLICIT_NAME_ALIASES,
                "current_point_origin_hashes_replayed": auxiliary["origin_hashes"]}
    mp = out / "application_input_manifest.json"
    mp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    outputs["application_input_manifest"] = {"path": str(mp), "sha256": sha256(mp), "bytes": mp.stat().st_size}
    neligible = int(packet[0].candidate_status.eq("eligible_for_fixed_independent_review").sum())
    receipt = {
        "status": "scoped_2001_to_2014_temporal_rule_packet_ready_for_independent_review_not_admitted",
        "rule_text": RULE_TEXT, "official_source_rows_reviewed": 27,
        "eligible_candidate_edges_for_review": neligible,
        "held_rows": len(packet[0])-neligible,
        "existing_2014_to_2021_associations_verified_in_baseline": 27,
        "candidate_retrospective_2001_point_uses": len(packet[2]),
        "fixed_risk_sample_rows": len(packet[4]),
        "explicit_official_Ukrainian_Russian_alias_cases": int(packet[0].name_normalization_class.eq("explicit_Ukrainian_Russian_place_name_alias").sum()),
        "observed_type_changes": int(packet[0].type_change_observed.sum()),
        "all_2001_source_rows_have_no_native_locality_code": True,
        "old_code_absence_as_identity_veto": False,
        "identity_admissions_2001_to_2014": 0, "identity_admissions_2001_to_2021": 0,
        "historical_coordinates_asserted": False, "historical_provider_binding_asserted": False,
        "administrative_boundary_comparability_asserted": False, "population_boundary_comparability_asserted": False,
        "current_QID_used_as_2001_identity_key": False,
        "P571_used_as_identity_evidence": 0,
        "Russian_2002_2010_census_coverage_or_gain": 0,
        "selected_frame_or_canonical_graph_mutated": False, "long_table_appended": False,
        "population_sums_computed_as_cross_year_comparisons": False,
        "input_hashes": input_hashes, "outputs": outputs,
        "limitations": "Only source-row continuity candidates are staged. Actual 2001 present-population, 2014 source values and 2021 selected populations have distinct measurement grains and are not claimed comparable. Current accepted point coordinates are modern and retrospective-use candidates only. The rule does not assert 2001 codes, P1365-derived historical identity, legal status dates, polygons, or Russian census coverage.",
    }
    rp = out / "review_packet_receipt.json"
    rp.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    receipt["receipt_sha256"] = sha256(rp)
    return receipt


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out-dir", type=Path, default=OUT_DEFAULT)
    ap.add_argument("--write-staged", action="store_true")
    args = ap.parse_args()
    input_hashes = {}
    input_hashes = verify_inputs()
    packet = build_packet()
    summary = {"status": "validated_temporal_rule_candidate_packet_no_admission", "source_rows": len(packet[0]),
               "candidate_edges": int(packet[0].candidate_status.eq("eligible_for_fixed_independent_review").sum()),
               "held_rows": int(packet[0].candidate_status.eq("held_not_rule_admissible").sum()),
               "existing_2014_to_2021_mappings": 27, "retrospective_point_candidates": len(packet[2]),
               "fixed_risk_sample_rows": len(packet[4]), "identity_admissions": 0,
               "graph_or_long_mutated": False, "Russian_2002_2010_2021_gain": 0}
    if args.write_staged:
        receipt = write_packet(args.out_dir, input_hashes, packet[:5], packet[5])
        summary.update(output_dir=str(args.out_dir), review_packet_receipt_sha256=receipt["receipt_sha256"])
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
