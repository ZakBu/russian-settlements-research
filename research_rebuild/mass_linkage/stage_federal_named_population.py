"""Stage reviewed 2021 St Petersburg/Sevastopol named locality populations.

Population evidence is a source-only candidate delta. The module preserves the
federal aggregate controls and does not select observations or admit points.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path('/workspace')
CONT = ROOT / 'settlements-work/continuation_20261003'
DEFAULT_REVIEW = CONT / 'federal_named_np_primary_review_v1/review.json'
DEFAULT_WORKBOOK = ROOT / 'settlements-raw/data/raw/rosstat_2021/Tom1_tab-5_VPN-2020.xlsx'
DEFAULT_RAW_2021 = ROOT / 'settlements-raw/data/interim/2021_tochno/data_allsettlements_anon_156_v20251217.parquet'
REVIEW_SHA256 = 'fcaa8194eb41d2d2a65d1d14b672f8f0f9480d2395633d52e739a3a0d488f012'
WORKBOOK_SHA256 = '0b232b3d2ab5daa231568acc719ac6fda4fb0979a01a0da6bacc69be6f252474'
RAW_2021_SHA256 = '86c197cd522e0b63669e9c6e7f43fd3d82b3704c6a126c800a9968ecd16cae14'
REVIEW_ID = 'federal_named_np_primary_grain_review_2021'
SHEET = 'таб. 5'
APPROVED = 'approve_primary_atomic_np_population'
HELD = 'hold_grain'
AGGREGATE_CONTROLS = {
    '2021_federal_aggregate_count': 19159843,
    'Moscow_aggregate': 13010112,
    'Saint_Petersburg_aggregate': 5601911,
    'Sevastopol_aggregate': 547820,
}
RAW_POINT_COLUMNS = [
    'object_level', 'object_name', 'region', 'mun_upper', 'settlement',
    'settlement_type_full_dadata', 'fias_id_dadata', 'fias_level_dadata',
    'latitude_dadata', 'longitude_dadata', 'okato_dadata', 'oktmo_dadata',
]


def sha(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(4 * 1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def norm(value) -> str:
    if value is None or pd.isna(value):
        return ''
    return ' '.join(str(value).casefold().replace('ё', 'е').split())


def label_parts(city: str, label: str) -> tuple[str, str | None]:
    """Parse literal settlement name/type from reviewed printed label only."""
    text = str(label).strip()
    typed = re.match(r'^(пос[её]лок|г\.)\s*(.+)$', text, flags=re.I)
    if typed:
        kind = 'поселок' if norm(typed.group(1)).startswith('пос') else 'город'
        return typed.group(2).strip(), kind
    if city == 'Севастополь':
        match = re.match(r'^Городское население\s*-\s*(?:пос[её]лок\s+)?(.+)$', text, flags=re.I)
        if match:
            name = match.group(1).strip()
            if norm(name).startswith('г. '):
                return name[2:].strip(), 'город'
            kind = 'поселок' if norm(text).find('поселок') >= 0 else None
            return name, kind
        match = re.search(r'-\s*г\.\s*([^;]+)$', text, flags=re.I)
        if match:
            return match.group(1).strip(), 'город'
    raise ValueError(f'cannot parse named locality from printed label: {label!r}')


def review_records(review_path: Path, workbook_path: Path,
                   expected_review_sha: str = REVIEW_SHA256) -> tuple[dict, list[dict], list[dict]]:
    if sha(review_path) != expected_review_sha:
        raise ValueError('independent federal named-NP review SHA-256 mismatch')
    review = json.loads(review_path.read_text(encoding='utf-8'))
    if review.get('review') != REVIEW_ID:
        raise ValueError('unexpected independent review ID')
    if sha(workbook_path) != WORKBOOK_SHA256:
        raise ValueError('official Table 5 workbook SHA-256 mismatch')
    source = review.get('source', {})
    if source.get('sha256') != WORKBOOK_SHA256 or source.get('sheet') != SHEET:
        raise ValueError('review workbook pin or sheet does not match expected source')
    if review.get('population_controls_preserved') != AGGREGATE_CONTROLS:
        raise ValueError('review aggregate control values differ from frozen controls')
    records = review.get('records')
    if not isinstance(records, list) or len(records) != 33:
        raise ValueError('review must contain the exact 33 bounded candidate rows')
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    if SHEET not in wb.sheetnames:
        raise ValueError('official Table 5 sheet missing')
    ws = wb[SHEET]
    approved, held = [], []
    seen_ids = set()
    for rec in records:
        rownum = int(rec['excel_row'])
        if rownum <= 0 or rownum > ws.max_row:
            raise ValueError(f'review Excel row outside worksheet: {rownum}')
        source_id = f'ROSSTAT2021:T5:{SHEET}:r{rownum}'
        if source_id != rec.get('source_record_id') or source_id in seen_ids:
            raise ValueError(f'bad or duplicate exact source record ID at row {rownum}')
        seen_ids.add(source_id)
        raw_cells = [ws.cell(rownum, col).value for col in range(1, 7)]
        if json.loads(rec.get('raw_cells_json', '[]')) != raw_cells:
            raise ValueError(f'official workbook raw cells differ from review at row {rownum}')
        if raw_cells[0] != rec.get('label_raw'):
            raise ValueError(f'official workbook label differs from review at row {rownum}')
        if raw_cells[1:4] != [rec.get('population'), rec.get('men'), rec.get('women')]:
            raise ValueError(f'official workbook population cells differ from review at row {rownum}')
        indent = float(ws.cell(rownum, 1).alignment.indent or 0)
        parent_row = None
        for candidate in range(rownum - 1, 0, -1):
            if ws.cell(candidate, 1).value is None:
                continue
            parent_indent = float(ws.cell(candidate, 1).alignment.indent or 0)
            if parent_indent < indent:
                parent_row = candidate
                break
        actual_parent_label = ws.cell(parent_row, 1).value if parent_row is not None else None
        if parent_row != rec.get('parent_row') or actual_parent_label != rec.get('parent_label_raw'):
            raise ValueError(f'official workbook parent row differs from review at row {rownum}')
        status = rec.get('decision')
        if status == APPROVED:
            name, kind = label_parts(str(rec['city']), str(raw_cells[0]))
            population, men, women = (int(raw_cells[i]) for i in (1, 2, 3))
            if men + women != population:
                raise ValueError(f'male+female counts fail for exact source row {rownum}')
            approved.append({
                'source_record_id': source_id,
                'census_year': 2021,
                'source_file': 'data/raw/rosstat_2021/Tom1_tab-5_VPN-2020.xlsx',
                'source_sheet': SHEET,
                'source_row_excel_1based': rownum,
                'source_locator': f'{SHEET}!A{rownum}:F{rownum}; worksheet row is 1-based',
                'source_sha256': WORKBOOK_SHA256,
                'federal_city': rec['city'],
                'label_raw': raw_cells[0],
                'settlement_name_from_label': name,
                'settlement_type_from_label': kind,
                'population': population,
                'men': men,
                'women': women,
                'raw_cells_json': json.dumps(raw_cells, ensure_ascii=False, separators=(',', ':')),
                'parent_row_excel_1based': rec.get('parent_row'),
                'parent_label_raw': rec.get('parent_label_raw'),
                'row_population_grain_review': rec.get('row_population_grain'),
                'independent_review_decision': status,
                'review_id': REVIEW_ID,
                'review_sha256': expected_review_sha,
                'candidate_population_observation': True,
                'selected_into_primary_population': False,
                'population_admission_allowed': False,
                'coordinates_or_identity_inferred': False,
            })
        elif status == HELD:
            held.append({
                'source_record_id': source_id,
                'census_year': 2021,
                'source_row_excel_1based': rownum,
                'federal_city': rec.get('city'),
                'label_raw': raw_cells[0],
                'population': raw_cells[1],
                'hold_reason': rec.get('grain_reason'),
                'review_id': REVIEW_ID,
                'review_sha256': expected_review_sha,
                'selected_into_primary_population': False,
            })
        else:
            raise ValueError(f'unexpected review decision at row {rownum}: {status!r}')
    if len(approved) != 32 or len(held) != 1 or int(held[0]['source_row_excel_1based']) != 7615:
        raise ValueError('review decisions differ from expected 32 approved and Kronstadt hold')
    if sum(row['population'] for row in approved) != 878643:
        raise ValueError('approved source population total differs from frozen review control')
    return review, approved, held


def raw_source_point_candidates(observations: list[dict], raw_path: Path,
                                expected_sha: str = RAW_2021_SHA256) -> list[dict]:
    """Find unique exact city-region named physical rows in the cached provider source."""
    if sha(raw_path) != expected_sha:
        raise ValueError('cached 2021 point-source SHA-256 mismatch')
    data = pd.read_parquet(raw_path, columns=RAW_POINT_COLUMNS)
    point_rows = []
    for obs in observations:
        name_key, region_key = norm(obs['settlement_name_from_label']), norm(obs['federal_city'])
        subset = data[
            data['object_level'].astype('string').map(norm).eq('населенный пункт')
            & data['region'].astype('string').map(norm).eq(region_key)
            & data['settlement'].astype('string').map(norm).map(
                lambda value: re.sub(r'^(?:пос[её]лок|г\.)\s*', '', value).strip()
            ).eq(name_key)
        ]
        if len(subset) == 1:
            ix = int(subset.index[0])
            row = subset.iloc[0]
            point_rows.append({
                'population_source_record_id': obs['source_record_id'],
                'federal_city': obs['federal_city'],
                'settlement_name_from_population_label': obs['settlement_name_from_label'],
                'point_candidate_status': 'unique_exact_name_region_raw_2021_physical_row_pending_review',
                'point_candidate_source_record_id': f'2021:data_allsettlements_anon_156_v20251217.parquet:parquet:{ix + 1}',
                'point_candidate_source_file': str(raw_path),
                'point_candidate_source_sha256': expected_sha,
                'point_candidate_row_1based': ix + 1,
                'point_candidate_object_level_raw': row['object_level'],
                'point_candidate_object_name_raw': row['object_name'],
                'point_candidate_settlement_raw': row['settlement'],
                'point_candidate_region_raw': row['region'],
                'point_candidate_parent_context_raw': row['mun_upper'],
                'point_candidate_fias_id_raw': row['fias_id_dadata'],
                'point_candidate_fias_level_raw': row['fias_level_dadata'],
                'point_candidate_okato_raw': row['okato_dadata'],
                'point_candidate_oktmo_raw': row['oktmo_dadata'],
                'point_candidate_latitude_raw': row['latitude_dadata'],
                'point_candidate_longitude_raw': row['longitude_dadata'],
                'point_binding_admitted': False,
                'identity_admitted': False,
                'coordinate_admitted': False,
            })
        else:
            point_rows.append({
                'population_source_record_id': obs['source_record_id'],
                'federal_city': obs['federal_city'],
                'settlement_name_from_population_label': obs['settlement_name_from_label'],
                'point_candidate_status': ('no_exact_name_region_raw_2021_physical_row_found'
                                           if len(subset) == 0 else 'multiple_exact_name_region_rows_hold'),
                'exact_candidate_row_count': int(len(subset)),
                'point_binding_admitted': False,
                'identity_admitted': False,
                'coordinate_admitted': False,
            })
    return point_rows


def stage(review_path: Path = DEFAULT_REVIEW, workbook_path: Path = DEFAULT_WORKBOOK,
          raw_2021_path: Path = DEFAULT_RAW_2021,
          output_dir: Path | None = None) -> dict:
    if output_dir is None:
        raise ValueError('an explicit new output directory is required')
    review_path, workbook_path = Path(review_path), Path(workbook_path)
    raw_2021_path, output_dir = Path(raw_2021_path), Path(output_dir)
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f'refusing to overwrite nonempty stage directory: {output_dir}')
    _, approved, held = review_records(review_path, workbook_path)
    point_candidates = raw_source_point_candidates(approved, raw_2021_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(approved).to_parquet(output_dir / 'population_observations.parquet', index=False)
    pd.DataFrame(held).to_parquet(output_dir / 'held_review_rows.parquet', index=False)
    pd.DataFrame(point_candidates).to_parquet(output_dir / 'point_binding_candidates.parquet', index=False)
    counts = {}
    for city in ('Санкт-Петербург', 'Севастополь'):
        rows = [row for row in approved if row['federal_city'] == city]
        counts[city] = {'rows': len(rows), 'population': sum(row['population'] for row in rows)}
    outputs = {}
    for filename in ('population_observations.parquet', 'held_review_rows.parquet',
                     'point_binding_candidates.parquet'):
        path = output_dir / filename
        outputs[filename] = {'path': str(path), 'rows': len(pd.read_parquet(path)), 'sha256': sha(path)}
    point_counts = {}
    for row in point_candidates:
        key = row['point_candidate_status']
        point_counts[key] = point_counts.get(key, 0) + 1
    receipt = {
        'status': 'federal_named_np_population_source_delta_staged_pending_selection_and_review',
        'review_id': REVIEW_ID,
        'review_sha256': REVIEW_SHA256,
        'review_path': str(review_path),
        'workbook_path': str(workbook_path),
        'workbook_sha256': WORKBOOK_SHA256,
        'workbook_sheet': SHEET,
        'raw_2021_point_source_path': str(raw_2021_path),
        'raw_2021_point_source_sha256': RAW_2021_SHA256,
        'approved_observation_rows': len(approved),
        'held_review_rows': len(held),
        'approved_population_total': sum(row['population'] for row in approved),
        'approved_population_by_federal_city': counts,
        'frozen_federal_aggregate_controls_preserved': AGGREGATE_CONTROLS,
        'point_candidate_status_counts': point_counts,
        'point_search_scope': 'Exact name after literal printed type prefix and exact federal-city region within cached 2021 raw provider rows whose object_level is Населенный пункт. Candidates are not identity or point admissions.',
        'primary_selection_or_admission_performed': False,
        'existing_federal_aggregate_controls_modified': False,
        'candidate_observations_may_be_added_without_root_selection': False,
        'outputs': outputs,
        'limits': [
            'Population candidates are a source-only delta and would double-count if added beside existing aggregate controls.',
            'Kronstadt remains held because the source row combines district and city and does not isolate physical locality population.',
            'No coordinates or identity are assigned by this staging rule.',
            'The separate point screen is a raw-source name/region lookup only; every candidate remains unadmitted.',
            'No Moscow city-core population or aggregate residual was inferred.',
        ],
    }
    (output_dir / 'receipt.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return receipt


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--review', type=Path, default=DEFAULT_REVIEW)
    parser.add_argument('--workbook', type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument('--raw-2021', type=Path, default=DEFAULT_RAW_2021)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args(argv)
    print(json.dumps(stage(args.review, args.workbook, args.raw_2021, args.output),
                     ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
