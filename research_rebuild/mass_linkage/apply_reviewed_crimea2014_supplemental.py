#!/usr/bin/env python3
"""Prepare a scoped supplement for the independently reviewed 2014 Crimea rows.

This utility combines three separately reviewed city rows with nineteen direct
settlement corridors. It keeps eleven literal dash rows as presence-only rows,
uses only the accepted current 2021 carrier points, and never changes the
canonical 2002/2010/2021 graph. The CLI defaults to validation/preview; writing
staged ledgers requires --write-staged. Appending to a long table is deliberately
not implemented here; the root integration step owns that action.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from openpyxl import load_workbook

CITY_RECEIPT_SHA256 = "461af33223ac50d19d1295da806edee38abf2849383287cac23e5e306891eaca"
DIRECT19_RECEIPT_SHA256 = "efd493d66a8b2fda324f35c2658856b0e9a9e580fb03a3547c442febc3e35e91"
WORKBOOK_SHA256 = "35e6acf1e5ecb66c23355591a0ddf630a70eefbedfca0ba470d6b0004baba06e"
CURRENT_RAW_SHA256 = "86c197cd522e0b63669e9c6e7f43fd3d82b3704c6a126c800a9968ecd16cae14"
POINT_LEDGER_SHA256 = "2215702e772417bb512da8c0a6c0d31ad9783c6a5345d4c0bbfe18403e11ad22"
BASE_LONG_SHA256 = "be04c25f2cdcc1614e2c17865ca8a96593e4308f36ad344281062c37c7bd6360"
IDENTITY_GRAPH_SHA256 = "18bf1edd259808c9e881ac11fa40be26c88be4fb27e1ac735c59e083d6dcd312"
CITY_OBSERVATIONS_SHA256 = "e302557b3d6b1b91b96b61d00c36ef2a344c610c7fa68c9230453289bb5ce0d0"
CITY_EDGES_SHA256 = "9d35acaadf86b7f8b4979734cfdbc9ef332873abd27dc95602b6c196a9726a67"
DIRECT19_KEYS_SHA256 = "9c6dfcc93faff6143d0f8533f31bfe73c8fbf9ff59e5d469695ff0b2256d0e20"
DIRECT19_ROWS_SHA256 = "dcae1038aad65e1f12148b309945ea67315dff92de315e849821d50e34e6c49f"

_DASHES = {"-", "—", "–", "−"}
_ACCEPTED_POINT_STATUSES = {"reviewed_rule_accepted", "reviewed_extension_rule_accepted"}


def sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _need(ok: bool, message: str) -> None:
    if not ok:
        raise ValueError(message)


def verify_hash(path: str | Path, expected: str, label: str) -> str:
    actual = sha256(path)
    _need(actual == str(expected), f"{label} hash mismatch: expected {expected}, got {actual}")
    return actual


def _read_csv(path: str | Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", engine="python")


def literal_integer(value: Any) -> int | None:
    """Parse only literal integer cells; blank and dash remain missing."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if math.isfinite(value) and value.is_integer() else None
    raw = str(value).strip()
    if not raw or raw in _DASHES:
        return None
    try:
        number = Decimal(raw.replace(" ", "").replace("\u00a0", ""))
    except InvalidOperation:
        return None
    return int(number) if number == number.to_integral_value() else None


def verify_manifest(manifest: dict[str, Any]) -> dict[str, Path]:
    """Validate frozen packet, workbook, baseline, selected rows and point ledger pins."""
    required = {
        "base_long", "selected_population", "current_raw", "identity_graph", "point_ledger", "source_workbook",
        "city_receipt", "city_observations", "city_edges", "direct19_receipt",
        "direct19_rows", "direct19_keys",
    }
    entries = manifest.get("inputs", {})
    _need(required <= set(entries), f"manifest missing input pins: {sorted(required - set(entries))}")
    paths: dict[str, Path] = {}
    for name in sorted(required):
        entry = entries[name]
        path = Path(entry["path"])
        _need(path.is_file(), f"missing input {name}: {path}")
        verify_hash(path, entry["sha256"], name)
        paths[name] = path
    for name, expected in (
        ("base_long", BASE_LONG_SHA256), ("current_raw", CURRENT_RAW_SHA256),
        ("identity_graph", IDENTITY_GRAPH_SHA256), ("point_ledger", POINT_LEDGER_SHA256), ("source_workbook", WORKBOOK_SHA256),
        ("city_receipt", CITY_RECEIPT_SHA256), ("city_observations", CITY_OBSERVATIONS_SHA256),
        ("city_edges", CITY_EDGES_SHA256), ("direct19_receipt", DIRECT19_RECEIPT_SHA256),
        ("direct19_rows", DIRECT19_ROWS_SHA256), ("direct19_keys", DIRECT19_KEYS_SHA256),
    ):
        _need(sha256(paths[name]) == expected, f"{name} does not match the reviewed frozen bytes")

    city_receipt = json.loads(paths["city_receipt"].read_text(encoding="utf-8"))
    _need(city_receipt.get("status") == "independently_reviewed_eligible_city_continuity_candidates_not_applied",
          "unexpected three-city review status")
    _need(city_receipt.get("counts", {}).get("eligible_2014_to_2021_physical_city_identity_candidates") == 3,
          "three-city candidate count mismatch")
    _need(city_receipt.get("counts", {}).get("admissions") == 0, "review packet must remain unapplied")
    direct_receipt = json.loads(paths["direct19_receipt"].read_text(encoding="utf-8"))
    _need(direct_receipt.get("status") == "independent_review_completed_candidate_corridors_only_no_application",
          "unexpected direct19 review status")
    counts = direct_receipt.get("counts", {})
    _need(counts.get("candidate_rows_reviewed") == counts.get("identity_corridors_eligible") == 19,
          "direct19 candidate count mismatch")
    _need(counts.get("literal_numeric_2014_population_rows") == 8 and counts.get("literal_dash_2014_population_rows") == 11,
          "direct19 numeric/dash split changed")
    _need(counts.get("sum_numeric_2014_source_cells_only") == 4164, "direct19 literal source-cell sum changed")
    return paths


def replay_source_cells(rows: pd.DataFrame, workbook_path: str | Path,
                        source_row_field: str, caption_field: str, population_field: str) -> dict[str, list[Any]]:
    """Read the pinned workbook sheet once and replay exact source captions/count cells."""
    _need(not rows.empty, "no source rows to replay")
    _need(set(rows.source_sheet) == {"pub-01-03"}, "unexpected source workbook sheet")
    row_numbers = [int(x) for x in rows[source_row_field]]
    _need(len(row_numbers) == len(set(row_numbers)), "duplicate 2014 workbook row locator")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        _need("pub-01-03" in workbook.sheetnames, "official 2014 workbook sheet missing")
        targets = set(row_numbers)
        found: dict[int, list[Any]] = {}
        ws = workbook["pub-01-03"]
        max_row = max(targets)
        for number, cells in enumerate(ws.iter_rows(values_only=True), 1):
            if number in targets:
                found[number] = list(cells)
            if number >= max_row:
                break
        _need(targets == set(found), "one or more 2014 source rows are absent from the workbook")
        out: dict[str, list[Any]] = {}
        for row in rows.itertuples(index=False):
            mapping = row._asdict()
            number = int(mapping[source_row_field])
            cells = found[number]
            caption = str(cells[0]).strip()
            _need(" ".join(caption.split()) == " ".join(str(mapping[caption_field]).strip().split()),
                  f"2014 source caption mismatch at row {number}")
            population_cell = cells[1]
            _need(str(population_cell).strip() == str(mapping[population_field]).strip(),
                  f"2014 source population literal mismatch at row {number}")
            out[str(mapping.get("source_record_id", mapping.get("candidate_2014_source_record_id")))] = cells
        return out
    finally:
        workbook.close()


def normalize_review_rows(city_obs: pd.DataFrame, city_edges: pd.DataFrame,
                          direct_rows: pd.DataFrame, direct_keys: pd.DataFrame) -> pd.DataFrame:
    """Create a shared reviewed-row vector while preserving original review columns as JSON."""
    _need(len(city_obs) == 3 and len(city_edges) == 3, "expected exactly three reviewed city rows and edges")
    _need(len(direct_rows) == 19 and len(direct_keys) == 19, "expected exactly nineteen direct reviewed rows and keys")
    _need(not city_obs.observation_id.duplicated().any() and not city_edges.from_source_record_id.duplicated().any(),
          "duplicate city source row or identity source")
    _need(not direct_rows.candidate_2014_source_record_id.duplicated().any() and
          not direct_keys.candidate_2014_source_record_id.duplicated().any(), "duplicate direct19 source row")
    _need(set(city_obs.observation_id) == set(city_edges.from_source_record_id), "city review rows and edges do not join exactly")
    _need(set(direct_rows.candidate_2014_source_record_id) == set(direct_keys.candidate_2014_source_record_id),
          "direct19 rows and eligible keys do not join exactly")
    city_edge_map = city_edges.set_index("from_source_record_id")
    city_records = []
    for row in city_obs.to_dict("records"):
        edge = city_edge_map.loc[row["observation_id"]].to_dict()
        _need(edge["to_source_record_id"] == row["current_2021_source_record_id"], "city edge target differs from source review")
        city_records.append({
            "source_record_id": row["observation_id"], "current_source_record_id": row["current_2021_source_record_id"],
            "year": 2014, "place_name": row["city"], "source_caption": row["raw_source_caption"],
            "source_type": row["source_type"], "source_region": row["source_region"],
            "source_row": row["source_row_1based"], "source_sheet": row["source_sheet"],
            "source_population_raw": row["source_2014_population_raw"],
            "source_population_value": row["source_2014_population_value"],
            "source_population_status": row["population_claim_status"],
            "event_scope_review": edge.get("event_scope_review", ""),
            "identity_review": edge.get("identity_review", ""),
            "current_native_OKTMO_literal": row["current_2021_raw_native_OKTMO_literal"],
            "current_2021_population_context": row["current_2021_raw_population"],
            "review_json": json.dumps({"observation": row, "edge": edge}, ensure_ascii=False, separators=(",", ":")),
            "row_family": "three_city_event_scoped",
        })
    direct_key_map = direct_keys.set_index("candidate_2014_source_record_id")
    direct_records = []
    for row in direct_rows.to_dict("records"):
        sid = row["candidate_2014_source_record_id"]
        key = direct_key_map.loc[sid].to_dict()
        _need(key["current_source_record_id"] == row["current_source_record_id"], "direct19 key points to wrong current source row")
        _need(key["review_status"] == "eligible_identity_corridor_only" and str(key["applied"]).lower() == "false",
              "direct19 row is not in the independently eligible unapplied set")
        direct_records.append({
            "source_record_id": sid, "current_source_record_id": row["current_source_record_id"],
            "year": 2014, "place_name": row["current_name"], "source_caption": row["source_caption_raw"],
            "source_type": row["source_type"], "source_region": row["current_region"],
            "source_row": row["source_row_1based"], "source_sheet": row["source_sheet"],
            "source_population_raw": row["source_population_raw"],
            "source_population_value": row["source_population_value"],
            "source_population_status": row["source_population_status"],
            "event_scope_review": "no event/successor scope exception reported in independently reviewed direct19 set",
            "identity_review": key["review_status"],
            "current_native_OKTMO_literal": row["raw_native_OKTMO_literal"],
            "current_2021_population_context": row["raw_population"],
            "review_json": json.dumps({"row": row, "eligible_key": key}, ensure_ascii=False, separators=(",", ":")),
            "row_family": "direct19",
        })
    combined = pd.DataFrame(city_records + direct_records)
    _need(len(combined) == 22 and not combined.source_record_id.duplicated().any(), "combined reviewed source rows collide")
    _need(not combined.current_source_record_id.duplicated().any(), "two historical rows share one current target")
    return combined


def validate_current_carriers(
    reviewed: pd.DataFrame,
    current_rows: pd.DataFrame,
    current_points: pd.DataFrame,
    point_origin_hashes: dict[str, str] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Require exact target join, matching physical type/region/code, and accepted point provenance."""
    _need(not current_rows.source_record_id.duplicated().any(), "current census carrier rows duplicate source IDs")
    _need(not current_points.target_source_record_id.duplicated().any(), "accepted point ledger has ambiguous current carriers")
    rowmap = current_rows.set_index("source_record_id", drop=False)
    pointmap = current_points.set_index("target_source_record_id", drop=False)
    _need(set(reviewed.current_source_record_id) <= set(rowmap.index), "reviewed row targets missing from current selected census rows")
    _need(set(reviewed.current_source_record_id) <= set(pointmap.index), "reviewed row targets missing accepted current point carriers")
    for record in reviewed.itertuples(index=False):
        row = rowmap.loc[record.current_source_record_id]
        point = pointmap.loc[record.current_source_record_id]
        _need(int(row.observation_year) == 2021 and str(row.record_type) == "census" and
              str(row.entity_category) == "settlement" and bool(str(row.entity_id)),
              f"target is not an identified current census settlement: {record.current_source_record_id}")
        _need(str(row.settlement_type) == str(record.source_type),
              f"source type does not match current physical settlement type: {record.source_record_id}")
        _need(str(row.region_raw) == str(record.source_region),
              f"source region does not match current row region: {record.source_record_id}")
        _need(str(row.source_native_id) == str(record.current_native_OKTMO_literal),
              f"reviewed current native code differs from pinned current row: {record.current_source_record_id}")
        _need(str(int(row.population_value)) == str(int(float(record.current_2021_population_context))),
              f"reviewed current population context differs from pinned census row: {record.current_source_record_id}")
        _need(str(point.coordinate_admission_status) in _ACCEPTED_POINT_STATUSES,
              f"current point carrier not accepted: {record.current_source_record_id}")
        _need(math.isfinite(float(point.latitude)) and math.isfinite(float(point.longitude)),
              f"non-finite accepted point: {record.current_source_record_id}")
        review = json.loads(record.review_json)
        if record.row_family == "three_city_event_scoped":
            expected_lat = float(review["observation"]["accepted_carrier_latitude"])
            expected_lon = float(review["observation"]["accepted_carrier_longitude"])
        else:
            expected_lat = float(review["row"]["accepted_current_latitude"])
            expected_lon = float(review["row"]["accepted_current_longitude"])
        _need(math.isclose(float(point.latitude), expected_lat, rel_tol=0, abs_tol=1e-10) and
              math.isclose(float(point.longitude), expected_lon, rel_tol=0, abs_tol=1e-10),
              f"reviewed coordinate does not match the actual current carrier: {record.current_source_record_id}")
        _need(str(point.point_origin_file).strip() and str(point.point_origin_sha256).strip() and str(point.point_origin_locator).strip(),
              f"accepted carrier is missing actual raw point-origin fields: {record.current_source_record_id}")
        if point_origin_hashes is not None:
            expected = point_origin_hashes.get(str(point.point_origin_file))
            _need(expected is not None and expected == str(point.point_origin_sha256),
                  f"accepted point source hash is not verified: {record.current_source_record_id}")
    return rowmap, pointmap


def validate_raw_current_rows(reviewed: pd.DataFrame, raw_rows: pd.DataFrame) -> None:
    """Reopen exact 2021 publisher rows; keep raw publisher point distinct from accepted carrier point."""
    _need(not raw_rows.row_number.duplicated().any(), "raw current row lookup is ambiguous")
    raw = raw_rows.set_index("row_number", drop=False)
    for item in reviewed.itertuples(index=False):
        row_number = int(str(item.current_source_record_id).rsplit(":", 1)[1])
        _need(row_number in raw.index, f"current raw row absent: {item.current_source_record_id}")
        actual = raw.loc[row_number]
        review = json.loads(item.review_json)
        if item.row_family == "three_city_event_scoped":
            witness = review["observation"]
            source_name = witness["current_2021_raw_source_name"]
            level = witness["current_2021_raw_object_level"]
            pop = witness["current_2021_raw_population"]
        else:
            witness = review["row"]
            source_name = witness["raw_object_name"]
            level = witness["raw_object_level"]
            pop = witness["raw_population"]
        _need(str(actual.object_level) == str(level), f"current raw object level differs: {item.current_source_record_id}")
        _need(str(actual.object_name) == str(source_name), f"current raw object name differs: {item.current_source_record_id}")
        _need(str(actual.region) == str(item.source_region), f"current raw region differs: {item.current_source_record_id}")
        _need(str(actual.oktmo) == str(item.current_native_OKTMO_literal), f"current raw OKTMO differs: {item.current_source_record_id}")
        _need(int(float(actual.population)) == int(float(pop)), f"current raw population differs: {item.current_source_record_id}")
        # A separate provider may supply the accepted point. Compare only the raw
        # publisher coordinates to the review's raw-publisher witness, never to
        # the accepted carrier point.
        if item.row_family == "three_city_event_scoped":
            raw_lat, raw_lon = witness.get("raw_current_source_point_latitude"), witness.get("raw_current_source_point_longitude")
        else:
            raw_lat, raw_lon = witness.get("raw_latitude_dadata"), witness.get("raw_longitude_dadata")
        if raw_lat not in (None, "") and raw_lon not in (None, ""):
            _need(math.isclose(float(actual.latitude_dadata), float(raw_lat), abs_tol=1e-10, rel_tol=0) and
                  math.isclose(float(actual.longitude_dadata), float(raw_lon), abs_tol=1e-10, rel_tol=0),
                  f"review raw-publisher point differs from current raw source row: {item.current_source_record_id}")


def create_supplement_rows(reviewed: pd.DataFrame, workbook_cells: dict[str, list[Any]],
                           current_rows: pd.DataFrame, current_points: pd.DataFrame,
                           workbook_path: str | Path = "PINNED_OFFICIAL_ROSSTAT_2014_WORKBOOK") -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Build 3 city observations, 19 direct place-presence rows, 22 links and 22 point contexts."""
    rowmap = current_rows.set_index("source_record_id", drop=False)
    pointmap = current_points.set_index("target_source_record_id", drop=False)
    population_rows: list[dict[str, Any]] = []
    presence_rows: list[dict[str, Any]] = []
    edges: list[dict[str, Any]] = []
    point_uses: list[dict[str, Any]] = []
    for item in reviewed.itertuples(index=False):
        current_id = str(item.current_source_record_id)
        current = rowmap.loc[current_id]
        point = pointmap.loc[current_id]
        cells = workbook_cells[str(item.source_record_id)]
        review = json.loads(item.review_json)
        witness = review["observation"] if item.row_family == "three_city_event_scoped" else review["row"]
        total = literal_integer(cells[1])
        if total is not None:
            _need(str(item.source_population_status) in {"literal_official_workbook_cell", "literal_integer_in_official_2014_source"},
                  f"numeric workbook cell is not marked observed: {item.source_record_id}")
            _need(str(literal_integer(item.source_population_value)) == str(total), f"reviewed population differs from raw workbook cell: {item.source_record_id}")
            male, female = literal_integer(cells[2]), literal_integer(cells[3])
            _need(male is None or female is None or male + female == total,
                  f"2014 gender totals do not sum to population: {item.source_record_id}")
            record_type = ("scoped_official_observation" if item.row_family == "three_city_event_scoped"
                           else "scoped_official_place_presence")
            payload = {
                "observation_id": (f"official2014:{item.source_record_id}" if item.row_family == "three_city_event_scoped"
                                   else f"official2014presence:{item.source_record_id}"),
                "record_type": record_type,
                "population_value": total,
                "population_raw": str(cells[1]),
                "population_male": male,
                "population_female": female,
                "population_male_raw": None if cells[2] is None else str(cells[2]),
                "population_female_raw": None if cells[3] is None else str(cells[3]),
                "population_value_quality": "literal_official_rosstat_2014_row_replayed",
            }
            (population_rows if item.row_family == "three_city_event_scoped" else presence_rows).append(payload)
        else:
            _need(item.row_family == "direct19", f"a reviewed city row cannot have a missing population cell: {item.source_record_id}")
            _need(str(cells[1]).strip() in _DASHES, f"non-numeric 2014 population cell is not a literal dash: {item.source_record_id}")
            _need(str(item.source_population_status) == "missing_source_dash_no_observation",
                  f"dash row is incorrectly marked as a population observation: {item.source_record_id}")
            _need(str(item.source_population_value).strip() == "", f"dash row has a fabricated numeric population: {item.source_record_id}")
            record_type = "scoped_official_place_presence"
            payload = {
                "observation_id": f"official2014presence:{item.source_record_id}",
                "record_type": record_type,
                "population_value": None,
                "population_raw": str(cells[1]),
                "population_male": None,
                "population_female": None,
                "population_male_raw": None if cells[2] is None else str(cells[2]),
                "population_female_raw": None if cells[3] is None else str(cells[3]),
                "population_value_quality": "missing_source_dash_no_population_observation",
            }
            presence_rows.append(payload)
        common = {
            **payload,
            "entity_id": current.entity_id,
            "associated_census_entity_id": current.entity_id,
            "current_place_entity_id": current.entity_id,
            "association_status": "reviewed_scoped_same_place_to_2021",
            "spatial_identity_status": "reviewed_scoped_same_place_to_2021",
            "identity_quality": "reviewed_scoped_same_place_to_2021",
            "observation_year": 2014,
            "observation_time_precision": "year_only_source_year; exact census date not claimed",
            "reference_date": None,
            "reference_date_basis": "source_publication_year_only",
            "actual_census_date_claimed": False,
            "source_record_id": item.source_record_id,
            "source_publication_row_id": item.source_record_id,
            "source_name_raw": item.source_caption,
            "settlement_name": current.settlement_name,
            "settlement_type": item.source_type,
            "region_raw": current.region_raw,
            "source_region_raw": item.source_region,
            "source_admin_heading_raw": (witness.get("source_city_group_header_reconstructed_from_workbook", "")
                                          if item.row_family == "three_city_event_scoped"
                                          else witness.get("source_district_heading_raw", "")),
            "source_municipality_heading_raw": witness.get("source_municipality_heading_raw", ""),
            "source_review_evidence_json": item.review_json,
            "population_raw_cells_json": json.dumps(cells, ensure_ascii=False, separators=(",", ":")),
            "population_scope": "one literal 2014 source row; 2021 boundary equivalence unknown",
            "source_population_status": item.source_population_status,
            "population_scope_comparability_asserted": False,
            "population_comparability_to_2021": False,
            "administrative_boundary_comparability_to_2021": "unknown_not_asserted",
            "boundary_comparability_asserted": False,
            "source_sheet": item.source_sheet,
            "source_row": int(item.source_row),
            "source_row_caption_literal": item.source_caption,
            "source_row_type_literal": item.source_type,
            "source_path": str(workbook_path),
            "source_workbook_path": str(workbook_path),
            "source_sha256": WORKBOOK_SHA256,
            "source_item_qid": "Q127387785",
            "source_item_title": "Russian Census 2014 Crimea publication; scoped source item only",
            "current_2021_source_record_id": current_id,
            "current_2021_source_native_id": current.source_native_id,
            "current_2021_population_context": current.population_value,
            # These are the accepted current carrier's actual point values and
            # graph7 origin, explicitly reused as modern context for the 2014
            # physical-place row; they are not 2014 measurements.
            "latitude": float(point.latitude),
            "longitude": float(point.longitude),
            "coordinate_quality": "reviewed_modern_representative_point_retrospective",
            "coordinate_admission_status": "reviewed_rule_accepted",
            "coordinate_temporal_basis": "modern_rep_point_spatial_continuity_inference_dateunknown",
            "coordinate_measurement_date_unknown": True,
            "direct_historical_coordinate_measurement": False,
            "coordinate_source": point.coordinate_source,
            "coordinate_source_record_id": None,
            "coordinate_carrier_target_source_record_id": current_id,
            "coordinate_provider": point.coordinate_provider,
            "coordinate_provider_id": None,
            "current_carrier_coordinate_provider_id_context_only": point.coordinate_provider_id,
            "coordinate_provenance": point.coordinate_provenance,
            "coordinate_admission_rule": "reviewed 2014-to-2021 physical-place continuity; accepted 2021 representative point reused with unknown historical measurement date",
            "point_origin_file": point.point_origin_file,
            "point_origin_sha256": point.point_origin_sha256,
            "point_origin_locator": point.point_origin_locator,
            "point_origin_kind": point.point_origin_kind,
            "point_origin_basis": "actual current carrier origin from the pinned graph7 accepted point ledger",
            "historical_provider_binding_asserted": False,
            "population_scope_comparability_asserted": False,
            "coordinate_population_boundary_comparability_asserted": False,
            "native_2014_okato_raw": None,
            "native_2014_oktmo_raw": None,
            "native_2014_identifier_binding_asserted": False,
            "historical_identity_admitted": True,
            "strict_Russian_2002_2010_2021_chain_eligible": False,
            "census_full_chain": False,
            "census_2002_status": "outside_scope_of_this_scoped_2014_layer",
            "census_2010_status": "outside_scope_of_this_scoped_2014_layer",
            "census_2021_status": "observed",
            "historical_coordinate_asserted": False,
            "event_scope_review": item.event_scope_review,
            "review_row_family": item.row_family,
        }
        (population_rows if item.row_family == "three_city_event_scoped" else presence_rows)[-1].update(common)
        edges.append({
            "from_source_record_id": item.source_record_id,
            "from_year": 2014,
            "to_source_record_id": current_id,
            "to_year": 2021,
            "relation": "same_physical_place_continuity_candidate",
            "decision_status": "accepted_scoped_identity_observation_only",
            "review_status": item.identity_review,
            "current_entity_id": current.entity_id,
            "source_population_observation_id": payload["observation_id"],
            "population_row_kind": record_type,
            "source_native_2014_identifier": None,
            "current_native_2021_identifier_literal": current.source_native_id,
            "current_native_code_year": 2021,
            "event_scope_review": item.event_scope_review,
            "boundary_comparability_asserted": False,
            "legal_effective_date": None,
            "strict_Russian_2002_2010_2021_chain_eligible": False,
            "canonical_2002_2010_2021_graph_modified": False,
        })
        point_uses.append({
            "target_source_record_id": item.source_record_id,
            "target_year": 2014,
            "latitude": float(point.latitude),
            "longitude": float(point.longitude),
            "coordinate_admission_status": "reviewed_rule_accepted",
            "point_use_application_status": "accepted_scoped_retrospective_point_use",
            "coordinate_temporal_basis": "modern_rep_point_spatial_continuity_inference_dateunknown",
            "coordinate_measurement_date_unknown": True,
            "direct_historical_coordinate_measurement": False,
            "historical_coordinate_asserted": False,
            "boundary_comparability_asserted": False,
            "population_scope_comparability_asserted": False,
            "current_carrier_target_source_record_id": current_id,
            "coordinate_source": point.coordinate_source,
            "coordinate_provider": point.coordinate_provider,
            "coordinate_provider_id_context_only": point.coordinate_provider_id,
            "coordinate_provenance": point.coordinate_provenance,
            "point_origin_file": point.point_origin_file,
            "point_origin_sha256": point.point_origin_sha256,
            "point_origin_locator": point.point_origin_locator,
            "point_origin_kind": point.point_origin_kind,
            "point_origin_basis": "actual current carrier origin copied from pinned graph7 accepted point ledger",
            "current_point_ledger_sha256": POINT_LEDGER_SHA256,
        })

    obs = pd.DataFrame(population_rows)
    presence = pd.DataFrame(presence_rows)
    edge_df = pd.DataFrame(edges)
    point_df = pd.DataFrame(point_uses)
    _need(len(obs) == 3 and len(presence) == 19, f"expected 3 city observations and 19 direct presence rows; got {len(obs)} and {len(presence)}")
    _need(len(edge_df) == len(point_df) == 22, "expected 22 reviewed identity links and point uses")
    _need(obs.population_value.notna().all() and int(presence.population_value.notna().sum()) == 8 and
          int(presence.population_value.isna().sum()) == 11,
          "direct presence rows must retain 8 literal counts and 11 missing dashes")
    _need(not obs.observation_id.duplicated().any() and not presence.observation_id.duplicated().any(), "duplicate scoped observation IDs")
    _need(not set(obs.observation_id) & set(presence.observation_id), "observation and presence ID namespaces collide")
    _need(not edge_df.from_source_record_id.duplicated().any() and not edge_df.to_source_record_id.duplicated().any(), "identity links are not one-to-one")
    _need(not point_df.target_source_record_id.duplicated().any(), "duplicate scoped point use")
    return obs, presence, edge_df, point_df


def ensure_no_observation_collisions(base_long_path: str | Path, new_ids: list[str], source_row_ids: list[str] | None = None) -> None:
    con = duckdb.connect(":memory:")
    try:
        n = con.execute("SELECT count(*) FROM read_parquet(?) WHERE observation_id IN (SELECT unnest(?))",
                        [str(base_long_path), new_ids]).fetchone()[0]
        _need(n == 0, f"{n} supplemental observation IDs already exist in the frozen 994-row application")
        if source_row_ids:
            columns = {str(r[0]) for r in con.execute("DESCRIBE SELECT * FROM read_parquet(?)", [str(base_long_path)]).fetchall()}
            predicates = []
            params: list[Any] = [str(base_long_path)]
            for column in ("source_record_id", "source_publication_row_id"):
                if column in columns:
                    predicates.append(f'"{column}" IN (SELECT unnest(?))')
                    params.append(source_row_ids)
            if predicates:
                source_collisions = con.execute(
                    "SELECT count(*) FROM read_parquet(?) WHERE " + " OR ".join(predicates), params
                ).fetchone()[0]
                _need(source_collisions == 0,
                      f"{source_collisions} supplemental publisher row IDs already exist in the frozen application")
    finally:
        con.close()


def typed_staged_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Apply stable nullable types to the fields consumed by the long-table integrator."""
    out = frame.copy()
    for column in ("observation_year", "source_row", "population_value", "population_male", "population_female",
                   "current_2021_population_context"):
        if column in out:
            out[column] = pd.to_numeric(out[column], errors="coerce").astype("Int64")
    if "observation_year" in out:
        out["observation_year"] = out["observation_year"].astype("Int16")
    bool_columns = [column for column in out.columns if column.endswith("_asserted") or column in {
        "actual_census_date_claimed", "direct_historical_coordinate_measurement",
        "coordinate_measurement_date_unknown", "historical_identity_admitted",
        "strict_Russian_2002_2010_2021_chain_eligible", "census_full_chain", "historical_coordinate_asserted",
    }]
    for column in bool_columns:
        out[column] = out[column].astype("boolean")
    # Keep source literals, identifiers, paths, hashes, and raw JSON as strings;
    # pandas nullable extension types above preserve NULL rather than inventing 0.
    return out


def _main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, help="manifest JSON with exact source and baseline pins")
    parser.add_argument("--write-staged", action="store_true", help="write candidate supplement CSV/Parquet/receipt; does not append to the accepted long table")
    args = parser.parse_args()
    manifest_path = Path(args.manifest)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    paths = verify_manifest(manifest)
    city_obs, city_edges = _read_csv(paths["city_observations"]), _read_csv(paths["city_edges"])
    direct_rows, direct_keys = _read_csv(paths["direct19_rows"]), _read_csv(paths["direct19_keys"])
    reviewed = normalize_review_rows(city_obs, city_edges, direct_rows, direct_keys)
    cells = replay_source_cells(reviewed, paths["source_workbook"], "source_row", "source_caption", "source_population_raw")

    target_ids = sorted(set(reviewed.current_source_record_id))
    con = duckdb.connect(":memory:")
    con.execute("SET memory_limit='512MB'")
    con.execute("SET threads=1")
    current = con.execute("""SELECT 'census' AS record_type, 2021 AS observation_year,
        source_record_id,entity_id,'settlement' AS entity_category,settlement_name,settlement_type,
        region_raw,population_value,source_native_id
        FROM read_parquet(?) WHERE record_type='census' AND observation_year=2021 AND source_record_id IN (SELECT unnest(?))""",
        [str(paths["base_long"]), target_ids]).fetchdf()
    selected = con.execute("""SELECT census_year,source_record_id,settlement_name,settlement_type,
        region_raw,population,source_native_id FROM read_parquet(?)
        WHERE census_year=2021 AND source_record_id IN (SELECT unnest(?))""",
        [str(paths["selected_population"]), target_ids]).fetchdf()
    _need(len(current) == len(selected) == 22, "current 2021 carrier source rows missing or duplicated")
    for field in ["source_record_id", "settlement_name", "settlement_type", "region_raw", "source_native_id"]:
        _need(current[field].astype(str).sort_values().tolist() == selected[field].astype(str).sort_values().tolist(),
              f"7th long baseline and selected source disagree in {field}")
    _need(current.population_value.astype(str).sort_values().tolist() == selected.population.astype(str).sort_values().tolist(),
          "7th long baseline and selected 2021 populations differ")
    raw_rownums = [int(x.rsplit(":", 1)[1]) for x in target_ids]
    raw_sql = """WITH source AS (SELECT row_number() OVER() row_number,object_level,object_name,oktmo,
        region,population,latitude_dadata,longitude_dadata FROM read_parquet(?))
        SELECT * FROM source WHERE row_number IN (SELECT unnest(?))"""
    raw_current = con.execute(raw_sql, [str(paths["current_raw"]), raw_rownums]).fetchdf()
    _need(len(raw_current) == 22, "current raw publisher rows missing or duplicated")
    validate_raw_current_rows(reviewed, raw_current)
    point_cols = ["target_source_record_id", "latitude", "longitude", "coordinate_source", "coordinate_provider",
                  "coordinate_provider_id", "coordinate_provenance", "coordinate_admission_status",
                  "point_origin_file", "point_origin_sha256", "point_origin_locator", "point_origin_kind"]
    point_sql = ",".join('"'+x+'"' for x in point_cols)
    points = con.execute(f"SELECT {point_sql} FROM read_parquet(?) WHERE regexp_replace(target_year,'\\.0$','')='2021' AND target_source_record_id IN (SELECT unnest(?))",
                         [str(paths["point_ledger"]), target_ids]).fetchdf()
    con.close()
    _need(len(points) == 22, "one or more accepted 2021 point carrier rows are missing")
    rowmap, pointmap = validate_current_carriers(reviewed, current, points)

    # Validate each actual raw point origin once. Hashes come from graph7 point_origin_* columns.
    origin_hashes: dict[str, str] = {}
    for row in points.itertuples(index=False):
        origin = str(row.point_origin_file)
        _need(Path(origin).is_file(), f"actual point-origin file is unavailable: {origin}")
        if origin not in origin_hashes:
            origin_hashes[origin] = sha256(origin)
        actual = origin_hashes[origin]
        _need(actual == str(row.point_origin_sha256), f"actual point-origin file hash mismatch: {origin}")
    validate_current_carriers(reviewed, current, points, origin_hashes)
    obs, presence, edges, point_uses = create_supplement_rows(reviewed, cells, current, points, paths["source_workbook"])
    candidate_ids = obs.observation_id.astype(str).tolist() + presence.observation_id.astype(str).tolist()
    ensure_no_observation_collisions(paths["base_long"], candidate_ids, reviewed.source_record_id.astype(str).tolist())
    _need(int(obs.population_value.sum() + presence.population_value.sum()) == 207999,
          "unexpected population sum across 3 city and 8 numeric direct observations")

    summary = {
        "status": "root_approved_scoped_crimea_2014_layer_ready_for_integration",
        "root_decision_reference": "2026-10-04 root approval of the independently reviewed 3-city plus 19-direct-row scoped 2014 layer",
        "candidate_source_rows": 22,
        "city_population_observation_rows": len(obs),
        "direct_place_presence_rows": len(presence),
        "population_value_cells_across_both_layers": int(obs.population_value.notna().sum()+presence.population_value.notna().sum()),
        "presence_only_rows_with_literal_dash": int(presence.population_value.isna().sum()),
        "typed_identity_links": len(edges),
        "point_context_rows": len(point_uses),
        "three_city_population_cells": int(obs.loc[obs.review_row_family.eq("three_city_event_scoped"), "population_value"].sum()),
        "direct19_numeric_population_cells": int(presence.loc[presence.population_value.notna(), "population_value"].sum()),
        "direct19_dash_rows": int(presence.population_value.isna().sum()),
        "all_point_origins_replayed": True,
        "historical_coordinates_asserted": False,
        "population_boundary_comparability_asserted": False,
        "strict_Russian_2002_2010_2021_chain_delta": 0,
        "canonical_graph_or_baseline_mutated": False,
        "applied_to_long_table": False,
        "Russian_census_chain_gain_claimed": False,
    }
    if args.write_staged:
        out = Path(manifest["staged_output_dir"])
        _need(not out.exists(), f"refusing to overwrite staged output directory: {out}")
        out.mkdir(parents=True)
        typed_obs = typed_staged_frame(obs)
        typed_presence = typed_staged_frame(presence)
        typed_edges = typed_staged_frame(edges)
        typed_points = typed_staged_frame(point_uses)
        all_observations = typed_staged_frame(pd.concat([typed_obs, typed_presence], ignore_index=True, sort=False))
        files = {
            "city_observations": typed_obs,
            "place_presence": typed_presence,
            "scoped_observations": all_observations,
            "identity_links": typed_edges,
            "retrospective_point_context": typed_points,
        }
        pins = {}
        for name, frame in files.items():
            csv_path = out / f"{name}.csv"
            parquet_path = out / f"{name}.parquet"
            frame.to_csv(csv_path, index=False, quoting=csv.QUOTE_MINIMAL)
            frame.to_parquet(parquet_path, index=False)
            pins[f"{name}_csv"] = {"path": str(csv_path), "sha256": sha256(csv_path), "bytes": csv_path.stat().st_size}
            pins[f"{name}_parquet"] = {"path": str(parquet_path), "sha256": sha256(parquet_path), "bytes": parquet_path.stat().st_size}
        manifest_copy = out / "frozen_input_manifest.json"
        manifest_copy.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        pins["frozen_input_manifest"] = {"path": str(manifest_copy), "sha256": sha256(manifest_copy), "bytes": manifest_copy.stat().st_size}
        receipt = {**summary, "input_hashes": {k: sha256(v) for k, v in paths.items()}, "outputs": pins,
                   "typed_parquet_schema": {name: {column: str(dtype) for column, dtype in frame.dtypes.items()}
                                            for name, frame in files.items()}}
        (out / "staged_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    _main()
