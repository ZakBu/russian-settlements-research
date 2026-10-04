"""Stage source-backed former-city/district trajectories from official census tables.

The output is an event and auxiliary-population candidate packet only. It does
not add records to selected census layers, identity graph or point ledger.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, re, subprocess, zipfile
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT=Path('/workspace')
P2002=ROOT/'settlements-raw/data/raw/2002_official_tom1/tom-1-chislennost-i-razmeshchenie-naseleniya.pdf'
P2002_XLS=ROOT/'settlements-raw/data/raw/2002_official_tom1/1_TOM_01_04.xls'
P2010_T1=ROOT/'settlements-raw/data/raw/2010_official_tom1/tom-1-chislennost-i-razmeshchenie-naseleniya.pdf'
P2010_T11=ROOT/'settlements-raw/data/raw/2010_official_tom11/pub-11-1-4.pdf'
X2021=ROOT/'settlements-raw/data/raw/rosstat_2021/Tom1_tab-5_VPN-2020.xlsx'
GNZIP=ROOT/'settlements-raw/data/raw/coordinate_candidates/geonames_RU_20260907.zip'
SEL=ROOT/'settlements-delivery/continuation-consolidated-20261003/selected_observations.parquet'
LEGACY=ROOT/'settlements-work/continuation_20261004/R4/temporal_transformation_path_recovery/held_event_candidates.csv'
LEGACY_TOP=ROOT/'settlements-work/continuation_20261004/R4/temporal_transformation_path_recovery/top_mass_event_inventory.csv'

CASES=[
 {'name':'Кайеркан','old_pop':27116,'2010_pop':22338,'2021_pop':21193,'old_source_row':8696,'gn_id':'1504139','gn_name':'Kayyerkan','gn_alias':'Kayyerkan','gn_raw_aliases':'Kajerkam,Kajerkan,Kayerkan,Kayyerkan,Кайеркан','gn_lat':69.37861,'gn_lon':87.74389,'gn_date':'2012-08-04'},
 {'name':'Талнах','old_pop':58654,'2010_pop':47307,'2021_pop':47216,'old_source_row':8697,'gn_id':'1490256','gn_name':'Talnakh','gn_alias':'Talnakh','gn_raw_aliases':'Talnakh,Талнах','gn_lat':69.4865,'gn_lon':88.3972,'gn_date':'2012-01-17'},
]

def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
 return h.hexdigest()

def pdf_text(path):
 return subprocess.run(['pdftotext','-layout',str(path),'-'],check=True,stdout=subprocess.PIPE).stdout.decode('utf-8',errors='replace')

def exact_line(pages, page_i, fragment):
 lines=pages[page_i-1].splitlines()
 matches=[(i+1,x.rstrip()) for i,x in enumerate(lines) if fragment in x]
 if len(matches)!=1: raise ValueError(f'Expected one {fragment!r} on PDF page {page_i}; got {matches}')
 return matches[0]

def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--output',required=True,type=Path); a=ap.parse_args()
 if a.output.exists(): raise FileExistsError(f'Use new immutable output path: {a.output}')
 a.output.mkdir(parents=True)
 t02=pdf_text(P2002).split('\f'); t10=pdf_text(P2010_T1).split('\f'); c10=pdf_text(P2010_T11).split('\f')
 # Reopened publisher lines and structural parent labels, all on exact pages.
 witnesses={
  '2002_parent': {'page':237,'fragment':'г. Норильск с подчиненными его администрации'},
  '2002_kayerkan': {'page':237,'fragment':'г. Кайеркан'},
  '2002_talnakh': {'page':237,'fragment':'г. Талнах'},
  '2010_parent': {'page':178,'fragment':'г. Норильск с подчиненными его администрации'},
  '2010_intracity_header': {'page':178,'fragment':'внутригородские районы'},
  '2010_kayerkan': {'page':178,'fragment':'Кайеркан'},
  '2010_talnakh': {'page':178,'fragment':'Талнах'},
  '2010_norilsk_comparison': {'page':13,'fragment':'Норильск'},
 }
 found={
  '2002_parent':exact_line(t02,237,'г. Норильск с подчиненными его администрации'),
  '2002_kayerkan':exact_line(t02,237,'г. Кайеркан'),
  '2002_talnakh':exact_line(t02,237,'г. Талнах'),
  '2010_parent':exact_line(t10,178,'г. Норильск с подчиненными его администрации'),
  '2010_intracity_header':exact_line(t10,178,'внутригородские районы'),
  '2010_kayerkan':exact_line(t10,178,'Кайеркан'),
  '2010_talnakh':exact_line(t10,178,'Талнах'),
  '2010_norilsk_comparison':exact_line(c10,13,'Норильск'),
 }
 if 'Кайеркан' in '\n'.join(c10) or 'Талнах' in '\n'.join(c10):
  raise ValueError('2010 comparative city table unexpectedly names a former city as a separate row')
 found['2002_parent_total']=(47,'населенными пунктами - городское население          221908     111036     110872        50,0           50,0')
 found['2002_norilsk_city']=(48,'    г. Норильск                                     134832      67316      67516        49,9           50,1')
 found['2002_snezhnogorsk']=(51,'    пгт Снежногорск                                   1306        658        648        50,4           49,6')
 found['2010_parent_total']=(39,'населенными пунктами - городское население            176252      88007      88245        49,9           50,1')
 found['2010_norilsk_city']=(40,'    г. Норильск                                       175365      87570      87795        49,9           50,1')
 found['2010_central_district']=(44,'           Центральный                                105720      52418      53302        49,6           50,4')
 found['2010_snezhnogorsk']=(45,'    пгт Снежногорск рп                                   887        437        450        49,3           50,7')
 # Current 2021 official primary table rows (printed admin grouping, not selected NP layer).
 wb=load_workbook(X2021,read_only=True,data_only=True); ws=wb['таб. 5']
 rowvals={i:[ws.cell(i,j).value for j in range(1,7)] for i in range(22432,22438)}
 expected={22432:'Городской округ город Норильск - городское население',22433:'г. Норильск',22434:'внутригородские районы',22435:'район Кайеркан',22436:'район Талнах',22437:'Центральный район'}
 for i,n in expected.items():
  if str(rowvals[i][0]).strip()!=n: raise ValueError(f'2021 source row changed at {i}: {rowvals[i]}')
 rowvals[22438]=[ws.cell(22438,j).value for j in range(1,7)]
 if str(rowvals[22438][0]).strip()!='пгт Снежногорск гп': raise ValueError('2021 Norilsk subordinate row changed')
 # Exact 2002 selected IDs/pop/type and absence of independent 2010/2021 selected settlement rows.
 s=pd.read_parquet(SEL,columns=['source_record_id','census_year','settlement_name','settlement_type','region_raw','population','source_file','source_sheet','source_row','source_sha256','source_locator','source_native_id','population_scope','is_additive_settlement_record'])
 ev=[]; aux=[]
 # GeoNames raw physical features only corroborate historical-place point candidates; they do not establish a census-date point.
 with zipfile.ZipFile(GNZIP) as z:
  lines=z.read('RU.txt').decode('utf-8').splitlines()
  gnrows={}; gn_alias_options={}
  wanted={c['name'].casefold() for c in CASES}
  for line in lines:
   f=line.split('\t')
   aliases=set(x.strip().casefold() for x in f[3].split(',')) if len(f)>3 else set()
   if len(f)>10 and f[6]=='P' and f[7]=='PPL' and f[10]=='91':
    for name in wanted:
     if name in aliases: gn_alias_options.setdefault(name,[]).append(f)
   if f[0] in {c['gn_id'] for c in CASES}: gnrows[f[0]]=f
 for c in CASES:
  old_id=f"2002:1_TOM_01_04.xls:0:{c['old_source_row']}"
  old=s[s.source_record_id.eq(old_id)]
  if len(old)!=1 or int(old.iloc[0].population)!=c['old_pop'] or old.iloc[0].settlement_type!='город': raise ValueError(f'2002 selected row mismatch for {c["name"]}')
  if len(s[(s.census_year.isin([2010,2021])) & s.settlement_name.eq(c['name'])]): raise ValueError(f'Unexpected independent selected NP row for {c["name"]}')
  gn=gnrows.get(c['gn_id'])
  options=gn_alias_options.get(c['name'].casefold(),[])
  if not gn or gn[1]!=c['gn_name'] or gn[3]!=c['gn_raw_aliases'] or gn[6]!='P' or gn[7]!='PPL' or gn[10]!='91' or not options or all(x[0]!=c['gn_id'] for x in options): raise ValueError(f'GeoNames witness mismatch for {c["name"]}: {gn}')
  # Page/line proof pins. The corresponding 2010/2021 values are subordinate
  # table observations and remain outside the selected settlement layer.
  oldline=found['2002_'+('kayerkan' if c['name']=='Кайеркан' else 'talnakh')][1]
  line10=found['2010_'+('kayerkan' if c['name']=='Кайеркан' else 'talnakh')][1]
  gny=gn
  ev.append({
   'candidate_event_id':'norilsk_city_to_intracity_district_'+('kayerkan' if c['name']=='Кайеркан' else 'talnakh'),
   'place_name':c['name'],'candidate_relation':'physical_place_continuity_with_city_to_district_scope_change',
   'event_type':'separate_city_row_reclassified_as_norilsk_intracity_district',
   'from_source_record_id_2002':old_id,'from_year':2002,'from_type':'город','from_population':c['old_pop'],
   'from_selected_source_file':old.iloc[0].source_file,'from_selected_sheet':old.iloc[0].source_sheet,'from_selected_row':int(old.iloc[0].source_row),
   'from_publisher_structural_row':oldline,'from_publisher_pdf':'2002 Tom 1, Table 5, PDF page 237',
   'to_auxiliary_source_2010':'2010 Tom 1, Table 5, PDF page 178', 'to_auxiliary_type_2010':'intracity district under Norilsk',
   'to_auxiliary_population_2010':c['2010_pop'],'to_publisher_structural_row_2010':line10,
   'current_auxiliary_source_2021':'2021 Tom 1, Table 5, worksheet таб. 5', 'to_auxiliary_type_2021':'район (under explicit внутригородские районы header)',
   'to_auxiliary_population_2021':c['2021_pop'],'to_publisher_row_2021':22435 if c['name']=='Кайеркан' else 22436,
   '2021_independent_NP_status':'not_applicable_as_a_separate_settlement_record; table gives intracity district observation only',
   'selected_Norilsk_city_population_2021':174453,'parent_record_id_2021':'2021:data_allsettlements_anon_156_v20251217.parquet:parquet:49239',
   'official_Norilsk_urban_group_population_2021':175237,'norilsk_subordinate_Snezhnogorsk_2021':784,
   'city_population_equals_intracity_district_sum_2021':True,
   'district_population_must_not_be_added_to_selected_city_population':True,
   'event_interval':'after 2002 census and by 2010 census; exact legal/statistical effective date not found in these tables',
   'event_date_claimed':False,'same_place_edge_to_Norilsk_claimed':False,'population_boundary_comparability_claimed':False,
   'old_physical_point_candidate_origin':'GeoNames RU snapshot, feature P/PPL, exact locality alias; not an exact census-date measurement',
   'old_point_geonameid':gny[0],'old_point_primary_name':gny[1],'old_point_feature_class':gny[6],'old_point_feature_code':gny[7],
   'old_point_aliases':gny[3],'old_point_latitude':float(gny[4]) if len(options)==1 else None,
   'old_point_longitude':float(gny[5]) if len(options)==1 else None,'geonames_source_modification_date':gny[18],
   'same_ADM1_exact_alias_PPL_candidate_count':len(options),
   'old_point_options_json':json.dumps([{'geonameid':x[0],'name':x[1],'aliases':x[3],'latitude':float(x[4]),'longitude':float(x[5]),'feature_class':x[6],'feature_code':x[7],'admin1_code':x[10],'population_field_not_used_for_choice':x[14],'modification_date':x[18]} for x in options],ensure_ascii=False),
   'old_point_admission_status':'point_choice_hold_multiple_exact_alias_PPL' if len(options)>1 else 'candidate_only_requires_independent_review','provider_identifier_binding_asserted':False,
   'candidate_only':True,'admission_status':'no_event_or_identity_admission'
  })
  for year,pop,table,loc,typ in [
   (2002,c['old_pop'],'2002 Tom 1 Table 5','PDF page 237; lines 46-50','city settlement row'),
   (2010,c['2010_pop'],'2010 Tom 1 Table 5','PDF page 178; lines 38-43','Norilsk intracity district row'),
   (2021,c['2021_pop'],'2021 Tom 1 Table 5','worksheet таб. 5; row '+str(22435 if c['name']=='Кайеркан' else 22436),'Norilsk intracity district row')]:
   aux.append({'place_name':c['name'],'year':year,'reported_population':pop,'source_table':table,'source_locator':loc,'published_label_type':typ,'population_layer':'auxiliary event/status observation; not appended to selected settlement population layer','included_in_current_selected_totals':year==2002,'row_subject':'named historical place / current intracity district; not the Norilsk parent'})
 # The 2010 comparative city table contains Norilsk proper but no K/T rows; its absence alone is not interpreted as an event.
 # Other legacy event pointers are kept held unless official sources expose a direct structural event witness.
 holds=[]
 for _,r in pd.read_csv(LEGACY).iterrows():
  name=r.get('place_name') or r.get('name') or r.get('event_id')
  eid=str(r.get('event_id',''))
  if 'talnakh' in eid or 'kayerkan' in eid: continue
  reason='No direct official comparative row/footnote linking the named former unit to the proposed successor was located in cached 2010/2021 tables; absence from a city list is not an event witness.'
  holds.append({'event_id':eid,'legacy_event_type':r.get('event_type'),'legacy_claimed_year':r.get('legacy_asserted_year'), 'legacy_name_or_id':name,'review_status':'held_unverified_event','hold_reason':reason,'candidate_only':True})
 for _,r in pd.read_csv(LEGACY_TOP).iterrows():
  if 'zheleznodorozhny' not in str(r.get('event_id','')): continue
  holds.append({'event_id':r.get('event_id'),'legacy_event_type':r.get('event_type'),'legacy_claimed_year':r.get('legacy_asserted_year'),'legacy_name_or_id':'Железнодорожный','review_status':'held_unverified_event',
   'hold_reason':'Official 2010 comparative table still lists this as a separate city; cached 2021 tables do not contain a direct inclusion/merger footnote for Balashikha. A cached Wikidata claim/legacy date is not promoted to an official event here.','candidate_only':True})
 # Source output tables, hashes, receipt.
 pd.DataFrame(ev).to_csv(a.output/'typed_transformation_candidates.csv',index=False)
 pd.DataFrame(aux).to_csv(a.output/'auxiliary_official_population_observations.csv',index=False)
 pd.DataFrame(holds).to_csv(a.output/'held_legacy_event_candidates.csv',index=False)
 sourcefiles=[P2002,P2002_XLS,P2010_T1,P2010_T11,X2021,GNZIP,SEL,LEGACY,LEGACY_TOP,Path(__file__)]
 receipt={
  'status':'candidate_only_official_source_backed_status_trajectories_no_admissions',
  'scope':'Two old Norilsk subordinate city settlements whose official census rows shift to intracity district rows by 2010 and remain districts in 2021.',
  'candidate_event_count':len(ev),'supported_place_names':[x['place_name'] for x in ev],
  'recorded_populations_by_place':{c['name']:{'2002_city':c['old_pop'],'2010_intracity_district':c['2010_pop'],'2021_intracity_district':c['2021_pop']} for c in CASES},
  'additional_hierarchy_controls':{
   '2002':{'Norilsk_administration_group':221908,'Norilsk_city_proper':134832,'Kayerkan_city':27116,'Talnakh_city':58654,'Snezhnogorsk_subordinate_pgt':1306,'rule':'component hierarchy as printed; do not sum children into parent a second time'},
   '2010':{'Norilsk_administration_group':176252,'Norilsk_city_proper':175365,'Kayerkan_intracity_district':22338,'Talnakh_intracity_district':47307,'Centralny_intracity_district':105720,'Snezhnogorsk_subordinate_pgt':887,'rule':'intracity districts sum to city proper; do not add to parent/city population'},
   '2021':{'Norilsk_urban_group':175237,'Norilsk_city_proper':174453,'Kayerkan_intracity_district':21193,'Talnakh_intracity_district':47216,'Centralny_intracity_district':106044,'Snezhnogorsk_subordinate_pgt':784,'rule':'intracity districts sum to city proper; city proper plus Snezhnogorsk equals urban group; never add districts to current city row'}},
  'additional_current_parent_context':{'Norilsk_2010_city_proper':175365,'Norilsk_2021_city_proper':174453,'district_counts_not_added_to_parent':True},
  'not_applicable_semantics':'2021 independent settlement row is not applicable for the former-city entity in the selected layer; the official auxiliary table still reports district populations, which remain separate and are not summed with parent.',
  'date_semantics':'Status/scope transition bounded after 2002 and by 2010 only; no exact legal effective date claimed.',
  'official_source_witnesses':{k:{'pdf_page_1based':v['page'],'line_1based':found[k][0],'text':found[k][1]} for k,v in witnesses.items()},
  '2010_city_table_semantics':'Table 1.4 has numeric 2002 and 2010 columns for city rows. Kayerkan/Talnakh do not occur as city rows there; no dash is used for them, so absence is not decoded as zero or as an event by itself.',
  'hard_nonclaims':['No same-place edge to Norilsk or parent population transfer.','No exact legal event date.','No population boundary comparability assertion.','No historical point admission; GeoNames locality point is a candidate witness only.','No auxiliary 2010/2021 district population is added to the selected settlement layer.'],
  'inputs':{str(p):{'sha256':sha(p),'bytes':p.stat().st_size} for p in sourcefiles},
  'outputs':{f.name:sha(f) for f in a.output.iterdir() if f.is_file()},
  'counts':{'direct_official_trajectory_candidates':len(ev),'held_legacy_event_candidates':len(holds),'auxiliary_population_rows':len(aux)}
 }
 (a.output/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'output':str(a.output),'events':len(ev),'held':len(holds),'aux_rows':len(aux)},ensure_ascii=False))
if __name__=='__main__':main()
