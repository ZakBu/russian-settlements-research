from __future__ import annotations

import json

import pandas as pd
import pytest
from openpyxl import Workbook

from research_rebuild.mass_linkage import apply_reviewed_crimea2014_scoped_observations as app


def _synthetic_tables():
    observations = []
    edges = []
    points = []
    current_rows = []
    current_points = []
    for i in range(994):
        old = f"2014:source:{i}"
        new = f"2021:source:{i}"
        observations.append({
            "historical_source_record_id": old,
            "observation_year_actual": "2014",
            "observation_time_precision": "year_only_as_P585_precision_9",
            "population": "10",
            "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256,
            "source_workbook_path": "/source.xlsx",
            "source_sheet": "pub-01-03",
            "source_row": str(i + 1),
            "source_row_caption_literal": f"Город {i}",
            "source_row_type": "город",
            "raw_row_cells_json": json.dumps([f"Город {i}", 10, 5, 5, 50]),
            "population_scope": "single settlement row; current type and source row type match",
            "population_value_quality": "direct_published_2014_official_workbook_row_replayed",
            "source_section": "Republic of Crimea",
            "admin_path_raw_json": "[]",
            "admin_heading_raw": "",
            "current_qid": f"Q{i+1}",
            "current_2021_source_record_id": new,
            "source_item_qid": "Q127387785",
            "source_item_title": "Rosstat 2014 publication",
            "review_status": app.EXPECTED_EDGE_STATUS,
        })
        edges.append({
            "from_source_record_id": old, "from_year": "2014", "to_source_record_id": new,
            "to_year": "2021", "relation": "same_place", "review_status": app.EXPECTED_EDGE_STATUS,
            "qid": f"Q{i+1}",
            "source_population_row_id": old, "source_population": "10",
            "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256,
            "boundary_comparability_asserted": "false",
            "strict_Russian_2002_2010_2021_chain_eligible": "false",
        })
        current_rows.append({
            "record_type": "census", "observation_year": 2021, "source_record_id": new,
            "entity_id": f"settlement:{i}", "entity_category": "settlement",
            "settlement_name": f"Город {i}", "settlement_type": "город", "region_raw": "Республика Крым",
            "population_value": 11, "source_native_id": str(i),
        })
        current_points.append({
            "target_source_record_id": new, "target_year": 2021,
            "latitude": 44.0 + i / 10000, "longitude": 34.0 + i / 10000,
            "coordinate_admission_status": "reviewed_rule_accepted",
            "coordinate_source": "tochno_dadata" if i % 2 else "wikidata_p625",
            "coordinate_provider": "tochno_dadata" if i % 2 else "wikidata_p625",
            "coordinate_provider_id": f"provider:{i}",
            "coordinate_quality": "accepted_direct_current_source_point",
            "coordinate_provenance": "actual ledger origin",
            "coordinate_provider_family": "tochno" if i % 2 else "wikidata",
            "source_file": f"source-{i}.parquet",
            "source_sha256": "b" * 64,
            "source_locator": f"row={i}",
        })
        if i < 993:
            points.append({
                "target_source_record_id": old, "target_year": "2014",
                "latitude": 44.0 + i / 10000, "longitude": 34.0 + i / 10000,
                "coordinate_use_review_status": app.EXPECTED_EDGE_STATUS,
                "coordinate_admission_status": "candidate_retrospective_current_point_use",
                "coordinate_source_record_id": new,
                "coordinate_source": "Wikidata P625 representative point",
                "source_2021_carrier_target_source_record_id": new,
                "coordinate_measurement_date_unknown": "true",
                "direct_historical_coordinate_measurement": "false",
                "boundary_comparability_asserted": "false",
                "population_scope_comparability_asserted": "false",
                "point_origin_kind": "wikidata_p625_current_2021_point",
                "point_origin_file": "cached-wikidata.jsonl.gz",
                "point_origin_sha256": "a" * 64,
                "point_origin_locator": f"line={i}",
            })
    return tuple(pd.DataFrame(x) for x in (observations, edges, points, current_rows, current_points))


def test_literal_dash_values_remain_unknown_and_numeric_gender_sums_are_checked(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "pub-01-03"
    ws.append(["Settlement", 10, "—", "—", 1])
    wb.save(path)
    monkeypatch.setattr(app, "EXPECTED_WORKBOOK_SHA256", app.sha256(path))
    observations = pd.DataFrame([{
        "historical_source_record_id": "2014:x", "observation_year_actual": "2014", "population": "10",
        "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256, "source_sheet": "pub-01-03", "source_row": "1",
        "source_workbook_path": str(path),
        "source_row_caption_literal": "Settlement", "source_row_type": "город",
        "raw_row_cells_json": '["Settlement",10,"—","—",1]',
        "population_scope": "single settlement row; current type and source row type match",
        "population_value_quality": "direct_published_2014_official_workbook_row_replayed",
        "review_status": app.EXPECTED_EDGE_STATUS,
    }])
    result = app.replay_workbook_rows(observations, path)["2014:x"]
    assert result["total"] == 10
    assert result["male"] is None and result["female"] is None
    assert result["male_raw"] == "—" and result["female_raw"] == "—"

    bad_scope = observations.copy()
    bad_scope.loc[0, "population_scope"] = "federal aggregate"
    with pytest.raises(ValueError, match="non-physical or aggregate"):
        app.replay_workbook_rows(bad_scope, path)

    bad_source_pin = observations.copy()
    bad_source_pin.loc[0, "source_workbook_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="workbook hash"):
        app.replay_workbook_rows(bad_source_pin, path)


def test_numeric_gender_cells_must_sum_to_total(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "pub-01-03"
    ws.append(["Settlement", 10, 4, 5, 1])
    wb.save(path)
    monkeypatch.setattr(app, "EXPECTED_WORKBOOK_SHA256", app.sha256(path))
    observations = pd.DataFrame([{
        "historical_source_record_id": "2014:x", "observation_year_actual": "2014", "population": "10",
        "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256, "source_workbook_path": str(path),
        "source_sheet": "pub-01-03", "source_row": "1", "source_row_caption_literal": "Settlement",
        "source_row_type": "город", "raw_row_cells_json": '["Settlement",10,4,5,1]',
        "population_scope": "single settlement row; current type and source row type match",
        "population_value_quality": "direct_published_2014_official_workbook_row_replayed",
        "review_status": app.EXPECTED_EDGE_STATUS,
    }])
    with pytest.raises(ValueError, match="male/female totals"):
        app.replay_workbook_rows(observations, path)


def test_workbook_population_mismatch_is_rejected(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "pub-01-03"
    ws.append(["Settlement", 11, 5, 6, 1])
    wb.save(path)
    monkeypatch.setattr(app, "EXPECTED_WORKBOOK_SHA256", app.sha256(path))
    observations = pd.DataFrame([{
        "historical_source_record_id": "2014:x", "observation_year_actual": "2014", "population": "10",
        "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256, "source_sheet": "pub-01-03", "source_row": "1",
        "source_workbook_path": str(path),
        "source_row_caption_literal": "Settlement", "source_row_type": "город",
        "raw_row_cells_json": '["Settlement",11,5,6,1]',
        "population_scope": "single settlement row; current type and source row type match",
        "population_value_quality": "direct_published_2014_official_workbook_row_replayed",
        "review_status": app.EXPECTED_EDGE_STATUS,
    }])
    with pytest.raises(ValueError, match="population mismatch"):
        app.replay_workbook_rows(observations, path)


def test_wrong_selected_year_is_rejected(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "pub-01-03"
    ws.append(["Settlement", 10, 5, 5, 1])
    wb.save(path)
    monkeypatch.setattr(app, "EXPECTED_WORKBOOK_SHA256", app.sha256(path))
    observations = pd.DataFrame([{
        "historical_source_record_id": "2002:x", "observation_year_actual": "2002", "population": "10",
        "source_workbook_sha256": app.EXPECTED_WORKBOOK_SHA256, "source_sheet": "pub-01-03", "source_row": "1",
        "source_workbook_path": str(path),
        "source_row_caption_literal": "Settlement", "source_row_type": "город",
        "raw_row_cells_json": '["Settlement",10,5,5,1]',
        "population_scope": "single settlement row; current type and source row type match",
        "population_value_quality": "direct_published_2014_official_workbook_row_replayed",
        "review_status": app.EXPECTED_EDGE_STATUS,
    }])
    with pytest.raises(ValueError, match="only 2014"):
        app.replay_workbook_rows(observations, path)


def test_point_uses_must_match_the_reviewed_current_2021_target_and_coordinates():
    obs, edges, points, current, current_points = _synthetic_tables()
    app.validate_review_tables(obs, edges, points, current, current_points)

    wrong_target = edges.copy()
    wrong_target.loc[0, "to_source_record_id"] = "2021:not-selected"
    with pytest.raises(ValueError, match="absent from the pinned 2021 core"):
        app.validate_review_tables(obs, wrong_target, points, current, current_points)

    wrong_point = points.copy()
    wrong_point.loc[0, "latitude"] = 1.0
    with pytest.raises(ValueError, match="differs from current accepted point"):
        app.validate_review_tables(obs, edges, wrong_point, current, current_points)


def test_long_rows_preserve_scope_unknowns_and_do_not_claim_2014_point_measurement(tmp_path):
    obs, edges, points, current, current_points = _synthetic_tables()
    current, current_points = app.validate_review_tables(obs, edges, points, current, current_points)
    cells = {str(r.historical_source_record_id): {"total": 10, "male": 5, "female": 5,
              "male_raw": 5, "female_raw": 5, "raw_cells": ["caption", 10, 5, 5, 50]}
             for r in obs.itertuples(index=False)}
    cells["2014:source:0"]["female_raw"] = "—"
    out = app.create_long_observations(obs, edges, points, current, current_points, cells, "/source.xlsx", app.EXPECTED_WORKBOOK_SHA256)
    assert len(out) == 994
    assert out.strict_Russian_2002_2010_2021_chain_eligible.eq(False).all()
    assert out.boundary_comparability_asserted.eq(False).all()
    missing = out.loc[out.source_record_id.eq("2014:source:993")].iloc[0]
    assert pd.isna(missing.latitude) and pd.isna(missing.longitude)
    assert pd.isna(missing.coordinate_admission_status)
    measured = out.loc[out.source_record_id.eq("2014:source:0")].iloc[0]
    assert measured.coordinate_measurement_date_unknown
    assert not bool(measured.historical_coordinate_asserted)
    assert measured.coordinate_source == "wikidata_p625"
    assert measured.coordinate_provider == "wikidata_p625"
    assert measured.point_source_file == "source-0.parquet"
    assert measured.population_male_raw == "5" and measured.population_female_raw == "—"
    dadata = out.loc[out.source_record_id.eq("2014:source:1")].iloc[0]
    assert dadata.coordinate_source == "tochno_dadata"
    assert dadata.coordinate_provider == "tochno_dadata"
    assert dadata.coordinate_provider_family == "tochno"
    # The actual source mixes numeric and dash gender cells; convenience raw
    # fields remain serializable while the typed original cells stay in JSON.
    out.to_parquet(tmp_path / "observations.parquet", index=False)


def test_union_by_name_preserves_unknown_baseline_columns(tmp_path):
    base = tmp_path / "base.parquet"
    output = tmp_path / "applied.parquet"
    pd.DataFrame([{
        "observation_id": "census:old", "record_type": "census", "population_value": 7,
        "mystery_existing_provenance": "leave-me-byte-value",
    }]).to_parquet(base, index=False)
    new = pd.DataFrame([{
        "observation_id": "official2014:new", "record_type": "scoped_official_observation",
        "population_value": 9, "new_scoped_provenance": "2014-only",
    }])
    summary = app.append_observations_by_name(base, new, output)
    result = pd.read_parquet(output)
    old = result.loc[result.observation_id.eq("census:old")].iloc[0]
    added = result.loc[result.observation_id.eq("official2014:new")].iloc[0]
    assert summary["input_rows"] == 1 and summary["output_rows"] == 2
    assert old.mystery_existing_provenance == "leave-me-byte-value"
    assert pd.isna(old.new_scoped_provenance)
    assert added.new_scoped_provenance == "2014-only"
    assert pd.isna(added.mystery_existing_provenance)


def test_union_refuses_any_change_to_original_long_values(tmp_path):
    base = tmp_path / "base.parquet"
    output = tmp_path / "output.parquet"
    pd.DataFrame([{"observation_id": "old:1", "record_type": "census", "population_value": 7,
                   "optional_old_field": None}]).to_parquet(base, index=False)
    new = pd.DataFrame([{"observation_id": "new:1", "record_type": "scoped", "population_value": 9}])
    app.append_observations_by_name(base, new, output)
    import duckdb
    con = duckdb.connect()
    tampered = tmp_path / "tampered.parquet"
    input_sql = "'" + str(output).replace("'", "''") + "'"
    output_sql = "'" + str(tampered).replace("'", "''") + "'"
    con.execute(f"COPY (SELECT observation_id, record_type, CASE WHEN observation_id='old:1' THEN 8 ELSE population_value END AS population_value, optional_old_field FROM read_parquet({input_sql})) TO {output_sql} (FORMAT PARQUET)")
    con.close()
    con = duckdb.connect()
    with pytest.raises(ValueError, match="union changed"):
        app._verify_original_rows_unchanged(con, str(base), str(tampered))
    con.close()


def test_wrong_review_receipt_hash_is_rejected(tmp_path, monkeypatch):
    receipt = tmp_path / "review.json"
    receipt.write_text(json.dumps({
        "status": app.EXPECTED_REVIEW_STATUS,
        "candidate_only": True,
        "not_an_admission_or_population_overlay_receipt": True,
        "checks": {
            "eligible_identity_observations": 994, "eligible_identity_edges": 994,
            "eligible_retrospective_point_uses": 993,
            "strict_Russian_2002_2010_2021_chain_contribution": 0,
        },
    }))
    fake = {}
    for name in ("full_long", "selected_population", "identity_graph", "point_ledger", "observations_csv",
                 "identity_edges_csv", "retrospective_points_csv", "source_workbook", "source_fetch_receipt"):
        p = tmp_path / name
        p.write_text("fixture")
        fake[name] = p
    sha = {k: app.sha256(v) for k, v in fake.items()}
    fetch = tmp_path / "fetch.json"
    workbook = fake["source_workbook"]
    fetch.write_text(json.dumps({"url": app.SOURCE_ARCHIVE_URL, "final_url": app.SOURCE_ARCHIVE_URL,
                                 "response_sha256": sha["source_workbook"], "response_bytes": workbook.stat().st_size}))
    fake["source_fetch_receipt"] = fetch
    sha["source_fetch_receipt"] = app.sha256(fetch)
    monkeypatch.setattr(app, "EXPECTED_REVIEW_SHA256", "0" * 64)
    monkeypatch.setattr(app, "EXPECTED_WORKBOOK_SHA256", sha["source_workbook"])
    monkeypatch.setattr(app, "EXPECTED_OBSERVATIONS_SHA256", sha["observations_csv"])
    monkeypatch.setattr(app, "EXPECTED_EDGES_SHA256", sha["identity_edges_csv"])
    monkeypatch.setattr(app, "EXPECTED_RETRO_POINTS_SHA256", sha["retrospective_points_csv"])
    monkeypatch.setattr(app, "EXPECTED_SOURCE_FETCH_RECEIPT_SHA256", sha["source_fetch_receipt"])
    manifest = {"inputs": {k: {"path": str(v), "sha256": sha[k]} for k, v in fake.items()}}
    manifest["inputs"]["review_receipt"] = {"path": str(receipt), "sha256": app.sha256(receipt)}
    with pytest.raises(ValueError, match="not the approved frozen receipt"):
        app.verify_review_packet(manifest)
