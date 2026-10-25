"""Official physical-retail regional series; local extraction is default.
--acquire downloads only the observed official workbook URLs.
No temporal shop opening/closing is inferred from mapping timestamps.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse,csv,hashlib,json,re
import openpyxl

OUT=Path(__file__).resolve().parent
RAW=OUT/'raw'
FILES=['obekty_rozn_2025.xlsx','retail_catering_2022-2024.xlsx','Fact_obesp_torg_2025.xls']
ROOT='https://rosstat.gov.ru/storage/mediabank/'
parser=argparse.ArgumentParser()
parser.add_argument('--acquire',action='store_true')
args=parser.parse_args()
if args.acquire:
    import requests
    receipts=[]
    for name in FILES:
        # Public Rosstat downloads; retain explicit TLS limitation in provenance.
        r=requests.get(ROOT+name,timeout=30,verify=False)
        r.raise_for_status()
        (RAW/name).write_bytes(r.content)
        receipts.append({'name':name,'url':r.url,'status':r.status_code,
                         'bytes':len(r.content),'tls_verification':False,
                         'retrieved_at_utc':datetime.now(timezone.utc).isoformat(),
                         'sha256':hashlib.sha256(r.content).hexdigest()})
    (RAW/'rosstat_download_receipts.json').write_text(json.dumps(receipts,ensure_ascii=False,indent=2))

def clean(x):
    return re.sub(r'\s+',' ',str(x or '')).strip()

def num(x):
    # Dash is deliberately not imputed as zero.
    return float(x) if isinstance(x,(int,float)) else None

def level(label):
    if label.startswith('Российская Федерация'):return 'country_published_total'
    if 'федеральный округ' in label:return 'federal_district_published_total'
    return 'regional_or_nested_regional_unit'

def write_csv(name,rows):
    with (OUT/name).open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]))
        w.writeheader();w.writerows(rows)

annual=[]
notes=[]
headers=[]
wb=openpyxl.load_workbook(RAW/FILES[0],data_only=True,read_only=True)
notes.extend([{'file':FILES[0],'sheet':'Оглавление','row':i,
               'text':clean(' '.join(str(c) for c in r if c is not None))}
              for i,r in enumerate(wb['Оглавление'].values,1) if i==9])
for year in range(2020,2026):
    s=wb[str(year)]
    allrows=list(s.values)
    cols={'shops_total':2,'shop_sales_area_m2':3}
    if year<=2023:
        cols.update({'hypermarkets_and_supermarkets':4,'minimarkets':6,'pavilions':8,'pharmacies_and_pharmacy_shops':10})
    elif year==2024:
        cols.update({'hypermarkets':4,'supermarkets':6,'specialized_food_stores':8,
                     'specialized_nonfood_stores':10,'minimarkets':12,'department_stores':14,
                     'other_stores':16,'pavilions':18,'pharmacies_and_pharmacy_shops':22})
    else:
        cols.update({'hypermarkets':4,'supermarkets':6,'minimarkets':8,
                     'universal_and_specialized_food_stores':10,
                     'universal_and_specialized_nonfood_stores':12,'pavilions':14,'pharmacy_organizations':18})
    header_rows=[2,3] if year<=2023 else [2,3,4]
    for measure,c in cols.items():
        headers.append({'file':FILES[0],'sheet':str(year),'measure':measure,
                        'column':openpyxl.utils.get_column_letter(c),
                        'header_text':' | '.join(clean(allrows[r][c-1]) for r in header_rows)})
    for i,row in enumerate(allrows,1):
        label=clean(row[0])
        if not label:continue
        if i>7 and num(row[1]) is None:
            if re.match(r'^\d\)',label): notes.append({'file':FILES[0],'sheet':str(year),'row':i,'text':label})
        if row[1] is None:continue
        for measure,c in cols.items():
            val=row[c-1]
            annual.append({'territory_raw':label,'territory_level':level(label),
                'observation_date':f'{year}-12-31','year':year,'measure':measure,
                'value':num(val),'value_raw':val,'unit':'m2' if measure.endswith('_m2') else 'objects',
                'source_file':FILES[0],'source_sheet':str(year),'source_row':i,
                'source_cell':f'{openpyxl.utils.get_column_letter(c)}{i}',
                'source_url':ROOT+FILES[0],
                'collection_form':'1-MO' if year<=2021 else '1-TORG(MO)',
                'scope_flag':'source_note_excludes_StP_Sevastopol_2020_2023_but_some_annual_rows_present_needs_audit' if year<=2023 else 'source_territorial_coverage_changed',
                'interpretation':'physical_object_inventory_not_enterprise_registration'})
quarterly=[]
wbq=openpyxl.load_workbook(RAW/FILES[1],data_only=True,read_only=True)
for year in range(2022,2025):
    s=wbq[str(year)];allrows=list(s.values)
    cat=None
    for c in range(2,66):
        r4=clean(allrows[3][c-1]);r5=clean(allrows[4][c-1])
        if r4 and 'в том числе' not in r4: cat=r4
        if r5:cat=r5
        if (c-2)%4==0:
            headers.append({'file':FILES[1],'sheet':str(year),'measure':cat,
                            'column':openpyxl.utils.get_column_letter(c),'header_text':cat})
        quarter=(c-2)%4
        date=f'{year}-'+['04-01','07-01','10-01','12-31'][quarter]
        for i,row in enumerate(allrows,1):
            if i<7:continue
            label=clean(row[0])
            if not label:continue
            if row[1] is None:
                if c==2 and re.match(r'^\d\)',label): notes.append({'file':FILES[1],'sheet':str(year),'row':i,'text':label})
                continue
            val=row[c-1]
            quarterly.append({'territory_raw':label,'territory_level':level(label),
                'observation_date':date,'year':year,'category_raw':cat,
                'value':num(val),'value_raw':val,'unit':'objects',
                'source_file':FILES[1],'source_sheet':str(year),'source_row':i,
                'source_cell':f'{openpyxl.utils.get_column_letter(c)}{i}',
                'source_url':ROOT+FILES[1],
                'interpretation':'physical_object_inventory_not_enterprise_registration'})

write_csv('rosstat_physical_retail_annual_2020_2025_long.csv',annual)
write_csv('rosstat_physical_retail_quarterly_2022_2024_long.csv',quarterly)
write_csv('source_headers.csv',headers)
(OUT/'source_notes.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2))

# Independently reconcile December totals in the two official workbooks.
annual_totals={(r['year'],r['territory_raw'].strip()):r['value']
               for r in annual if r['measure']=='shops_total'}
reconcile=[]
for r in quarterly:
    if r['category_raw']!='Магазины (всего)' or not r['observation_date'].endswith('12-31'):continue
    key=(r['year'],r['territory_raw'].strip())
    # Do not manufacture joins across spelling variants; exact labels only.
    if key not in annual_totals:continue
    av=annual_totals[key]
    if av is None or r['value'] is None:continue
    reconcile.append({'year':r['year'],'territory_raw':key[1],
                      'annual_value':av,'quarterly_December_value':r['value'],
                      'difference':r['value']-av})
assert reconcile
assert all(r['difference']==0 for r in reconcile)
write_csv('annual_quarterly_total_reconciliation.csv',reconcile)

keys=[(r['territory_raw'],r['observation_date'],r['measure']) for r in annual]
assert len(keys)==len(set(keys))
keysq=[(r['territory_raw'],r['observation_date'],r['category_raw']) for r in quarterly]
assert len(keysq)==len(set(keysq))
assert all(r['value'] is None or r['value']>=0 for r in annual+quarterly)
# Published country totals are not re-summed from nested regional entities.
country=[r for r in annual if r['territory_level']=='country_published_total' and r['measure']=='shops_total']
manifest={'retrieval_date':'2026-10-24','files':[{'name':n,'url':ROOT+n,
    'sha256':hashlib.sha256((RAW/n).read_bytes()).hexdigest(),
    'bytes':(RAW/n).stat().st_size} for n in FILES],
    'annual_rows':len(annual),'quarterly_rows':len(quarterly),
    'annual_territories_by_year':{str(y):len({r['territory_raw'] for r in annual if r['year']==y}) for y in range(2020,2026)},
    'reconciled_exact_label_year_totals':len(reconcile),
    'nonzero_reconciliation_differences':sum(r['difference']!=0 for r in reconcile),
    'annual_missing_numeric_cells':sum(r['value'] is None for r in annual),
    'quarterly_missing_numeric_cells':sum(r['value'] is None for r in quarterly),
    'country_total_not_constant_territorial_scope':[{k:r[k] for k in ['year','value','source_cell']} for r in country],
    'scope':'regions and nested regional aggregates; NOT municipalities or individual settlements',
    'source_scope_conflict':'Cover note excludes StP/Sevastopol2020-2023, but annual2020/2021 contain numeric federal-city rows; do not infer invariant national scope',
    'methodological_break':'1-MO through2021 to1-TORG(MO) from2022; category revisions2024/2025; federal-city coverage changes2024',
    'dash_policy':'raw dash preserved; missing numeric, never imputed as zero',
    'checks':'PASS unique keys, nonnegative observed values, exact-label December totals cross-source equal',
    'tls':'Rosstat public downloads required verify=False after local CA validation failure; URLs/file hashes/raw responses retained'}
(OUT/'acquisition_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False,indent=2))
