# coding: utf-8
from pathlib import Path
import csv,json,hashlib,re
from collections import Counter
from openpyxl import load_workbook
base=Path(__file__).resolve().parent;raw=base/'acquisition_raw'
def region_norm(text):
 text=str(text or '').lower().replace('ё','е')
 text=re.split(r',?\s+(?:включая|без учета|без автоном|в том числе)',text)[0]
 text=text.replace('кемеровская область - кузбасс','кемеровская область')
 text=re.sub(r'\([^)]*(?:автоном|кроме)[^)]*\)','',text)
 text=re.sub(r'\b(?:республика|область|край|автономный|автономная|округ)\b',' ',text)
 return re.sub(r'[^а-яa-z0-9]+',' ',text).strip()
files=[(2022,'Chisl_nasel_RF_MO_01-01-2022.xlsx'),(2023,'BUL_MO_2023.xlsx'),(2024,'BUL_MO_2024.xlsx'),(2025,'original_cyrillic_Сhisl_MO_01-01-2025.xlsx')]
inventory=[];allrecords=[]
for year,fn in files:
 f=raw/fn;w=load_workbook(f,data_only=True,read_only=True);sheet=w[w.sheetnames[-1]];ni=1 if year>=2024 else 0
 region_sheet=w[w.sheetnames[-2]];regions={str(r[1]).strip():str(r[1]).strip() for r in region_sheet.values if len(r)>1 and isinstance(r[0],(int,float)) and r[1]}
 region_keys={region_norm(x):x for x in regions};current_region=None;previous_municipality=None;rows=[]
 for lineno,r in enumerate(sheet.values,1):
  name=r[ni]
  if not isinstance(name,str):continue
  name=name.strip();total=r[ni+1];urban=r[ni+2];rural=r[ni+3]
  if name in regions:current_region=regions[name];continue
  if region_norm(name) in region_keys:current_region=region_keys[region_norm(name)];continue
  if not isinstance(total,(int,float)):continue
  if 'город федерального значения' in name.lower() or any(s in name.lower() for s in ['городской округ','муниципальный','городское поселение','сельское поселение']):previous_municipality=name;continue
  m=re.match(r'^(г|пгт|рп|п)\s+(.+)',name)
  if not m:continue
  if urban!=total or rural not in [0,None,'-','–','—']:continue
  row={'year':year,'reference_date':f'{year}-01-01','region':current_region,'name_raw':name,'name_without_type':m[2].strip(),'type_raw':m[1],'population_persons':total,'urban_population':urban,'rural_population':rural,'source_territory_code_raw':str(r[0]) if year>=2024 else None,'preceding_municipal_context':previous_municipality,'sheet':sheet.title,'source_excel_row':lineno,'source_filename':fn,'source_url':'https://rosstat.gov.ru/storage/mediabank/'+fn.removeprefix('original_cyrillic_'),'classification':'explicit_urban_name_candidate_requires_registry_dedup','census_basis':'without_VPN2020_results' if year==2022 else ('with_VPN2020_results' if year==2023 else 'latest_census_methodology_header_not_explicitly_rebased')}
  rows.append(row)
 out=base/f'rosstat_urban_np_candidates_{year}.csv'
 with out.open('w',encoding='utf-8-sig',newline='') as h:
  writer=csv.DictWriter(h,fieldnames=rows[0].keys());writer.writeheader();writer.writerows(rows)
 counts=Counter(x['type_raw'] for x in rows)
 inv={'observation_year':year,'reference_date':f'{year}-01-01','source_filename':fn,'source_url':'https://rosstat.gov.ru/storage/mediabank/'+fn.removeprefix('original_cyrillic_'),'local_original':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'size_bytes':f.stat().st_size,'source_sheet':sheet.title,'source_sheet_total_rows':sheet.max_row,'urban_named_candidate_rows':len(rows),'candidate_prefix_counts':dict(counts),'region_labels_recognized':len(regions),'region_labels_present_in_candidates':len({x['region'] for x in rows if x['region']}),'output_csv':str(out),'census_basis_warning':'2022предисловие:безучетаВПН2020;2023предисловие:сучетомВПН2020. Не трактовать2022–2023какчистуюдинамику.','coverage_warning':'Это строки с явным префиксомг/пгт/рп/п и городским населением=всему. Не утверждать число уникальных НП до сверки с реестром: включены некоторые внутригородские муниципальные образования федеральных городов и возможны повторения/особенности статуса.'}
 inventory.append(inv);allrecords.extend(rows)
(base/'rosstat_acquisition_inventory.json').write_text(json.dumps({'checked_on':'2026-10-24','files':inventory,'notes':['УровеньМОнеподменяетНП:выписанытолькострокисявнымименемгородскогоНП.','Коды2024обозначеныТЕРСОН-МО,несчитатьихобычнымОКТМО11знаков.','2023/2022неткодовнаосновномлистевпроверенныхвыпусках.','Сельскиепоселениянеявляютсясельскиминаселённымипунктами;всельскуюпанельнепереносить.']},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(inventory,ensure_ascii=False,indent=2))
