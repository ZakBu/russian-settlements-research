"""Convert prepared wiki publications to editable DOCX with Pandoc and one Word profile.

The caller supplies already numbered figures, citation links and a shared bibliography.
Images must be captured or downloaded to local files before this step.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.image.image import Image
from lxml import html


def create_reference(path: Path):
    """The project's report profile; a university template may replace it explicitly."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(30), Mm(15)
    section.top_margin = section.bottom_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    styles = doc.styles
    names = ["Normal", "Body Text", "First Paragraph", "Compact", "Title", "Subtitle",
             "Caption", "Image Caption", "Table Caption", "Figure Caption", "Figure Source",
             "Bibliography", "Footnote Text", "Block Text", "TOC Heading", "Source Code"]
    names += [f"Heading {n}" for n in range(1, 10)] + [f"TOC {n}" for n in range(1, 10)]
    for name in names:
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != "Normal":
            style.base_style = styles["Normal"]
        style.font.name, style.font.size = "Times New Roman", Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Times New Roman")
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "ru-RU")
        style.element.get_or_add_rPr().append(lang)
        p = style.paragraph_format
        p.line_spacing, p.first_line_indent = 1.5, Cm(1.25)
        p.space_before = p.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if name.startswith("Heading "):
            style.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.space_before, p.space_after, p.keep_with_next = Pt(14), Pt(7), True
            p.page_break_before = name == "Heading 1"
        if name in {"Title", "Subtitle", "TOC Heading"}:
            p.alignment, p.first_line_indent = WD_ALIGN_PARAGRAPH.CENTER, Cm(0)
            style.font.bold = name != "Subtitle"
            p.space_after = Pt(14)
        if name in {"Figure Caption", "Image Caption", "Caption", "Table Caption", "Figure Source", "Footnote Text"}:
            style.font.size = Pt(12)
            p.line_spacing, p.first_line_indent = 1, Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if name in {"Figure Source", "Table Caption", "Footnote Text"} else WD_ALIGN_PARAGRAPH.CENTER
            p.keep_with_next = name in {"Figure Caption", "Image Caption", "Table Caption"}
            p.space_after = Pt(6)
        if name.startswith("TOC ") or name == "Compact":
            p.first_line_indent = Cm(0)
        if name == "Bibliography":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if name == "Source Code":
            style.font.name, style.font.size = "Courier New", Pt(10)
            p.line_spacing, p.first_line_indent = 1, Cm(0)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    section.footer.paragraphs[0]._p.append(page)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.save(path)


def _style_block(element, name):
    wrapper = html.Element("div", {"custom-style": name})
    element.addprevious(wrapper)
    wrapper.append(element)


def prepare_fragment(source: str, identity: str, title: str, resource_root: Path):
    """Expand native disclosures, keep only publication content and pin local images."""
    fragment = html.fragment_fromstring(source or "", create_parent="div")
    for element in fragment.xpath('.//script | .//style | .//nav | .//aside | .//button'):
        element.drop_tree()
    for element in fragment.xpath('.//*[@aria-hidden="true"] | .//*[contains(concat(" ", normalize-space(@class), " "), " heading-anchor ")]'):
        element.drop_tree()
    for element in fragment.xpath('.//details'):
        element.tag = "div"
        element.attrib.pop("open", None)
        for summary in element.xpath('./summary'):
            summary.tag = "p"
            strong = html.Element("strong")
            strong.text = summary.text
            summary.text = None
            for child in list(summary):
                strong.append(child)
            summary.append(strong)
    # Native HTML figure captions become explicit editable paragraphs, not a second image.
    for element in fragment.xpath('.//figure | .//figcaption'):
        element.tag = "div"
    for element in list(fragment.xpath('.//p')):
        text = element.text_content().strip()
        if re.match(r"^Рисунок\s+\d+\s*[—–-]", text):
            _style_block(element, "Figure Caption")
        elif text.startswith("Источник:"):
            _style_block(element, "Figure Source")
        elif re.match(r"^Таблица\s+\d+\s*[—–-]", text):
            _style_block(element, "Table Caption")
    for image in fragment.xpath('.//img'):
        source_url = image.get("src", "")
        parsed = urlparse(source_url)
        if parsed.scheme not in {"", "file"} or parsed.netloc:
            raise ValueError(f"Capture images locally before Word export: {source_url}")
        image_path = Path(unquote(parsed.path if parsed.scheme else source_url))
        if not image_path.is_absolute():
            image_path = resource_root / image_path
        image_path = image_path.resolve()
        if not image_path.is_file():
            raise ValueError(f"Image file does not exist: {image_path}")
        picture = Image.from_file(str(image_path))
        max_width, max_height = 165 / 25.4 * 96, 200 / 25.4 * 96
        width = min(max_width, max_height * picture.px_width / picture.px_height)
        image.set("src", str(image_path))
        image.set("width", f"{width:.2f}")
        image.attrib.pop("height", None)
        image.attrib.pop("style", None)
    if title and not fragment.xpath('.//h1'):
        heading = html.Element("h1")
        heading.text = title
        fragment.insert(0, heading)
    prefix = "chapter-" + re.sub(r"[^a-zA-Z0-9_-]+", "-", identity) + "-"
    identifiers = {element.get("id"): prefix + element.get("id") for element in fragment.xpath('.//*[@id]')}
    for element in fragment.xpath('.//*[@id]'):
        element.set("id", identifiers[element.get("id")])
    for link in fragment.xpath('.//a[starts-with(@href, "#")]'):
        target = unquote(link.get("href")[1:])
        if target in identifiers:
            link.set("href", "#" + identifiers[target])
    return html.tostring(fragment, encoding="unicode")


def native_footnotes(ast):
    """Map Marked's endnote HTML to Pandoc's native Note nodes before DOCX writing."""
    notes = {}
    dropped = object()

    def strip_backlinks(value):
        if isinstance(value, list):
            return [result for child in value if (result := strip_backlinks(child)) is not dropped]
        if isinstance(value, dict):
            if value.get("t") == "Link":
                attrs = value["c"][0]
                if "footnote-back" in attrs[1] or any(k.endswith("footnote-backref") or (k == "role" and v == "doc-backlink") for k, v in attrs[2]):
                    return dropped
            return {key: strip_backlinks(child) for key, child in value.items()}
        return value

    def collect(value):
        if isinstance(value, dict):
            if value.get("t") == "Div":
                attrs, blocks = value["c"]
                if attrs[0] and not ("footnotes" in attrs[1]):
                    notes[attrs[0]] = strip_backlinks(deepcopy(blocks))
            for child in value.values():
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)

    footnote_sections = [block for block in ast.get("blocks", []) if block.get("t") == "Div" and "footnotes" in block["c"][0][1]]
    # A section may be inside a chapter wrapper rather than at document root.
    def find_sections(value):
        if isinstance(value, dict):
            if value.get("t") == "Div" and "footnotes" in value["c"][0][1] and value not in footnote_sections:
                footnote_sections.append(value)
            for child in value.values():
                find_sections(child)
        elif isinstance(value, list):
            for child in value:
                find_sections(child)
    find_sections(ast)
    for section in footnote_sections:
        collect(section["c"][1])

    def replace(value):
        if isinstance(value, list):
            return [result for child in value if (result := replace(child)) is not dropped]
        if not isinstance(value, dict):
            return value
        if value.get("t") == "Div" and "footnotes" in value["c"][0][1]:
            return dropped
        if value.get("t") == "Link" and value["c"][2][0].startswith("#"):
            target = unquote(value["c"][2][0][1:])
            if target in notes:
                return {"t": "Note", "c": deepcopy(notes[target])}
        result = {key: replace(child) for key, child in value.items()}
        if result.get("t") == "Superscript" and len(result["c"]) == 1 and result["c"][0].get("t") == "Note":
            return result["c"][0]
        return result
    return replace(ast)


def build_html(job, resource_root):
    documents = job.get("documents")
    if not isinstance(documents, list) or not documents:
        raise ValueError("Provide at least one document in documents[]")
    chunks = []
    identities = [str(item.get("id", index + 1)) for index, item in enumerate(documents) if isinstance(item, dict)]
    if len(identities) != len(set(identities)):
        raise ValueError("Document IDs must be unique within a Word publication")
    for index, item in enumerate(documents):
        if not isinstance(item, dict) or not isinstance(item.get("html"), str):
            raise ValueError("Every document requires an HTML string")
        chunks.append(prepare_fragment(item["html"], str(item.get("id", index + 1)), item.get("title", ""), resource_root))
    bibliography = job.get("bibliographyHtml", "")
    if bibliography:
        fragment = html.fragment_fromstring(bibliography, create_parent="div")
        for heading in fragment.xpath('.//h1 | .//h2 | .//h3'):
            heading.tag = "h1"
            heading.set("class", "unnumbered")
        fragment.set("custom-style", "Bibliography")
        chunks.append(html.tostring(fragment, encoding="unicode"))
    return "<!doctype html><html lang=\"ru\"><head><meta charset=\"utf-8\"></head><body>" + "\n".join(chunks) + "</body></html>"


def run_pandoc(command, source):
    result = subprocess.run(command, input=source, text=True, capture_output=True)
    if result.returncode:
        raise ValueError("Pandoc conversion failed: " + result.stderr.strip())
    return result.stdout


def export_job(job, output: Path, resource_root: Path, reference_doc: Path):
    markup = build_html(job, resource_root)
    parsed = run_pandoc(["pandoc", "--from=html", "--to=json"], markup)
    ast = native_footnotes(json.loads(parsed))
    if job.get("title") and len(job["documents"]) > 1:
        ast["meta"]["title"] = {"t": "MetaInlines", "c": [{"t": "Str", "c": str(job["title"])}]}
    output.parent.mkdir(parents=True, exist_ok=True)
    command = ["pandoc", "--from=json", "--to=docx", "--standalone", "--reference-doc=" + str(reference_doc),
               "--metadata=lang:ru-RU", "--metadata=toc-title:Содержание", "--output=" + str(output)]
    if job.get("toc", True):
        command += ["--toc", "--toc-depth=3"]
    if job.get("numberSections", False):
        command += ["--number-sections"]
    run_pandoc(command, json.dumps(ast, ensure_ascii=False))
    # Images are separate from editable captions. Keep the picture with its caption.
    document = Document(output)
    for paragraph in document.paragraphs:
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.right_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith("Источник:"):
            paragraph.paragraph_format.keep_with_next = False
    if job.get("bibliographyHtml"):
        bibliography = html.fragment_fromstring(job["bibliographyHtml"], create_parent="div")
        titles = {heading.text_content().strip() for heading in bibliography.xpath('.//h1 | .//h2 | .//h3')}
        starts = [index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip() in titles]
        if starts and "Bibliography" in document.styles:
            for paragraph in document.paragraphs[starts[-1] + 1:]:
                paragraph.style = document.styles["Bibliography"]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
    for previous in list(document.settings.element.findall(qn("w:updateFields"))):
        document.settings.element.remove(previous)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    cache_toc_entries(document)
    document.save(output)
    return output


def cache_toc_entries(document):
    """Keep a useful linked outline before Word refreshes its native TOC field.

    Page numbers remain Word's responsibility; no guessed pagination is stored.
    """
    instructions = [node for node in document.element.iter(qn("w:instrText")) if (node.text or "").lstrip().startswith("TOC ")]
    if not instructions:
        return
    entries = []
    numeric_ids = [int(node.get(qn("w:id"))) for node in document.element.iter(qn("w:bookmarkStart")) if (node.get(qn("w:id")) or "").isdigit()]
    bookmark_id = max(numeric_ids, default=0) + 1
    for paragraph in document.paragraphs:
        match = re.fullmatch(r"Heading([1-3])", paragraph.style.style_id, re.I)
        if not match or not paragraph.text.strip():
            continue
        name = f"AtlasHeading{bookmark_id}"
        begin, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        begin.set(qn("w:id"), str(bookmark_id))
        begin.set(qn("w:name"), name)
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, begin)
        paragraph._p.append(end)
        entries.append((int(match[1]), paragraph.text, name))
        bookmark_id += 1
    if not entries:
        return
    instruction = instructions[0]
    paragraph = instruction.getparent()
    while paragraph.tag != qn("w:p"):
        paragraph = paragraph.getparent()
    command = instruction.text
    for child in list(paragraph):
        paragraph.remove(child)
    def field_run(kind):
        run, field = OxmlElement("w:r"), OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            field.set(qn("w:dirty"), "true")
        run.append(field)
        return run
    paragraph.append(field_run("begin"))
    run, text = OxmlElement("w:r"), OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = command
    run.append(text)
    paragraph.append(run)
    paragraph.append(field_run("separate"))
    current = paragraph
    for index, (depth, title, anchor) in enumerate(entries):
        if index:
            following = OxmlElement("w:p")
            current.addnext(following)
            current = following
        properties, style = OxmlElement("w:pPr"), OxmlElement("w:pStyle")
        style.set(qn("w:val"), f"TOC{depth}")
        properties.append(style)
        current.insert(0, properties)
        link, run, text = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:t")
        link.set(qn("w:anchor"), anchor)
        text.text = title
        run.append(text)
        link.append(run)
        current.append(link)
    current.append(field_run("end"))


def main(argv=None):
    parser = argparse.ArgumentParser(description="Export prepared wiki publications to formal editable DOCX with Pandoc")
    parser.add_argument("--input", required=True, help="JSON publication manifest or {jobs:[...]} batch")
    parser.add_argument("--output", help="DOCX path for a single manifest")
    parser.add_argument("--reference-doc", help="Existing university DOCX template instead of the default report profile")
    args = parser.parse_args(argv)
    if not shutil.which("pandoc"):
        parser.error("Install Pandoc before exporting Word documents")
    manifest_path = Path(args.input).resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    jobs = manifest.get("jobs", [manifest])
    if args.output and len(jobs) != 1:
        parser.error("--output is only valid for a single publication")
    with tempfile.TemporaryDirectory(prefix="atlas-word-") as temporary:
        reference = Path(args.reference_doc).resolve() if args.reference_doc else Path(temporary) / "reference.docx"
        if not args.reference_doc:
            create_reference(reference)
        for job in jobs:
            target = args.output if len(jobs) == 1 and args.output else job.get("output")
            if not target:
                parser.error("Provide --output for a single document or output in each batch job")
            result = export_job(job, Path(target).resolve(), manifest_path.parent, reference)
            print(result)


if __name__ == "__main__":
    main()
