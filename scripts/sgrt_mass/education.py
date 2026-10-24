#!/usr/bin/env python3
"""Read primary preschool/workforce books and national OSM education evidence.

The Rosstat publication edition is 2025, but Grok's CSV export drops workbook
headers and takes the final nonempty cell. Reference year and subcategory are
therefore NOT asserted from the filename. Preschool `oktmo` is often a numeric
measurement: invalid codes are retained as raw fields, never settlement keys.
"""
from __future__ import annotations
import argparse
import calendar
from datetime import date
import hashlib
import json
from pathlib import Path
import re
import subprocess

import numpy as np
import pandas as pd
from match import build_reference, match_addresses, normalize_name

SOURCE_COMMIT='3b37253baa283dd2e3fbb7b0d7e9c4f8fe27e6cb'
SOURCE_RAR_URL='https://rosstat.gov.ru/storage/mediabank/stat.info_NP_2025.rar'
SOURCE_RAR_SHA='520b79e3a7a9263f89712df3715b7f1fe1b121055c8bf56d1d6d0835b13cf9e3'


def extract_primary(rar_path,output):
    """Read real archive members rather than the lossy derived CSV columns."""
    import libarchive
    if sha(rar_path)!=SOURCE_RAR_SHA:raise ValueError('Original Rosstat RAR differs from pinned snapshot')
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    marker=output/'primary_extraction.json'
    if marker.exists():
        metadata=json.loads(marker.read_text())
        if metadata.get('archive_sha256')==SOURCE_RAR_SHA and all((output/p).exists() for p in metadata['selected_paths']):return metadata
    selected=[];errors=[];total=0;temporary=0
    with libarchive.file_reader(str(rar_path)) as archive:
        for entry in archive:
            total+=1;name=str(entry.pathname);base=name.split('/')[-1]
            if not base.endswith('_НП.xlsx'):continue
            if base.startswith('~$'):temporary+=1;continue
            head='2-10_Труд' in name and re.match(r'п\.?\s*9_',base)
            education='Образован' in name
            if not (head or education):continue
            try:
                data=b''.join(entry.get_blocks())
                if not data.startswith(b'PK'):raise ValueError('Expected an XLSX ZIP member')
                path=output/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data);selected.append(name)
            except Exception as e:errors.append({'path':name,'error':str(e)})
    metadata={'archive_sha256':SOURCE_RAR_SHA,'archive_bytes':Path(rar_path).stat().st_size,'archive_entries':total,
              'selected_paths':selected,'extraction_errors':errors,'office_temporary_members_excluded':temporary}
    marker.write_text(json.dumps(metadata,ensure_ascii=False,indent=2)+'\n')
    return metadata


def primary_cell(value):
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        return float(value),'observed'
    text='' if value is None else str(value).strip()
    if not text:return np.nan,'missing'
    if text in {'...','…'}:return np.nan,'suppressed'
    if text in {'-','—','–'}:return np.nan,'dash_not_assumed_zero'
    # Literal numeric strings occur in some books; commas are decimal marks.
    if re.fullmatch(r'-?\d+(?:[,.]\d+)?',text):return float(text.replace(',','.')),'observed'
    return np.nan,'unparsed_literal'


def parse_primary_books(primary_root):
    """Preserve totals and components as separate cells, with actual headers."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    books=[];rows=[];cells=[];errors=[]
    root=Path(primary_root)
    for path in sorted(root.rglob('*.xlsx')):
        if path.name.startswith('~$'):continue
        match=re.search(r'п(?:п)?\.?\s*(\d+)',path.name)
        if match is None:continue
        metric=int(match.group(1));relative=path.relative_to(root).as_posix()
        if metric not in {9,*range(36,52)}:continue
        try:
            wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
            for sheet in wb:
                matrix=list(sheet.iter_rows(values_only=True))
                header=None;letter_code_column=False;first_data_override=None
                for i,row in enumerate(matrix[:15]):
                    first='' if not row or row[0] is None else re.sub(r'\s+',' ',str(row[0])).strip().lower()
                    if first.startswith(('населенные пункты','населённые пункты')) or first=='наименование':header=i;break
                    if first in {'а','a'}:
                        header=i
                        letter_code_column=len(row)>1 and str(row[1]).strip().lower() in {'б','b'}
                        break
                if header is None:
                    for i,row in enumerate(matrix[:20]):
                        first='' if not row or row[0] is None else str(row[0]).strip()
                        if re.match(r'^(?:\d{11}\s*[-–]\s*)?(?:г|д|с|п|пгт|рп|нп|аул|х|ст-ца|ж/д|ст)\.?\s+',first,re.I):
                            header=i-1;first_data_override=i;break
                if header is None:
                    errors.append({'path':relative,'sheet':sheet.title,'error':'settlement_header_not_identified'});continue
                first_data=header+1
                if first_data_override is not None:first_data=first_data_override
                else:
                    for i in range(header+1,min(header+8,len(matrix))):
                        first=matrix[i][0] if matrix[i] else None
                        if first is not None and str(first).strip() not in {'А','Б','A','B'}:first_data=i;break
                preamble=' '.join(str(v) for row in matrix[:header+1] for v in row if v is not None)
                years=sorted(set(map(int,re.findall(r'\b20[0-3]\d\b',preamble))))
                year=years[0] if len(years)==1 else None
                width=max(map(len,matrix)) if matrix else 0
                headers={j:' | '.join(str(row[j]).strip().replace('\n',' ') for row in matrix[:first_data] if j<len(row) and row[j] is not None) for j in range(1,width)}
                code_columns={j for j,literal in headers.items() if 'октмо' in literal.lower()}
                letter_code_column=letter_code_column or any(len(row)>1 and str(row[0]).strip().lower() in {'а','a'} and str(row[1]).strip().lower() in {'б','b'} for row in matrix[:first_data])
                if letter_code_column:code_columns.add(1)
                book_id=hashlib.sha256((relative+'|'+sheet.title).encode()).hexdigest()[:20]
                books.append({'book_id':book_id,'source_file':relative,'sheet':sheet.title,'metric_number':metric,
                    'source_region_literal':path.relative_to(root).parts[2], 'source_title':str(matrix[0][0]) if matrix and matrix[0] else '',
                    'observation_year':year,'header_year_candidates':'|'.join(map(str,years)),
                    'settlement_header_excel_row':header+1,'first_data_excel_row':first_data+1,
                    'header_by_column':json.dumps(headers,ensure_ascii=False),'primary_sha256':sha(path)})
                current_district=''
                for i,row in enumerate(matrix[first_data:],first_data+1):
                    if not row or row[0] is None:continue
                    name=str(row[0]).strip()
                    if not name or name in {'А','Б','A','B'} or name.startswith(('...','…','Данные','Примечание','*')):continue
                    # A source footnote is not an enormous place name.
                    if len(name)>180:continue
                    district_heading=bool(re.search(r'\b(?:район|округ)\b',name,re.I)) and not re.match(r'^(?:г|д|с|п|пгт|рп)\.?\s+',name,re.I)
                    if district_heading:current_district=name
                    raw_name=name
                    embedded_code=re.match(r'^(\d{11})\s*[-–]\s*(.+)$',name)
                    raw_code=next((str(row[j]).strip() for j in sorted(code_columns) if j<len(row) and row[j] is not None),'')
                    if embedded_code:
                        raw_code=embedded_code.group(1);name=embedded_code.group(2).strip()
                    code=raw_code.zfill(11) if re.fullmatch(r'\d{10,11}',raw_code) else raw_code if re.fullmatch(r'\d{8}',raw_code) else ''
                    row_id=hashlib.sha256(f'{book_id}|{i}'.encode()).hexdigest()[:24]
                    rows.append({'source_row_id':row_id,'book_id':book_id,'source_file':relative,'source_excel_row':i,
                        'source_sheet':sheet.title,'source_region_literal':path.relative_to(root).parts[2],
                        'settlement':name,'source_settlement_literal':raw_name,'source_oktmo':code,'raw_source_code':raw_code,'metric_number':metric,
                        'source_district_literal':current_district,
                        'source_row_grain':'administrative_header_or_aggregate' if district_heading else 'named_settlement',
                        'observation_year':year,'publication_edition_year':2025,'column_semantics_verified':True,
                        'source_id':'ROSSTAT_NP2025_PRIMARY_BOOKS'})
                    for j in range(1,len(row)):
                        if j in code_columns:continue
                        raw=row[j]
                        if raw is None and not headers.get(j):continue
                        value,status=primary_cell(raw)
                        total=(j==min(k for k in range(1,len(row)) if k not in code_columns)) and metric not in {42,47}
                        cells.append({'source_row_id':row_id,'source_column':get_column_letter(j+1),'column_index':j,
                            'source_header_literal':headers.get(j,''),'reported_value':value,'raw_value_literal':None if raw is None else str(raw),
                            'value_status':status,'is_total_column':total})
            wb.close()
        except Exception as e:errors.append({'path':relative,'error':str(e)})
    return pd.DataFrame(books),pd.DataFrame(rows),pd.DataFrame(cells),errors


def classify_primary_grain(source_rows):
    """Keep a city/municipality heading apart from its immediately nested NP."""
    rows=source_rows.copy()
    rows['source_row_grain']=rows.source_row_grain.replace({'administrative_header_or_aggregate':'municipal_aggregate'})
    rows['source_grain_basis']='original_named_row'
    extended_aggregate=rows.raw_source_code.astype(str).str.fullmatch(r'\d{8}000')
    rows.loc[extended_aggregate,'source_row_grain']='city_aggregate'
    rows.loc[extended_aggregate,'source_grain_basis']='explicit_extended_municipal_code_ending000'
    order=rows.sort_values(['book_id','source_excel_row'])
    names=order.settlement.map(normalize_name)
    typed=order.settlement.str.match(r'^(?:г|с|д|п|пгт|рп|аул|х)\.?\s+',case=False)
    paired=(order.book_id.eq(order.book_id.shift(-1))&names.eq(names.shift(-1))&
        (order.source_excel_row.shift(-1)-order.source_excel_row).le(2)&~typed&typed.shift(-1,fill_value=False))
    rows.loc[order.index[paired],'source_row_grain']='city_aggregate'
    rows.loc[order.index[paired],'source_grain_basis']='same_name_heading_immediately_before_explicit_settlement_row'
    return rows


def attach_primary(primary_root,reference,out,*,cached_parse=False):
    if cached_parse:
        books=pd.read_parquet(out/'primary_books.parquet')
        previous=pd.read_parquet(out/'primary_settlement_rows.parquet')
        original_columns=['source_row_id','book_id','source_file','source_excel_row','source_sheet','source_region_literal',
            'settlement','source_settlement_literal','source_oktmo','raw_source_code','metric_number','source_district_literal',
            'source_row_grain','observation_year','publication_edition_year','column_semantics_verified','source_id']
        source_rows=previous[original_columns].copy()
        cells=pd.read_parquet(out/'primary_metric_cells.parquet')
        previous_summary=json.loads((out/'primary_summary.json').read_text())
        errors=previous_summary['book_errors']+[{'path':x['path'],'sheet':x['sheet'],'error':'settlement_header_not_identified'} for x in previous_summary.get('non_data_sheet_review',[])]
    else:books,source_rows,cells,errors=parse_primary_books(primary_root)
    # Inspect the few extra Excel sheets rather than reporting blank/helper
    # tabs as missing measurement data.
    import openpyxl
    data_errors=[];nondata=[]
    for error in errors:
        if error.get('error')!='settlement_header_not_identified':data_errors.append(error);continue
        workbook=openpyxl.load_workbook(Path(primary_root)/error['path'],read_only=True,data_only=True)
        matrix=list(workbook[error['sheet']].iter_rows(values_only=True));workbook.close()
        literals=[str(v).strip() for row in matrix for v in row if v is not None and str(v).strip()]
        if not literals:reason='empty_extra_worksheet'
        elif literals[0]=='ОКТМО' and all(re.fullmatch(r'\d{8,11}',v) for v in literals[1:]):reason='code_only_helper_sheet_no_measurement'
        else:data_errors.append(error);continue
        nondata.append({'path':error['path'],'sheet':error['sheet'],'reason':reason,'nonempty_cells':len(literals)})
    errors=data_errors
    source_rows=classify_primary_grain(source_rows)
    if len(source_rows)==0:raise ValueError('No primary Rosstat rows parsed')
    geokeys=['source_region_literal','settlement','source_oktmo','source_district_literal','source_row_grain']
    unique=source_rows[geokeys].drop_duplicates().copy();unique['activity_address']=''
    matches=match_addresses(unique,reference,address_col='activity_address',region_col='source_region_literal',name_col='settlement',oktmo_col='source_oktmo',district_col='source_district_literal')
    joined=source_rows.merge(matches.drop(columns='activity_address'),on=geokeys,how='left',validate='many_to_one')
    # Store long numeric cells without repeated prose/coordinates. These can be
    # joined by source_row_id; no totals/components are summed together.
    save(books,out/'primary_books')
    save(joined,out/'primary_settlement_rows')
    cells.to_parquet(out/'primary_metric_cells.parquet',index=False,compression='zstd')
    combined=cells.merge(joined,on='source_row_id',validate='many_to_one')
    workforce=combined[(combined.metric_number.eq(9))&combined.is_total_column].copy()
    workforce['value']=workforce.reported_value
    workforce['partial_organisation_workforce_value']=workforce.reported_value
    workforce['year']=workforce.observation_year
    workforce['settlement_jobs']=np.nan
    workforce['headcount_scope']='annual_average_organisation_workers_excluding_external_parttimers_and_small_business'
    save(workforce,out/'partial_workforce')
    # Explicit total seats, pupils, staff posts and persons retain different
    # units. Attending children do not measure residential potential demand.
    relevant=combined[combined.metric_number.isin([39,40,44,45,46])&combined.is_total_column].copy()
    relevant['measure']=relevant.metric_number.map({39:'preschool_places_total',40:'preschool_enrolled_children_total',44:'preschool_staff_posts_total',45:'preschool_filled_posts_total',46:'preschool_workers_total'})
    relevant['unit']=relevant.metric_number.map({39:'places',40:'children_enrolled',44:'staff_posts',45:'filled_posts',46:'persons_excluding_external_parttimers_and_civil_contracts'})
    relevant['is_residential_demand']=False
    save(relevant,out/'preschool_primary_totals')
    eligible=relevant[relevant.match_status.eq('matched')&relevant.observation_year.notna()&relevant.value_status.eq('observed')&relevant.reported_value.gt(0)]
    supplies=[]
    known_groups={key:g for key,g in relevant[relevant.value_status.eq('observed')&relevant.source_record_id.notna()].groupby(['source_record_id','observation_year'])}
    for (sid,year),g in eligible.groupby(['source_record_id','observation_year']):
        all_known=known_groups[(sid,year)]
        cap=all_known[all_known.metric_number.eq(39)].reported_value.unique()
        enrolled=all_known[all_known.metric_number.eq(40)].reported_value.unique()
        posts=all_known[all_known.metric_number.eq(44)].reported_value.unique()
        workers=all_known[all_known.metric_number.eq(46)].reported_value.unique()
        supplies.append({'provider_id':f'rosstat_np{int(year)}:preschool:'+sid,'source_record_id':sid,'service_code':'preschool_official',
            'supply_status':'observed','supply_basis':'positive_named_preschool_total_in_primary_workbook_with_explicit_year',
            'source_id':'ROSSTAT_NP2025_PRIMARY_BOOKS','observation_year':int(year),'publication_edition_year':2025,
            'capacity_total':float(cap[0]) if len(cap)==1 else np.nan,'capacity_unit':'total_preschool_places_not_vacant_places',
            'reported_enrolled_children':float(enrolled[0]) if len(enrolled)==1 else np.nan,
            'reported_staff_posts':float(posts[0]) if len(posts)==1 else np.nan,
            'reported_preschool_workers':float(workers[0]) if len(workers)==1 else np.nan,
            'capacity_conflicting_source_totals':len(cap)>1,'known_positive_source_rows':len(g),
            'local_residential_demand_known':False,'capacity_feasibility':'not_estimated',
            'institution_inventory_is_complete':False,'geometry_role':'settlement_aggregate_not_institution_entrance','established_flow':False})
    supply=pd.DataFrame(supplies)
    save(supply,out/'official_preschool_supply')
    summary={'primary_books':len(books),'regions':int(books.source_region_literal.nunique()),'source_rows':len(joined),'numeric_cell_rows':len(cells),
        'book_errors':errors,'non_data_sheet_review':nondata,'book_count_by_metric':books.metric_number.value_counts().sort_index().to_dict(),
        'book_periods':json.loads(books.groupby(['metric_number','observation_year'],dropna=False).size().rename('books').reset_index().to_json(orient='records')),
        'matched_source_rows':int(joined.match_status.eq('matched').sum()),'official_preschool_supply_rows':len(supply),
        'city_aggregate_source_rows_held':int(joined.source_row_grain.eq('city_aggregate').sum()),
        'preschool_known_capacity_rows':int(supply.capacity_total.notna().sum()),
        'partial_workforce_rows':len(workforce),'partial_workforce_numeric_observations':int(workforce.value_status.eq('observed').sum()),
        'partial_workforce_matched_numeric':int((workforce.value_status.eq('observed')&workforce.match_status.eq('matched')).sum()),
        'all_columns_preserved':True,'totals_not_summed_with_components':True}
    (out/'primary_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=str)+'\n')
    return supply,summary


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda:f.read(1<<20),b''):h.update(b)
    return h.hexdigest()


def save(frame,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    parquet=path.with_suffix('.parquet');csv=path.with_suffix('.csv')
    ptmp=parquet.with_name(parquet.name+'.partial');ctmp=csv.with_name(csv.name+'.partial')
    frame.to_parquet(ptmp,index=False,compression='zstd');ptmp.replace(parquet)
    frame.to_csv(ctmp,index=False);ctmp.replace(csv)


def read_rosstat(raw):
    frames=[];inputs=[]
    for kind in ['headcount','preschool','preschool_children']:
        for path in sorted((raw/kind).glob('*.csv')):
            d=pd.read_csv(path,dtype=str,keep_default_na=False)
            required={'region','settlement','oktmo','value','value_status'}
            if not required<=set(d):raise ValueError(f'CSV schema missing required fields: {path}')
            d['source_file']=path.relative_to(raw).as_posix()
            d['source_row']=np.arange(2,len(d)+2)
            d['source_kind']=kind
            d['source_region_literal']=d.region
            d['source_settlement_literal']=d.settlement
            d['source_raw_oktmo_field']=d.oktmo
            valid_code=d.oktmo.str.fullmatch(r'\d{11}')
            if kind=='headcount':valid_code|=d.oktmo.str.fullmatch(r'\d{8}')
            d['source_oktmo']=d.oktmo.where(valid_code,None)
            d['source_code_verified_by_workbook_header']=False
            d['activity_address']=d.settlement
            d['reported_value']=pd.to_numeric(d.value.where(d.value.ne('')),errors='coerce')
            if d.loc[d.value_status.eq('observed'),'reported_value'].isna().any():raise ValueError('Non-numeric observed value')
            d['publication_edition_year']=2025
            d['observation_year']=None
            d['column_semantics_verified']=False
            d['indicator']=d.get('indicator',pd.Series('headcount_last_source_cell' if kind=='headcount' else 'preschool_children_last_source_cell',index=d.index))
            d['source_id']='ROSSTAT_NP2025_GROK_CSV_'+kind.upper()
            d['source_commit']=SOURCE_COMMIT
            d['raw_value_status']=d.value_status
            d['value_scope']='source_last_nonempty_cell_unverified_subcategory_and_year'
            d['jobs_observed_at_settlement']=False
            frames.append(d)
            inputs.append({'path':path.relative_to(raw).as_posix(),'bytes':path.stat().st_size,'sha256':sha(path),'rows':len(d)})
    if not frames:raise ValueError('No Rosstat source CSVs')
    return pd.concat(frames,ignore_index=True),inputs


def attach_rosstat(raw,reference,out):
    frame,inputs=read_rosstat(raw)
    # Codes syntactically valid in the CSV are still checked against name/region.
    # Values accidentally exported under `oktmo` never pass this screen.
    matched=match_addresses(frame,reference,address_col='activity_address',region_col='source_region_literal',name_col='source_settlement_literal',oktmo_col='source_oktmo')
    save(matched,out/'rosstat_all_rows')
    matched['positive_reported_value']=matched.raw_value_status.eq('observed')&matched.reported_value.gt(0)
    supply=[]
    for sid,g in matched[matched.match_status.eq('matched')&matched.positive_reported_value&matched.source_kind.ne('headcount')].groupby('source_record_id'):
        supply.append({'provider_id':'rosstat_np2025:preschool:'+sid,'source_record_id':sid,'service_code':'preschool',
            'supply_status':'proxy','supply_basis':'positive_preschool_indicator_last_cell_unverified_year_and_subcategory',
            'source_id':'ROSSTAT_NP2025_GROK_CSV_EDUCATION','observation_year':None,'publication_edition_year':2025,
            'capacity_total':np.nan,'capacity_unit':None,'known_positive_source_rows':len(g),
            'reported_capacity_not_admitted':True,'institution_inventory_is_complete':False,
            'geometry_role':'settlement_aggregate_not_institution_entrance','established_flow':False})
    # Even positive exported cells cannot establish current service presence
    # before the workbook's reference year and column category are inspected.
    # Keep candidate evidence outside the routing supply pool.
    candidates=pd.DataFrame(supply)
    if len(candidates):
        candidates['admitted_to_supply']=False
        candidates['exclusion_reason']='primary_workbook_header_year_subcategory_not_verified'
    save(candidates,out/'rosstat_positive_candidates_not_admitted')
    work=matched[matched.source_kind.eq('headcount')].copy()
    work['partial_organisation_workforce_value']=work.reported_value
    work['headcount_scope']='organisations_excluding_small_business_and_external_parttimers; final_source_column_unverified'
    work['settlement_jobs']=np.nan
    save(work,out/'partial_workforce_rows')
    coverage=matched.groupby(['source_kind','source_region_literal','raw_value_status','match_status'],dropna=False).size().rename('rows').reset_index()
    save(coverage,out/'rosstat_coverage')
    summary={'files':len(inputs),'rows':len(frame),'source_files_by_kind':{kind:sum(p['path'].startswith(kind+'/') for p in inputs) for kind in frame.source_kind.unique()},
        'rows_by_kind':frame.source_kind.value_counts().to_dict(),'raw_observed_rows_by_kind':frame[frame.value_status.eq('observed')].source_kind.value_counts().to_dict(),
        'matched_rows_by_kind':matched[matched.match_status.eq('matched')].source_kind.value_counts().to_dict(),
        'positive_export_candidate_settlements_not_admitted':len(supply),'reference_year_established':False,
        'source_header_error':'The source parser labels the first nonempty cell oktmo; many education books have no settlement-code column. Invalid keys are never used.',
        'source_value_limit':'The exporter kept the last nonempty cell without source header/year/subcategory. Values retained but capacities and workforce are not certified.',
        'original_rar':{'url':SOURCE_RAR_URL,'sha256_from_grok_build_log':SOURCE_RAR_SHA,'bytes_from_grok_build_log':230053990,'primary_header_freshly_rechecked':False},'inputs':inputs}
    return pd.DataFrame(columns=['provider_id','source_record_id','service_code','supply_status']),summary


def explicit_date_interval(value):
    """Only literal operational start/end dates; no construction-date inference."""
    if value is None or pd.isna(value) or not str(value).strip():return None
    literal=str(value).strip().lower()
    literal=re.sub(r'\s*(?:г\.|года|г)$','',literal).strip()
    months={'января':1,'февраля':2,'марта':3,'апреля':4,'мая':5,'июня':6,'июля':7,'августа':8,'сентября':9,'октября':10,'ноября':11,'декабря':12}
    words=re.fullmatch(r'(\d{1,2})\s+([а-я]+)\s+(\d{4})',literal)
    if words and words[2] in months:literal=f'{int(words[3]):04d}-{months[words[2]]:02d}-{int(words[1]):02d}'
    if re.fullmatch(r'\d{8}',literal):literal=f'{literal[:4]}-{literal[4:6]}-{literal[6:]}'
    if re.fullmatch(r'\d{4}\.\d{2}\.\d{2}',literal):literal=literal.replace('.','-')
    local_day=re.fullmatch(r'(\d{2})[.-](\d{2})[.-](\d{4})',literal)
    if local_day:literal=f'{local_day[3]}-{local_day[2]}-{local_day[1]}'
    if not re.fullmatch(r'\d{4}(?:-\d{2}(?:-\d{2})?)?',literal):return None
    parts=list(map(int,literal.split('-')))
    try:
        if len(parts)==1:return date(parts[0],1,1),date(parts[0],12,31),'year'
        if len(parts)==2:return date(*parts,1),date(*parts,calendar.monthrange(*parts)[1]),'month'
        day=date(*parts);return day,day,'day'
    except ValueError:return None


def osm_temporal_status(start,end,snapshot='2021-01-01'):
    """Exclude only dates whose entire precision interval contradicts snapshot."""
    snap=date.fromisoformat(str(snapshot));a=explicit_date_interval(start);b=explicit_date_interval(end)
    if b and b[1]<snap:return 'ended_before_snapshot'
    if a and a[0]>snap:return 'started_after_snapshot'
    if any(x is not None and not pd.isna(x) and str(x).strip() and parsed is None for x,parsed in [(start,a),(end,b)]):return 'unresolved_date_literal'
    if a or b:return 'no_demonstrated_temporal_conflict'
    return 'no_explicit_operational_dates'


def osm_lifecycle_reason(tags_json):
    tags=json.loads(tags_json) if isinstance(tags_json,str) and tags_json else {}
    return '|'.join(key+'=yes' for key in ['disused','abandoned','ruins'] if str(tags.get(key,'')).strip().lower()=='yes')


def osm_supply(osm_path,legacy_path,reference,out):
    objects=pd.read_parquet(osm_path)
    legacy=pd.read_parquet(legacy_path)
    # The old object key is a legacy node ID, not the current census root.
    bridge=legacy[legacy.root.notna()&legacy.source_record_id_2021.notna()].merge(
        reference[reference.source_record_id.notna()][['root','source_record_id','name','region','lat','lon']],
        left_on=['root','source_record_id_2021'],right_on=['root','source_record_id'],how='inner',validate='one_to_one',suffixes=('_legacy','_census'))
    if not bridge.id.is_unique:raise ValueError('Duplicate legacy identity')
    rows=objects[objects.education_group.isin(['school','preschool','college','university'])].copy()
    # Previously detected node/way duplicates remain in raw source, but not in
    # the accepted presence pool. No object count is used as model weight.
    rows=rows[rows.possible_duplicate_of.isna()].copy()
    rows['legacy_id']=rows.baseline_uid.fillna(rows.candidate_uid)
    linked=rows.merge(bridge[['id','source_record_id','root','name_census','region_census','lat_census','lon_census']],left_on='legacy_id',right_on='id',how='left',validate='many_to_one')
    explicit=linked.link_status.eq('explicit_address_locality_unique_within5km')
    address_agrees=pd.Series([normalize_name(a)==normalize_name(b) for a,b in zip(linked.source_address_locality.fillna(''),linked.name_census.fillna(''))],index=linked.index)
    linked['current_identity_status']=np.where(linked.source_record_id.isna(),'unmatched_current_census_identity',np.where(explicit&address_agrees,'explicit_locality_and_exact_legacy_census_bridge',np.where(linked.link_status.eq('spatial_candidate_within2_5km_separated_not_identity_proof'),'spatial_proxy_not_identity_proof','held_address_or_source_status')))
    linked['supply_status']=np.where(linked.current_identity_status.eq('explicit_locality_and_exact_legacy_census_bridge'),'observed',np.where(linked.current_identity_status.eq('spatial_proxy_not_identity_proof'),'proxy','unknown'))
    linked['supply_status_before_temporal_check']=linked.supply_status
    linked['temporal_status']=[osm_temporal_status(a,b,s) for a,b,s in zip(linked.source_start_date,linked.source_end_date,linked.source_snapshot)]
    linked['temporal_eligible']=~linked.temporal_status.isin(['ended_before_snapshot','started_after_snapshot'])
    linked['lifecycle_exclusion_reason']=linked.tags_json.map(osm_lifecycle_reason)
    linked['lifecycle_eligible']=linked.lifecycle_exclusion_reason.eq('')
    linked.loc[~linked.lifecycle_eligible,'supply_status']='excluded_lifecycle'
    linked.loc[~linked.temporal_eligible,'supply_status']='excluded_temporal'
    linked['service_code']=linked.education_group.map({'school':'school_general','preschool':'preschool','college':'vocational_education','university':'higher_education'})
    linked['source_id']='OSM_GEOFABRIK_EDUCATION_20210101'
    linked['observation_year']=2021
    linked['capacity_raw_not_normalized']=linked.source_capacity
    linked['grades_known']=False
    linked['geometry_role_scope']='OSM_object_location; settlement_identity_or_proxy_kept_separately'
    save(linked,out/'osm_object_links')
    save(linked[linked.temporal_status.ne('no_explicit_operational_dates')],out/'osm_operational_date_audit')
    save(linked[~linked.lifecycle_eligible],out/'osm_inactive_lifecycle_audit')
    # Sparse explicit ISCED tags remain a separate evidence layer. Do not guess
    # school grades for the rest of the country or assume that a tag3 implies
    # that levels1 and2 are also offered.
    tagged=[]
    for r in linked[linked.education_group.eq('school')&linked.source_isced_level.notna()].itertuples():
        literal=str(r.source_isced_level).strip()
        levels=set()
        if re.fullmatch(r'[0-3](?:[;,][0-3])*',literal):levels=set(map(int,re.split('[;,]',literal)))
        elif re.fullmatch(r'[0-3]-[0-3]',literal):
            first,last=map(int,literal.split('-'))
            if first<=last:levels=set(range(first,last+1))
        for level in sorted(levels&{1,2,3}):
            tagged.append({'site_id':r.site_id,'source_record_id':r.source_record_id,'supply_status':r.supply_status,
                'source_isced_level_literal':literal,'isced_level':level,
                'service_code':{1:'school_primary_isced1',2:'school_lower_secondary_isced2',3:'school_upper_secondary_isced3'}[level],
                'grade_basis':'literal_OSM_ISCED_tag_not_Russian_grade_inference','source_id':'OSM_GEOFABRIK_EDUCATION_20210101',
                'observation_year':2021,'admitted_to_current_routing_supply':False,'capacity_total':np.nan})
    save(pd.DataFrame(tagged),out/'tagged_school_orders_not_in_main_supply')
    supply=[]
    admissible=linked[linked.source_record_id.notna()&linked.supply_status.isin(['observed','proxy'])]
    for (sid,service),g in admissible.groupby(['source_record_id','service_code']):
        strict=g[g.supply_status.eq('observed')]
        chosen=strict if len(strict) else g
        supply.append({'provider_id':'osm2021:'+service+':'+sid,'source_record_id':sid,'service_code':service,
            'supply_status':'observed' if len(strict) else 'proxy','supply_basis':'explicit_address_locality_and_exact_census_bridge' if len(strict) else 'spatial_proxy_not_identity_proof',
            'source_id':'OSM_GEOFABRIK_EDUCATION_20210101','observation_year':2021,'publication_edition_year':2021,
            'capacity_total':np.nan,'capacity_unit':None,'known_positive_source_rows':len(chosen),
            'reported_capacity_not_admitted':True,'institution_inventory_is_complete':False,
            'geometry_role':'settlement_aggregate_not_institution_entrance','established_flow':False})
    frame=pd.DataFrame(supply)
    return frame,{'source_objects':len(objects),'selected_nonduplicate_objects':len(rows),'object_status_counts':linked.current_identity_status.value_counts().to_dict(),
        'temporal_status_counts':linked.temporal_status.value_counts().to_dict(),
        'lifecycle_excluded_objects':int((~linked.lifecycle_eligible).sum()),
        'temporal_exclusions':linked.loc[~linked.temporal_eligible,['site_id','service_code','source_start_date','source_end_date','source_snapshot','temporal_status','source_record_id','supply_status_before_temporal_check']].to_dict('records'),
        'settlement_service_counts':frame.groupby(['service_code','supply_status']).size().rename('rows').reset_index().to_dict('records'),
        'explicit_tagged_school_order_rows_separate_layer':len(tagged),
        'source_sha256':sha(osm_path),'legacy_bridge_sha256':sha(legacy_path),'source_geometry_is_dated_measurement':False}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--raw',type=Path,help='Optional lossy Grok CSVs retained for diagnostics, never a capacity source')
    p.add_argument('--primary-rar',type=Path,help='Pinned primary Rosstat archive; reads real headers, totals and components')
    p.add_argument('--population',type=Path,required=True)
    p.add_argument('--osm',type=Path,required=True);p.add_argument('--legacy',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();a.output.mkdir(parents=True,exist_ok=True)
    reference=build_reference(a.population)
    if a.primary_rar:
        extraction=extract_primary(a.primary_rar,a.output/'primary')
        official,rosstat=attach_primary(a.output/'primary',reference,a.output)
        rosstat['extraction']=extraction
    elif a.raw:
        official,rosstat=attach_rosstat(a.raw,reference,a.output)
    else:
        p.error('Provide the primary archive or diagnostic CSV directory')
    osm,osmsummary=osm_supply(a.osm,a.legacy,reference,a.output)
    supply=pd.concat([official,osm],ignore_index=True)
    if supply.provider_id.duplicated().any():raise ValueError('Repeated provider identifier')
    supply_name='supply_final_corrected' if a.primary_rar else 'supply'
    save(supply,a.output/supply_name)
    summary={'rosstat':rosstat,'osm':osmsummary,'supply_rows':len(supply),'matched_distinct_settlements':int(supply.source_record_id.nunique()),
        'supply_counts':supply.groupby(['source_id','service_code','supply_status']).size().rename('rows').reset_index().to_dict('records'),
        'routing_supply_file':supply_name+'.parquet','population_sha256':sha(a.population),'limitations':['Presence is not a journey or complete local supply.',
            'No school grade guessed from the general OSM school label.','Primary totals and subcategories are separate; capacities and workforce are not automatic model weights.',
            'Unknown and suppressed values are not zero.','Preschool is distinct from general school and higher/vocational education.']}
    (a.output/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in summary.items() if k!='rosstat'},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
